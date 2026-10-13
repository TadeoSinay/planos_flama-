"""MRP de insumos de FLAMA (salida/mrp/FLAMA_MRP.xlsx), compacto: sólo números y fórmulas.

Cada fila lleva un código (Ref.) que remite al Word «Justificación_MRP», donde está de dónde sale cada valor, con su
captura. No hay supuestos: todo sale de los procesos de manuales, carros y recargas, del BOM, de las Proyecciones, de
los dos Dimensionamientos, de la planilla de MP, de las cotizaciones o de fichas técnicas. Lo que no tiene consumo
documentado no entra al MRP y se compra por punto de pedido (lista al final del Word).

Hojas
  DEMANDA          unidades por año de cada producto (Indexado Dim Técnico y Reservas y Sustitutos); las recargas son
                   las de la planta industrial
  ESTACIONALIDAD   perfiles y producción mensual del año 2035, como el Dimensionamiento Técnico FLAMA
  MANUALES, RODANTES, RECARGAS, REVENDIDOS
                   arriba la demanda de cada producto; abajo el consumo por unidad de cada material (una columna por
                   producto) y lo que hay que comprar por año (consumo por unidad × demanda)
  CONSOLIDADO      una fila por material: lo que hay que comprar sumando las cuatro líneas
  POLITICA_STOCK   año 2035: demanda de diseño, stock de seguridad, punto de pedido, lote, stock máximo y posiciones
  PROY_INDEXADO, PROY_RESERVAS   copias de las hojas de las Proyecciones (si se pegan valores nuevos, todo se recalcula)
  PROY_USADOS      usados y recargas por año y rubro (suma) y debajo las filas copiadas de las siete hojas (sin 2025)

Uso: python generar.py --mrp-insumos <Proyecciones.xlsx>   (necesita salida/bom/FLAMA_BOM.xlsx ya recalculado)
"""
import json
import os

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from .bom_simple import BLANCO, GRIS_C, GRIS_O, NEGRO, recalcular

AÑOS = list(range(2026, 2036))
MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
LINEAS = ("MANUALES", "RODANTES", "RECARGAS", "REVENDIDOS")
AR = "Arial"
FINO = Side(style="thin", color="BFBFBF")

# ------------------------------------------------------------------ productos
FAB = [("ABC_1kg", "ABC1", "ABC 1 kg"), ("ABC_2.5kg", "ABC2.5", "ABC 2,5 kg"), ("ABC_5kg", "ABC5", "ABC 5 kg"),
       ("ABC_10kg", "ABC10", "ABC 10 kg"), ("ABC_25kg", "ABC25", "ABC 25 kg"), ("ABC_50kg", "ABC50", "ABC 50 kg"),
       ("ABC_70kg", "ABC70", "ABC 70 kg"), ("ABC_100kg", "ABC100", "ABC 100 kg")]
MAN, ROD = FAB[:4], FAB[4:]
REV = [("AGUA_10l", "Agua 10 l"), ("AFFF_10l", "AFFF 10 l"), ("AFFF_50l", "AFFF 50 l"), ("BC_5kg", "BC 5 kg"),
       ("SALESK_6l", "Clase K 6 l"), ("CO2_2kg", "CO₂ 2 kg"), ("CO2_5kg", "CO₂ 5 kg"), ("HCFC-HFC_5kg", "HCFC 5 kg"),
       ("CLASED_9l", "Clase D 9 l")]
NOMBRE = {c: n for c, _, n in FAB} | dict(REV)
# filas de las Proyecciones (Indexado Dim Técnico) → código del producto, en el orden de las recargas
ETIQ = {"ABC 1 kg": "ABC_1kg", "ABC 2,5 kg": "ABC_2.5kg", "ABC 5 kg": "ABC_5kg", "ABC 10 kg": "ABC_10kg",
        "Agua 10 l": "AGUA_10l", "AFFF 10 l": "AFFF_10l", "AFFF 50 l s/ruedas": "AFFF_50l", "BC 5 kg": "BC_5kg",
        "Sales K 6 l": "SALESK_6l", "CO₂ 2 kg": "CO2_2kg", "CO₂ 5 kg": "CO2_5kg", "HCFC/HFC 5 kg": "HCFC-HFC_5kg",
        "Clase D 9 l": "CLASED_9l", "ABC 25 kg s/ruedas": "ABC_25kg", "ABC 50 kg s/ruedas": "ABC_50kg",
        "ABC 70 kg s/ruedas": "ABC_70kg", "ABC 100 kg s/ruedas": "ABC_100kg"}
INV = {v: k for k, v in ETIQ.items()}
REC = list(ETIQ.values())
# pool de sustitutos (Reservas y Sustitutos): fila del pool → tipo que cubre
POOL = {"ABC 5 kg (cubre ABC 1; 2,5 y 5 kg)": "ABC_5kg", "ABC 10 kg": "ABC_10kg", "Agua 10 l": "AGUA_10l",
        "AFFF 10 l": "AFFF_10l", "AFFF 50 l s/ruedas": "AFFF_50l", "BC 5 kg": "BC_5kg", "Sales K 6 l": "SALESK_6l",
        "CO₂ 5 kg (cubre CO₂ 2 y 5 kg)": "CO2_5kg", "HCFC/HFC 5 kg": "HCFC-HFC_5kg", "Clase D 9 l": "CLASED_9l",
        "ABC 50 kg s/ruedas (cubre 25 y 50 kg)": "ABC_50kg", "ABC 100 kg s/ruedas (cubre 70 y 100 kg)": "ABC_100kg"}
# zonas que atiende la planta industrial (Dimensionamiento Técnico V2, Almacen PT, filas 28 a 34); su parte de las
# recargas sale de la distribución de recargas por zona del Indexado, año por año
ZONAS_PI_NOM = ("CABA Sur", "PBA Sur 1er", "PBA Sur 2do", "PBA Sur 3er")
# hojas de las Proyecciones que separan las recargas de equipos usados (columnas Año, Mat Usados y RECARGAS TOT)
HOJAS_USADOS = ("SALUD", "EDUCACION", "INDUSTRIASALM", "HOTELES")
# hojas de proyección que no separan los usados: (hoja, columna del año, columna de recargas), base 0
HOJAS_SIN_USADOS = (("SUPERMERCADOS", 1, 10), ("AUTOMOTOR", 0, 12), ("EDIFICIOS", 1, 10))
# perfiles de estacionalidad: Dimensionamiento Técnico FLAMA, DATOS-BASE filas 102 a 104; recargas /12 (filas 135+)
PERFILES = [("Automovilístico (1 kg)", [.12, .09, .08, .075, .07, .07, .075, .075, .075, .075, .075, .12]),
            ("Manual comercial (2,5 / 5 / 10 kg)", [.10, .08, .085, .085, .08, .075, .08, .08, .08, .075, .07, .11]),
            ("Industrial s/ruedas (25 – 100 kg)", [.085, .075, .09, .09, .09, .08, .09, .09, .085, .075, .07, .08]),
            ("Recargas (sin estacionalidad)", [1 / 12] * 12)]


def perfil(c, linea):
    if linea == "RECARGAS":
        return 3
    if c == "ABC_1kg":
        return 0
    return 2 if c in ("ABC_25kg", "ABC_50kg", "ABC_70kg", "ABC_100kg", "AFFF_50l") else 1


# ------------------------------------------------------------------ catálogo
# planilla de MP («MP Abastecimiento», opción D): lead time, días entre pedidos (T), stock de seguridad (días) y
# unidad de carga por posición, por código de la planilla
PLAN = {"MP-01": (5, 5, 5, 61), "MP-02": (10, 20, 10, 136), "MP-03": (10, 20, 10, 158), "MP-04": (10, 20, 10, 56),
        "MP-05": (10, 5, 10, 1000), "MP-06": (10, 10, 10, 492), "MP-07": (10, 5, 10, 615), "MP-08": (10, 5, 10, 492),
        "MP-09": (10, 20, 10, 738), "MP-10": (10, 20, 10, 498), "MP-11": (10, 20, 10, 34), "MP-12": (10, 20, 10, 22),
        "MP-13": (10, 5, 10, 1000), "MP-15": (60, 40, 20, 200), "MP-16": (60, 40, 20, 160),
        "MP-17": (60, 40, 20, 120), "MP-18": (28, 20, 10, 12000), "MP-19": (28, 20, 10, 8000),
        "MP-20": (28, 40, 20, 400), "MP-21": (28, 40, 20, 100), "MP-22": (10, 60, 20, None),
        "MP-23": (5, 20, 5, 900), "MP-24": (5, 40, 10, 900), "MP-25": (1, 5, 2, 115.5), "MP-27": (5, 20, 10, 250),
        "MP-28": (7, 60, 20, 1000), "MP-31": (10, 20, 10, 600), "MP-33": (7, 5, 10, 2000), "MP-34": (1, 10, 3, 103),
        "MP-35": (7, 60, 20, 1000), "MP-36": (15, 120, 20, 1000), "MP-37": (2, 10, 5, 25), "MP-38": (30, 60, 20, 50),
        "MP-40": (7, 60, 20, 480), "MP-41": (7, 20, 10, 480), "MP-42": (20, 20, 15, 3000),
        "MP-43": (20, 20, 15, 1500), "MP-44": (20, 20, 10, 200), "MP-45": (20, 20, 15, 3000),
        "MP-46": (15, 20, 10, 6000), "MP-47": (15, 20, 10, 200),
        "MP-48": (15, 20, 10, 1500), "MP-49": (15, 20, 10, 120), "MP-50": (10, 20, 10, 20000),
        "MP-51": (10, 20, 10, 20000), "MP-52": (15, 20, 10, 3000), "MP-54": (15, 20, 10, 144),
        "MP-55": (15, 60, 20, 2000), "MP-56": (10, 20, 10, 20000), "MP-57": (10, 20, 10, 20000),
        "MP-58": (10, 20, 10, 10000), "MP-59": (10, 10, 5, 600), "MP-60": (5, 10, 5, 15), "MP-62": (10, 20, 10, 40000),
        "MP-65": (5, 20, 10, 200)}

Z_CHM, Z_CANT, Z_CHR = "Chapa y fleje (manuales)", "Cantilever de caños y barras", "Calle de chapa (carros)"
Z_CASQ, Z_INR, Z_INM = "Rack de casquetes y cuplas (carros)", "Rack de insumos de montaje (carros)", \
    "Insumos de montaje (manuales)"
Z_POL, Z_GAS, Z_SOL = "Depósito de polvo químico", "Jaula de gases", "Insumos de soldadura y superficie"
Z_ID, Z_EMB, Z_REC = "Identificación", "Embalaje", "Insumos de recargas"
Z_PT, Z_SERV = "Almacén de producto terminado", "Servicio (no se almacena)"
# capacidad de las zonas que dimensionó el proceso de carros (puesto 1), en posiciones
CAPACIDAD = {Z_CHR: 3, Z_CASQ: 12, Z_INR: 8}


def I(cod, nom, um, alm, prov, uc, cont, plan, fuc, umpos=None, fpos=None):
    """Ítem: código, nombre, UM, almacén, proveedor, unidad de compra y contenido, código de la planilla de MP de la
    que sale el lead time / T / stock de seguridad, figura de la unidad de compra y UM por posición si no es la de la
    planilla."""
    lt, t, ss, pos = PLAN[plan]
    return dict(cod=cod, nom=nom, um=um, alm=alm, prov=prov, uc=uc, cont=cont, plan=plan, fuc=fuc, lt=lt, t=t, ss=ss,
                umpos=umpos if umpos is not None else pos, fpos=fpos or "X07")


# tubo pescante: aparte de la válvula, como la planilla de MP (MP-46 plástico 2,5 a 10 kg; MP-47 acero, carros);
# el del 1 kg viene en la válvula M-22. Dos proveedores que lo publican o lo fabrican, ninguno con plazo publicado
PESCA_PVC = "Alt. 1 Esinex (Billinghurst) / Alt. 2 Securitá"
PESCA_ACERO = "Alt. 1 Mozart (V. Alsina) / Alt. 2 Georgia (CABA)"
# revendidos que figuran en el catálogo de Maxiseguridad (W22)
REV_MAXI = ("AGUA_10l", "AFFF_10l", "AFFF_50l", "BC_5kg", "CO2_2kg")

ITEMS = [
    I("MP-01", "Caño c/costura Ø76,2 × 1,25 × 6 m (cuerpo 1 kg)", "barra", Z_CANT, "Metalpri / Prilux",
      "Atado de 16 barras", 16, "MP-01", "P01"),
    I("MP-02", "Hoja LAF 1,25 × 1220 × 2440 mm", "hoja", Z_CHM, "Pradecon", "Paquete de 50 hojas", 50, "MP-02", "P01"),
    I("MP-03", "Hoja LAF 1,6 × 1000 × 2000 mm", "hoja", Z_CHM, "Pradecon", "Paquete de 80 hojas", 80, "MP-03", "P01"),
    I("MP-04", "Hoja LAF 2,0 × 1500 × 3000 mm", "hoja", Z_CHM, "Pradecon", "Paquete de 20 hojas", 20, "MP-04", "P01"),
    I("MP-05", "Fleje LAF 0,9 × 114 mm (cúpula 1 kg)", "kg", Z_CHM, "Pradecon", "Rollo de 1.900 kg", 1900, "MP-05",
      "Q01"),
    I("MP-13", "Fleje LAF 1,25 × 135 mm (fondo 1 kg)", "kg", Z_CHM, "Pradecon", "Rollo de 1.900 kg", 1900, "MP-13",
      "Q01"),
    I("MP-06", "Fleje LAF 1,25 × 200 mm (cúpula y fondo 2,5 kg)", "kg", Z_CHM, "Pacheco Chapas", "Rollo de 492 kg",
      492, "MP-06", "Q02"),
    I("MP-07", "Fleje LAF 1,6 × 250 mm (cúpula 5 kg)", "kg", Z_CHM, "Pacheco Chapas", "Rollo de 615 kg", 615, "MP-07",
      "Q02"),
    I("MP-08", "Fleje LAF 1,6 × 200 mm (fondo 5 kg)", "kg", Z_CHM, "Pacheco Chapas", "Rollo de 492 kg", 492, "MP-08",
      "Q02"),
    I("MP-09", "Fleje LAF 1,6 × 300 mm (cúpula 10 kg)", "kg", Z_CHM, "Pacheco Chapas", "Rollo de 590 kg", 590, "MP-09",
      "Q02"),
    I("MP-10", "Fleje LAF 2,0 × 249 mm (fondo 10 kg)", "kg", Z_CHM, "Pacheco Chapas", "Rollo de 498 kg", 498, "MP-10",
      "Q02"),
    I("MP-11", "Hoja LAC 3,2 × 1500 × 3000 mm", "hoja", Z_CHR, "Pradecon", "Paquete de 20 hojas", 20, "MP-11", "P07",
      40, "P07"),
    I("MP-12", "Hoja LAC 4,75 × 1500 × 3000 mm", "hoja", Z_CHR, "Pradecon", "Paquete de 15 hojas", 15, "MP-12", "P07",
      30, "P07"),
    I("MP-15", "Casquete 10\" × 3,2 mm (carro 25 kg)", "u", Z_CASQ, "Importado (China)", "Pallet de 400 u", 400, "MP-15",
      "P07", 400, "P07"),
    I("MP-16", "Casquete 12\" × 3,2 mm (carro 50 kg)", "u", Z_CASQ, "Importado (China)", "Pallet de 300 u", 300, "MP-16",
      "P07", 300, "P07"),
    I("MP-17", "Casquete 15\" × 4,75 mm (carros 70 y 100 kg)", "u", Z_CASQ, "Importado (China)", "Pallet de 150 u", 150,
      "MP-17", "P07", 150, "P07"),
    I("MP-18", "Asiento de válvula M22 × 1,5 (1 kg)", "u", Z_INM, "Eli-Met", "Caja de 1.000 u", 1000, "MP-18", "X07"),
    I("MP-19", "Asiento de válvula M30 × 1,5 (2,5 a 10 kg)", "u", Z_INM, "Eli-Met", "Caja de 1.000 u", 1000, "MP-19",
      "X07"),
    I("MP-20", "Cupla M30 × 1,5 (carros 25 y 50 kg)", "u", Z_CASQ, "Eli-Met", "Caja de 100 u", 100, "MP-20", "P07",
      600, "P07"),
    I("MP-21", "Cupla 2½\" BSP (carros 70 y 100 kg)", "u", Z_CASQ, "Eli-Met", "Caja de 40 u", 40, "MP-21", "P07",
      160, "P07"),
    I("MP-66", "Caño SAE 1010 1\" × 2,0 × 6 m (manija; caño del eje 25 kg)", "barra", Z_CANT, "Centrosider",
      "Barra de 6 m", 1, "MP-22", "P15", 60, "X07"),
    I("MP-67", "Caño SAE 1010 1¼\" × 2,0 × 6 m (caño del eje 50 a 100 kg)", "barra", Z_CANT, "Centrosider",
      "Barra de 6 m", 1, "MP-22", "P15", 60, "X07"),
    I("MP-68", "Barra trefilada SAE 1045 3/4\" × 4 m (eje 25 kg)", "barra", Z_CANT, "Vladimirsky", "Barra de 4 m",
      1, "MP-22", "P15", 60, "X07"),
    I("MP-69", "Barra trefilada SAE 1045 1\" × 6 m (eje 50 a 100 kg)", "barra", Z_CANT, "Vladimirsky",
      "Barra de 6 m", 1, "MP-22", "P15", 60, "X07"),
    I("MP-70", "Arandela de tope del eje", "u", Z_INR, "Würth Argentina", "Bolsa de 1.000 u", 1000, "MP-51", "X07"),
    I("MP-71", "Pasador partido del eje", "u", Z_INR, "Würth Argentina", "Bolsa de 1.000 u", 1000, "MP-51", "X07"),
    I("MP-54/300", "Rueda Ø300 × 50, buje 3/4\" (carro 25 kg)", "u", Z_INR, "Ruedar", "Pallet de 72 u", 72, "MP-54",
      "X07"),
    I("MP-54/350", "Rueda Ø350 × 60, buje 1\" (carro 50 kg)", "u", Z_INR, "Ruedar", "Pallet de 72 u", 72, "MP-54",
      "X07"),
    I("MP-54/400", "Rueda Ø400 × 60, buje 1\" (carros 70 y 100 kg)", "u", Z_INR, "Escanort", "Pallet de 72 u", 72,
      "MP-54", "X07"),
    I("MP-23", "Alambre MAG ER70S-6 Ø0,9 mm (manuales)", "kg", Z_SOL, "Soldanet", "Pallet de 900 kg", 900,
      "MP-23", "X07"),
    I("MP-24", "Alambre MAG ER70S-6 Ø1,0 mm (carros)", "kg", Z_INR, "Soldanet", "Pallet de 900 kg", 900,
      "MP-24", "X07"),
    I("MP-25", "Gas MAG Ar + 8 % CO₂ (Arcal 21)", "m³", Z_GAS, "Air Liquide", "Batería de 11 cilindros", 115.5,
      "MP-25", "Q03"),
    I("MP-92", "Disco de sensitiva 355 mm", "u", Z_SOL, "Würth Argentina", "Caja de 25 u", 25, "MP-27", "X07"),
    I("MP-27", "Disco de corte 115 × 1,0 mm", "u", Z_SOL, "Espacio Industria", "Caja de 25 u", 25, "MP-27", "X07"),
    I("MP-93", "Disco flap 115 mm grano 60", "u", Z_SOL, "Würth Argentina", "Caja de 25 u", 25, "MP-27", "X07"),
    I("MP-94", "Mecha de acero rápido 4,5 mm", "u", Z_SOL, "Espacio Industria", "Caja de 25 u", 25, "MP-27",
      "X07"),
    I("MP-95", "Spray antiproyecciones 400 ml", "u", Z_SOL, "Würth Argentina", "Caja de 25 u", 25, "MP-27", "X07"),
    I("MP-28", "Granalla de acero S330 / S390", "kg", Z_SOL, "CyM Materiales", "Pallet de 1.000 kg", 1000, "MP-28",
      "X07"),
    I("MP-31", "Pintura en polvo poliéster rojo 03-1-050", "kg", Z_POL, "Sinteplast", "Pallet de 600 kg", 600,
      "MP-31", "X07"),
    I("MP-33", "Polvo ABC 60 (manuales y sus recargas)", "kg", Z_POL, "DEMSA", "Big bag de 1.000 kg", 1000, "MP-33",
      "P06"),
    I("MP-33B", "Polvo ABC 90 (carros y sus recargas)", "kg", Z_POL, "DEMSA", "Big bag de 1.000 kg", 1000, "MP-33",
      "P06"),
    I("MP-34", "Nitrógeno seco", "m³", Z_GAS, "Air Liquide", "Pack de 12 cilindros", 110, "MP-34", "P04"),
    I("MP-35", "Polvo BC (recargas)", "kg", Z_POL, "Polvex", "Bolsa de 25 kg", 25, "MP-35", "Q07"),
    I("MP-36", "Polvo clase D (recargas)", "kg", Z_POL, "DEMSA", "Bolsa de 25 kg", 25, "MP-36", "X07"),
    I("MP-37", "CO₂ líquido (recargas)", "kg", Z_GAS, "Air Liquide", "Cilindro de 25 kg", 25, "MP-37", "X07"),
    I("MP-38", "Agente limpio HCFC-123 (recargas)", "kg", Z_REC, "Polvex / DEMSA", "Cilindro de 50 kg", 50, "MP-38",
      "X07"),
    I("MP-40", "Concentrado AFFF 3 % (recargas)", "L", Z_REC, "Polvex", "Bidón de 20 L", 20, "MP-40", "Q07"),
    I("MP-41", "Acetato de potasio Gastro K (recargas)", "L", Z_REC, "Polvex", "Bidón de 20 L", 20, "MP-41", "Q07"),
    I("MP-42", "Válvula HZ M-22 con tubo de pesca (1 kg)", "u", Z_INM, "Polvex (HZ)", "Pallet de 3.000 u", 3000,
      "MP-42", "X07"),
    I("MP-43", "Válvula HZ M-30 con resorte, sin tubo pescante (2,5 a 10 kg)", "u", Z_INM, "Polvex (HZ)",
      "Pallet de 1.500 u", 1500, "MP-43", "X07"),
    I("MP-44/25", "Válvula HZ de carro 25 kg, sin tubo pescante", "u", Z_INR, "Polvex (HZ)", "Caja de 10 u", 10,
      "MP-44", "X07"),
    I("MP-44/50", "Válvula HZ de carro 50-70-100 kg, sin tubo pescante", "u", Z_INR, "Polvex (HZ)", "Caja de 10 u",
      10, "MP-44", "X07"),
    I("MP-45", "Manómetro 1,4 MPa", "u", Z_INM, "Yukon", "Caja de 100 u", 100, "MP-45", "X07"),
] + [I(f"MP-46/{p}", f"Tubo pescante de PVC rígido, {n}", "u", Z_INM, PESCA_PVC, "Caja de 200 u", 200, "MP-46",
       "X07") for _, p, n in MAN[1:]] + [
    I(f"MP-47/{p}", f"Tubo pescante de acero SAE 1010, carro {n[4:]}", "u", Z_INR, PESCA_ACERO, "Caja de 20 u", 20,
      "MP-47", "X07") for _, p, n in ROD] + [
    I(f"MP-48/{p}", f"{'Tobera' if p == 'ABC1' else 'Manguera con tobera'} {n}", "u", Z_INM, "Yukon",
      "Caja de 50 u", 50, "MP-48", "X07") for _, p, n in MAN] + [
    I(f"MP-49/{p}", f"Manguera con válvula esférica y tobera, carro {n[4:]}", "u", Z_INR, "Yukon", "Caja de 10 u", 10,
      "MP-49", "X07") for _, p, n in ROD] + [
    I("MP-50/M22", "Junta tórica del cuello M22", "u", Z_INM, "Argensold", "Bolsa de 1.000 u", 1000, "MP-50", "X07"),
    I("MP-50/M30", "Junta tórica del cuello M30", "u", Z_INM, "Argensold", "Bolsa de 1.000 u", 1000, "MP-50", "X07"),
    I("MP-50/2½", "Junta tórica del cuello 2½\"", "u", Z_INR, "Argensold", "Bolsa de 1.000 u", 1000, "MP-50", "X07"),
    I("MP-50/OT", "Junta tórica del cuello (recargas, surtido por medida)", "u", Z_REC, "Würth Argentina",
      "Bolsa de 1.000 u", 1000, "MP-50", "X07"),
] + [I(f"MP-52/{p}", f"Soporte de pared {n}", "u", Z_INM, "Matafuegos Aconcagua (J. C. Paz)", "Caja de 100 u", 100,
       "MP-52", "X07")
     for _, p, n in MAN[1:]] + [
    I("MP-87", "Soporte vehicular (1 kg)", "u", Z_INM, "Matafuegos Aconcagua (J. C. Paz)", "Caja de 100 u", 100,
      "MP-52", "X07"),
    I("MP-86", "Tornillo y tarugo de fijación", "u", Z_INM, "Würth Argentina", "Bolsa de 1.000 u", 1000, "MP-51",
      "X07"),
    I("MP-55", "Conjunto de seguridad de CO₂ (recargas)", "u", Z_REC, "Georgia", "Bolsa de 100 u", 100, "MP-55",
      "X07"),
] + [I(f"MP-56/{p}", f"Etiqueta de características {n}", "u", Z_ID,
         "Rotolan (Lanús)" if p in ("ABC1", "ABC2.5", "ABC5") else "RotulArte (CABA)", "Rollo de 1.000 u", 1000,
         "MP-56", "X07") for _, p, n in FAB] + [
    I("MP-78", "Oblea de fabricación (PBA)", "u", Z_ID, "DPS (PBA)", "Paquete de 500 u", 500, "MP-58", "X07"),
    I("MP-79", "Estampilla IRAM de conformidad", "u", Z_ID, "IRAM", "Paquete de 500 u", 500, "MP-58", "X07"),
    I("MP-77", "Tarjeta de identificación AGC (CABA)", "u", Z_ID, "AGC (CABA)", "Paquete de 500 u", 500, "MP-58",
      "X07"),
    I("MP-80", "Etiqueta de n° de serie GS1", "u", Z_ID, "Rotolan (Lanús)", "Rollo de 1.000 u", 1000, "MP-56", "X07"),
    I("MP-81", "Faja de garantía válvula-cuello", "u", Z_ID, "Rotolan (Lanús)", "Rollo de 1.000 u", 1000, "MP-56",
      "X07"),
    I("MP-51", "Precinto de fábrica", "u", Z_ID, "SonFuertes", "Bolsa de 1.000 u", 1000, "MP-51", "X07"),
    I("MP-88", "Etiqueta del cilindro suelto", "u", Z_ID, "Rotolan (Lanús)", "Rollo de 1.000 u", 1000, "MP-56", "X07"),
    I("MP-62/M22", "Tapón protector de rosca M22", "u", Z_EMB, "ACPLIND (CABA)", "Caja de 5.000 u", 5000,
      "MP-62", "X07"),
    I("MP-62/M30", "Tapón protector de rosca M30", "u", Z_EMB, "ACPLIND (CABA)", "Caja de 5.000 u", 5000,
      "MP-62", "X07"),
    I("MP-62/2½", "Tapón protector de rosca 2½\"", "u", Z_EMB, "ACPLIND (CABA)", "Caja de 5.000 u", 5000,
      "MP-62", "X07"),
    I("MP-85", "Impresos (protocolo de ensayo, remito de préstamo)", "u", Z_ID, "Formas Visuales (Lanús)",
      "Paquete de 500 u", 500, "MP-58", "X07"),
    I("MP-57", "Marbete del año", "u", Z_REC, "Rotolan (Lanús)", "Paquete de 1.000 u", 1000, "MP-57", "X07"),
    I("MP-97", "Precinto del recargador", "u", Z_REC, "SonFuertes", "Bolsa de 1.000 u", 1000, "MP-51", "X07"),
    I("MP-72", "Etiqueta de servicio", "u", Z_REC, "Rotolan (Lanús)", "Rollo de 1.000 u", 1000, "MP-56", "X07"),
    I("MP-75", "Estampilla y oblea de recarga (PBA)", "u", Z_REC, "DPS (PBA)", "Paquete de 500 u", 500, "MP-58",
      "X07"),
    I("MP-76", "Tarjeta de identificación DPS (PBA)", "u", Z_REC, "DPS (PBA)", "Paquete de 500 u", 500, "MP-58",
      "X07"),
    I("MP-82", "Faja amarilla de sustituto", "u", Z_REC, "Rotolan (Lanús)", "Rollo de 1.000 u", 1000, "MP-56", "X07"),
    I("MP-83", "Tarjeta DPS «sustituto habilitado»", "u", Z_REC, "DPS (PBA)", "Paquete de 500 u", 500, "MP-58",
      "X07"),
    I("MP-84", "Tarjeta AGC «es sustituto»", "u", Z_REC, "AGC (CABA)", "Paquete de 500 u", 500, "MP-58", "X07"),
    I("MP-65", "Esmalte sintético rojo RAL 3000 (recargas)", "L", Z_REC, "Sinteplast", "Lata de 4 L", 4, "MP-65", "X07"),
] + [I(f"MP-59/{p}", f"Caja de cartón {'doble' if p == 'ABC10' else 'simple'} {n}", "u", Z_EMB, "Sellerpack (Pompeya)",
       "Pallet de 600 cajas", 600, "MP-59", "X07") for _, p, n in MAN] + [
    I("MP-89", "Etiqueta de caja", "u", Z_EMB, "Rotolan (Lanús)", "Rollo de 1.000 u", 1000, "MP-56", "X07"),
    I("MP-60", "Pallet de madera 1200 × 1000", "u", Z_EMB, "IndusPallets", "Pila de 15 u", 15, "MP-60", "X07"),
    I("MP-90", "Etiqueta de pallet", "u", Z_EMB, "Rotolan (Lanús)", "Rollo de 1.000 u", 1000, "MP-56", "X07"),
] + [I(f"MP-96/{p}", f"Servicio de pintura, carro {n[4:]} (tercerizado)", "u", Z_SERV, "Pintor de carros",
       "Unidad", 1, "MP-60", "Q09") for _, p, n in ROD] + [
    I(f"RV-{c}", f"Matafuego {n} (reventa)", "u", Z_PT,
      "Matafuegos Rodo" if c in ("CO2_5kg", "SALESK_6l", "HCFC-HFC_5kg") else
      "Maxiseguridad (Munro)" if c in REV_MAXI else "Fabricante nacional", "Unidad", 1,
      "MP-43", "X07", 60 if c == "AFFF_50l" else 180, "X13") for c, n in REV]
ITEM = {it["cod"]: it for it in ITEMS}
# plazo de entrega en días hábiles: (mínimo, máximo, de dónde sale, figura). Primero las cotizaciones que lo informan;
# después el plazo que publica en su web el proveedor o, si el elegido no lo publica, el de otro vendedor del mismo
# producto o rubro en el país, primero PBA (W4 a W35, buscados del 12 al 13/10/2026), marcado «Aproximado» al proveedor
# directo. Lo que nadie publica toma el de la planilla de MP y queda «sin plazo publicado». Corridos → hábiles: × 5/7.
PRAD = (7, 10, "Cotización Pradecon: 7 a 10 días", "Q01")
PRAD_H = (7, 10, "Aproximado: cotización Pradecon (plazo del fleje; para hojas no lo informa)", "Q01")
PRAD_P = (7, 10, "Aproximado: cotización Pradecon (Pacheco cotizó el mismo fleje sin plazo)", "Q01")
AL = (1, 1, "Cotización Air Liquide: día + 1", "Q03")
ELI = (25, 29, "Cotización Eli-Met: 35 a 40 días corridos = 25 a 29 hábiles", "Q03")
METAPOL = (1, 3, "Aproximado: web Metapol (CABA), barras SAE 1045: entrega en 24 a 72 h en el AMBA (Vladimirsky no "
                  "lo publica)", "W4")
CHINA = (54, 68, "Aproximado: web Importa Fácil: importar de China lleva 75 a 95 días corridos desde el pedido = 54 "
                 "a 68 hábiles", "W5")
ESPACIO = (1, 1, "Web Espacio Industria (Munro), discos de corte y mechas: entrega en CABA y GBA en 24 h", "W6")
ROTOLAN = (3, 7, "Web Rotolan (Lanús), etiquetas en rollo de 20 a 210 mm de ancho, desde 500 u: 3 a 7 días hábiles",
           "W7")
# la etiqueta del 10 kg (224 × 343 mm) y la de los carros (213 mm de alto) no entran en 210 mm: vinilo laminado
ROTULARTE = (1, 2, "Web RotulArte (CABA), vinilo autoadhesivo full color con laminado, hasta 155 cm de ancho: 24 a 48 h "
                   "hábiles desde que se aprueba el arte", "W16")
DAMIAN = (3, 5, "Web Papelera Damián, que vende los SonFuertes: localidades limítrofes de CABA, 3 a 5 días hábiles", "W8")
SELLERPACK = (10, 15, "Web Sellerpack (Pompeya, CABA), cajas a medida: 10 a 15 días desde que se aprueba la prueba "
                       "final (la primera vez, +5 a 10 si lleva troquel nuevo)", "W9")
DEPALLET = (1, 3, "Aproximado: web Depallet, pallets: entrega en 24 a 72 h (IndusPallets no lo publica)", "W10")
WURTH = (3, 5, "Web Würth Argentina: despacha al día hábil siguiente a la factura y entrega en hasta 5 días hábiles "
               "(unos 3 en CABA y AMBA)", "W11")
SOLDANET = (4, 5, "Web Soldanet, alambre MIG ER70S-6 de 15 kg: 4 a 5 días hábiles después de acreditado el pago", "W12")
RODO = (1, 2, "Web Matafuegos Rodo: entrega en 24/48 h en CABA y GBA (ABC, CO₂ 3,5 y 5 kg, K 6 l y HCFC 5 kg)", "W13")
ARGENSOLD = (1, 3, "Argensold (Lomas del Mirador): entrega en 72 h a todo el país (su Facebook); 24 h según un "
                   "directorio antiguo", "W14")
MAXISEG = (2, 2, "Web Maxiseguridad (Munro): entrega a domicilio a las 48 h del pedido por la web (nota antigua); su "
                 "catálogo tiene agua y AFFF 10 l, AFFF 50 l, BC 5 kg y CO₂ 2 kg", "W15")
OXICAS = (3, 6, "Aproximado: web Grupo Oxicas (Villa Constitución, Santa Fe), polvo ABC 60 de DEMSA en bolsa de "
                "25 kg: despacha a las 24 a 48 h del pago por OCA, que entrega en 48 a 96 h hábiles = 3 a 6 hábiles "
                "(DEMSA no lo publica)", "W24")
VERNASSA = (1, 1, "Aproximado: web Vernassa (CABA), concentrado AFFF de baja expansión al 1, 3 o 6 % en bidón de "
                  "20 L: plazo de entrega inmediato (Polvex no lo publica)", "W25")
# aproximados al proveedor directo, que no publica el suyo: el que publica otro vendedor del mismo producto o rubro
GALPON = (4, 7, "Aproximado: El Galpón (Corrientes), válvulas de 1 kg y de 5 y 10 kg, manómetro de 1,4 MPa y manguera "
                "de 5 kg: prepara el pedido en 48 h y despacha por Correo Argentino (2 a 5 días hábiles)", "W27")
GALPON_R = (4, 7, "Aproximado con los repuestos de matafuego de El Galpón (Corrientes): 48 h de preparación + Correo "
                  "Argentino en 2 a 5 días hábiles (nadie publica el plazo de esta pieza)", "W27")
OXICAS_R = (3, 6, "Aproximado con el polvo de DEMSA en Grupo Oxicas: despacho de 24 a 48 h + OCA en 48 a 96 h hábiles "
                  "(nadie publica el plazo de este agente)", "W24")
ACEROSYA = (5, 5, "Aproximado: AcerosYa (Sarandí), caños estructurales en barras de 6 m: entrega dentro de los 7 días "
                  "del pago = 5 hábiles (radio de 30 km)", "W28")
MORSHOP = (3, 5, "Aproximado: Morshop, ruedas industriales y de goma maciza para carros: despacha a las 24 a 72 h "
                 "hábiles del pago y el correo entrega en 48 h en Buenos Aires", "W29")
EGM = (1, 1, "Aproximado: EGM Trader (Florida, PBA), importador de granalla de acero esférica en bolsa de 25 kg y big "
             "bag: su publicación sale con envío en el día en CABA y GBA", "W30")
ENAMELS = (1, 1, "Aproximado: Enamels (Lanús), distribuidora de pintura en polvo: entrega inmediata por stock en el "
                 "GBA", "W31")
MAFER = (2, 10, "Aproximado: Pinturerías Mafer, que vende los esmaltes sintéticos de Sinteplast: 2 a 10 días "
                "hábiles en CABA y GBA", "W32")
DPS_T = (1, 1, "Trámite DPS (Res. 349/07): la Dirección Provincial las expide al representante de la firma inscripta "
               "al pagar el arancel (aproximado con el trámite de la AGC)", "W33")
IRAM_T = (21, 21, "IRAM, procedimiento DC-PG-129 de las marcas de conformidad (al que remite el Anexo R de "
                  "extintores): las estampillas se piden con 30 días de anticipación = 21 hábiles", "W34")
MARTIN = (3, 10, "Aproximado: Martin Morris (Coronel Suárez, PBA), tapones plásticos de rosca macho y hembra (BSP y "
                 "NPT, polipropileno, ½\" a 3\"): entrega en 3 a 10 días hábiles desde el despacho (ACPLIND no lo "
                 "publica)", "W35")
ACONCAGUA_D = (3, 7, "Aproximado: Matafuegos Aconcagua vende el extintor clase D de 5 kg (se encarga indicando los "
                     "metales): 3 a 7 días hábiles después de acreditado el pago", "W26")
ACONCAGUA = (3, 7, "Web Matafuegos Aconcagua (José C. Paz), soportes de pared de 2,5 a 10 kg y de vehículo de 1 kg: "
                   "3 a 7 días hábiles después de acreditado el pago", "W26")
FORMAS = (1, 3, "Guía Todo Lanús, Formas Visuales (Lanús Oeste), imprenta digital y offset: entrega en 24 a 72 h",
          "W17")
AGC_T = (1, 1, "Trámite AGC: FLAMA genera e imprime las tarjetas en el sistema de la AGC al pagar la boleta; no "
               "espera a un proveedor", "W18")
# link de cada fuente web del plazo: (página donde está el plazo, otra página: producto, correo o catálogo)
# nombre corto de cada fuente web, para la columna «De dónde sale el plazo» (el texto completo va en el comentario)
NOMBRE_FUENTE = {"W4": "Metapol (CABA)", "W5": "Importa Fácil (China)", "W6": "Espacio Industria (Munro)",
                 "W7": "Rotolan (Lanús)", "W8": "Papelera Damián", "W9": "Sellerpack (Pompeya)", "W10": "Depallet",
                 "W11": "Würth Argentina", "W12": "Soldanet", "W13": "Matafuegos Rodo",
                 "W14": "Argensold (Lomas del Mirador)", "W15": "Maxiseguridad (Munro)", "W16": "RotulArte (CABA)",
                 "W17": "Formas Visuales (Lanús)", "W18": "Trámite AGC", "W24": "Grupo Oxicas (Santa Fe)",
                 "W25": "Vernassa (CABA)", "W26": "Matafuegos Aconcagua (J. C. Paz)", "W27": "El Galpón (Corrientes)",
                 "W28": "AcerosYa (Sarandí)", "W29": "Morshop", "W30": "EGM Trader (Florida)", "W31": "Enamels (Lanús)",
                 "W32": "Pinturerías Mafer", "W33": "Trámite DPS (Res. 349/07)", "W34": "IRAM (DC-PG-129)",
                 "W35": "Martin Morris (Coronel Suárez)"}
LINKS = {
    "W4": ("https://metapol.com.ar/productos/sae-1045.html", ""),
    "W5": ("https://importa-facil.com.ar/cuanto-cuesta-importar-de-china-argentina", ""),
    "W6": ("https://espacio-industria.com/", ""),
    "W7": ("https://rotolan.com.ar/etiquetas-en-rollo.html", ""),
    "W8": ("https://papeleradamian.com/presta/content/delivery.html", ""),
    "W9": ("https://sellerpack.com.ar/industrias/construccion-industria/",
           "https://sellerpack.com.ar/cajas/exhibidor/"),
    "W10": ("https://depallet.com.ar/logistica/", ""),
    "W11": ("https://www.wurth.com.ar/contenidos/preguntas-frecuentes/tiempos-de-entrega.html",
            "https://www.wurth.com.ar/es/tienda/corte-y-taladro/discos-de-corte/disco-de-corte-premium-355x3-5x25"
            "-4.html"),
    "W12": ("https://soldanet.com/preguntasfrecuentes", ""),
    "W13": ("https://www.matafuegosrodo.com/", ""),
    "W14": ("https://www.facebook.com/people/Argensold-el-nombre-del-Oring/100091676384453/",
            "https://editores.com.ar/empresa/argensold"),
    "W15": ("https://www.maxiseguridad.com.ar/detalle-noticias-maxiseguridad/169/Venta-y-Recarga-de-Matafuegos-en-Forma-"
            "Online", "https://catalogo.maxiseguridad.com.ar/index.php"),
    "W16": ("https://rotularte.com.ar/impresion-vinilo-autoadhesivo/", ""),
    "W17": ("https://www.todolanus.com.ar/rubro/imprentas-talleres-graficos", ""),
    "W18": ("https://www.buenosaires.gob.ar/tramites/registro-de-fabricacion-reparacion-recarga-instalacion-y-con"
            "trol-de-extintores-matafuegos-y",
            ""),
    "W24": ("https://www.grupooxicas.com/content/9-envios",
            "https://www6.oca.com.ar/PlataformaEnvios/PreguntasFrecuentes"),
    "W25": ("https://vernassa.com.ar/pt/producto/concentrado-espumigeno-de-baja-expansion-afff/",
            "https://vernassa.com.ar/producto/concentrado-espumigeno-contra-incendio-afff-ar/"),
    "W26": ("https://www.matafuegosaconcagua.com/preguntas-frecuentes/",
            "https://www.matafuegosaconcagua.com/productos/soporte-pared-para-matafuego/"),
    "W27": ("https://seguridad.elgalponcorrientes.com.ar/informacion/",
            "https://www.correoargentino.com.ar/MiCorreo/public/infopaqarpymes"),
    "W28": ("https://www.acerosya.com/preguntas_frecuentes/",
            "https://www.acerosya.com/canos-estructurales/cano-redondo/"),
    "W29": ("https://www.morshop.com.ar/envios-y-retiros/",
            "https://www.morshop.com.ar/ruedas/ruedas-industriales-y-de-carga/"),
    "W30": ("https://egmtrader.mercadoshops.com.ar/MLA-907638018-granalla-de-acero-al-carbono-esferica-varias-med"
            "idas-x-25-kg-_JM",
            "https://www.egmtrader.com/empresa/"),
    "W31": ("https://www.enamels.com.ar/la-empresa/", ""),
    "W32": ("https://www.pintureriasmafer.com.ar/terminos-y-condiciones/",
            "https://www.pintureriasmafer.com.ar/hogar-y-obra/esmaltes-sinteticos/"),
    "W33": ("https://www.ambiente.gba.gob.ar/sites/default/files/RESOLUCION%20349%2007.pdf", ""),
    "W34": ("https://quesosespeciales.com/resources/img/quality/DC-PG_129_IRAM.pdf",
            "https://sellos.iram.org.ar/manuales/Anexo_R-Proteccion_contra_Incendios.pdf"),
    "W35": ("https://www.martinmorris.ar/shipping-policies",
            "https://martinmorris.ar/products/tapa-o-tapon-rosca-hembra-de-1-arag-5344"),
}
LT_COT = {"MP-02": PRAD_H, "MP-03": PRAD_H, "MP-04": PRAD_H, "MP-11": PRAD_H, "MP-12": PRAD_H,
          "MP-05": PRAD, "MP-13": PRAD, "MP-06": PRAD_P, "MP-07": PRAD_P, "MP-08": PRAD_P, "MP-09": PRAD_P,
          "MP-10": PRAD_P, "MP-25": AL, "MP-34": AL, "MP-37": AL, "MP-18": ELI, "MP-19": ELI, "MP-20": ELI,
          "MP-21": ELI, "MP-68": METAPOL, "MP-69": METAPOL, "MP-15": CHINA, "MP-16": CHINA, "MP-17": CHINA,
          "MP-27": ESPACIO, "MP-94": ESPACIO, "MP-51": DAMIAN, "MP-97": DAMIAN, "MP-60": DEPALLET,
          **{c: ROTOLAN for c in ("MP-80", "MP-81", "MP-88", "MP-72", "MP-82", "MP-89", "MP-90", "MP-57")},
          **{f"MP-56/{p}": ROTOLAN if p in ("ABC1", "ABC2.5", "ABC5") else ROTULARTE for _, p, _ in MAN + ROD},
          **{f"MP-59/{p}": SELLERPACK for _, p, _ in MAN},
          **{c: WURTH for c in ("MP-86", "MP-95", "MP-93", "MP-92", "MP-70", "MP-71", "MP-50/OT")},
          "MP-23": SOLDANET, "MP-24": SOLDANET, "MP-85": FORMAS, "MP-77": AGC_T, "MP-84": AGC_T,
          **{f"RV-{c}": RODO for c in ("CO2_5kg", "SALESK_6l", "HCFC-HFC_5kg")},
          **{f"MP-50/{m}": ARGENSOLD for m in ("M22", "M30", "2½")},
          **{c: ACONCAGUA for c in ("MP-52/ABC2.5", "MP-52/ABC5", "MP-52/ABC10", "MP-87")},
          "MP-33": OXICAS, "MP-40": VERNASSA,
          **{c: GALPON for c in ("MP-42", "MP-43", "MP-45")}, **{f"MP-48/{p}": GALPON for _, p, _ in MAN},
          **{c: GALPON_R for c in ("MP-44/25", "MP-44/50", "MP-55")},
          **{f"MP-49/{p}": GALPON_R for _, p, _ in ROD}, **{f"MP-47/{p}": GALPON_R for _, p, _ in ROD},
          **{f"MP-46/{p}": GALPON_R for _, p, _ in MAN[1:]},
          **{c: OXICAS_R for c in ("MP-33B", "MP-35", "MP-36", "MP-38", "MP-41")},
          **{c: ACEROSYA for c in ("MP-01", "MP-66", "MP-67")},
          **{f"MP-54/{d}": MORSHOP for d in ("300", "350", "400")},
          "MP-28": EGM, "MP-31": ENAMELS, "MP-65": MAFER, "MP-79": IRAM_T, "RV-CLASED_9l": ACONCAGUA_D,
          **{f"MP-62/{m}": MARTIN for m in ("M22", "M30", "2½")},
          **{c: DPS_T for c in ("MP-75", "MP-76", "MP-78", "MP-83")},
          **{f"RV-{c}": MAXISEG for c in REV_MAXI}}
# sin cotización, pero el proceso de carros los nombra (tabla 3 del puesto 1, tomados de la planilla de MP)
LT_CARROS = {"MP-15", "MP-16", "MP-17", "MP-33", "MP-33B", "MP-44/25", "MP-44/50", "MP-45", "MP-24", "MP-54/300",
             "MP-54/350", "MP-54/400"} | {f"MP-49/{p}" for _, p, _ in ROD}


def plazo(cod):
    if cod in LT_COT:
        return LT_COT[cod]
    lt = ITEM[cod]["lt"]
    if cod in LT_CARROS:
        return (lt, lt, "Sin plazo publicado: el del proceso de carros (de la planilla de MP)", "P06")
    return (lt, lt, "Sin plazo publicado: el de la planilla de MP", "X07")

# ------------------------------------------------------------------ datos de fabricación
CUERPO = {"ABC1": ("MP-01", 23, "P03"), "ABC2.5": ("MP-02", 27, "P02"), "ABC5": ("MP-03", 12, "P02"),
          "ABC10": ("MP-04", 15, "P02")}
# disco de cúpula / fondo: (ítem, ancho del fleje, Ø del disco, espesor); puente entre discos 3 mm (planilla de MP)
DISCO = {("ABC1", "Cúpula"): ("MP-05", 114, 95, 0.9), ("ABC1", "Fondo"): ("MP-13", 135, 112, 1.25),
         ("ABC2.5", "Cúpula"): ("MP-06", 200, 179, 1.25), ("ABC2.5", "Fondo"): ("MP-06", 200, 148, 1.25),
         ("ABC5", "Cúpula"): ("MP-07", 250, 221, 1.6), ("ABC5", "Fondo"): ("MP-08", 200, 177, 1.6),
         ("ABC10", "Cúpula"): ("MP-09", 300, 258, 1.6), ("ABC10", "Fondo"): ("MP-10", 249, 205, 2.0)}
PUENTE = 3
# planilla de MP, Geometria Acero filas 67 y 71: alambre (kg) y gas (m³) por recipiente de manuales
ALAMBRE_MAN = {"ABC1": 0.01499, "ABC2.5": 0.0304, "ABC5": 0.03675, "ABC10": 0.04575}
GAS_MAN = {"ABC1": 0.01671, "ABC2.5": 0.03388, "ABC5": 0.04096, "ABC10": 0.051}
# proceso de manuales: producción del pico con la que da 48 m³ de gas por día (tramos y desarrollos)
PICO_DOC_MAN = {"ABC1": 1216, "ABC2.5": 99, "ABC5": 446, "ABC10": 27}
GAS_DOC = 48
POR_PALLET_MAN = {"ABC1": 400, "ABC2.5": 167, "ABC5": 90, "ABC10": 40}      # proceso de manuales, embalaje
POR_PALLET_ROD = {"ABC25": 24, "ABC50": 9, "ABC70": 6, "ABC100": 6}         # planilla de MP, Pallets PT
# proceso de carros: (puesto, gas en L, alambre en g) por modelo 25 / 50 / 70 / 100
CARRO = [("Puesto 4", (3.4, 4.1, 9.5, 10.2), (6.0, 7.2, 16.8, 18.0)),
         ("Puesto 5", (18, 22, 32, 42), (48, 63, 120, 158)),
         ("Puesto 6", (10, 10, 17, 17), (10.5, 10.5, 29, 29)),
         ("Puesto 7", (66, 75, 122, 122), (190, 218, 463, 463)),
         ("Puesto 8", (96, 98, 109, 109), (141, 146, 199, 199))]
CASQ = {"ABC25": "MP-15", "ABC50": "MP-16", "ABC70": "MP-17", "ABC100": "MP-17"}
CUPLA = {"ABC25": "MP-20", "ABC50": "MP-20", "ABC70": "MP-21", "ABC100": "MP-21"}
MANIJA = {"ABC25": 1440, "ABC50": 1520, "ABC70": 1590, "ABC100": 1590}
CANO_EJE = {"ABC25": ("MP-66", 354), "ABC50": ("MP-67", 354), "ABC70": ("MP-67", 434), "ABC100": ("MP-67", 434)}
EJE = {"ABC25": ("MP-68", 490, 4000), "ABC50": ("MP-69", 510, 6000), "ABC70": ("MP-69", 590, 6000),
       "ABC100": ("MP-69", 590, 6000)}
RUEDA = {"ABC25": "MP-54/300", "ABC50": "MP-54/350", "ABC70": "MP-54/400", "ABC100": "MP-54/400"}
VALV = {"ABC1": "MP-42", "ABC2.5": "MP-43", "ABC5": "MP-43", "ABC10": "MP-43", "ABC25": "MP-44/25",
        "ABC50": "MP-44/50", "ABC70": "MP-44/50", "ABC100": "MP-44/50"}
JUNTA = {"ABC1": "MP-50/M22", "ABC70": "MP-50/2½", "ABC100": "MP-50/2½"}
TAPON = {"ABC1": "MP-62/M22", "ABC70": "MP-62/2½", "ABC100": "MP-62/2½"}
# recargas: Dimensionamiento V2, Parametros Recargas (B14 descargados, B15 vaciados, B22 pintura) y matriz (fila 74)
P_DESCARG, P_VACIA, P_PINTA = 0.10, 0.35, 0.15
CAMBIO_POLVO = 0.10       # proceso de recargas: polvo a descarte cuando el año termina en el dígito de fabricación
AFFF_CONC = 0.03          # proceso de recargas: premezcla al 3 %
ESMALTE_REND = 8.8        # Sinteplast Sintesmal, esmalte sintético industrial (ficha 741): m² por litro con 40 µm secos
PINT_REND, PINT_APROV = 9.46, 0.95   # WEG Politherm 66 (poliéster): m²/kg a 70 µm; aprovechamiento con recuperación


# ------------------------------------------------------------------ estilo
def _f(bold=False, color=NEGRO, size=10, italic=False):
    return Font(name=AR, size=size, bold=bold, color=color, italic=italic)


def _titulo(ws, texto):
    ws["A1"] = texto
    ws["A1"].font = _f(True, size=12)


def _cab(ws, r, textos, anchos=None, c0=1):
    for j, t in enumerate(textos, c0):
        c = ws.cell(r, j, t)
        c.font = _f(True, BLANCO)
        c.fill = PatternFill("solid", fgColor=NEGRO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 30
    if anchos:
        for j, w in enumerate(anchos, c0):
            ws.column_dimensions[L(j)].width = w


def _fila(ws, r, n, estilo=None, c0=1):
    for j in range(c0, c0 + n):
        c = ws.cell(r, j)
        c.border = Border(bottom=FINO)
        if estilo == "gris":
            c.font = _f(True)
            c.fill = PatternFill("solid", fgColor=GRIS_C)
        elif estilo == "sec":
            c.font = _f(True, BLANCO)
            c.fill = PatternFill("solid", fgColor=GRIS_O)
        elif not (c.font and c.font.bold):
            c.font = _f()


def _imprimir(ws, titulos):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = titulos
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.outlinePr.summaryRight = False


def q(hoja):
    return f"'{hoja}'"


def num(x):
    return f"{x:g}".replace(".", ",") if isinstance(x, (int, float)) else str(x)


# ------------------------------------------------------------------ lectura de fuentes
def leer_bom(ruta):
    """Filas del FLAMA_BOM.xlsx recalculado por producto, y la hoja MODELOS por código."""
    wb = load_workbook(ruta, data_only=True)
    prods, modelos = {}, {}
    for hoja in ("Fabricados", "Revendidos", "Cilindros", "Recargas"):
        ws = wb[hoja]
        cab = [c.value for c in ws[1]]
        prod = None
        for r, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if not fila[1]:
                continue
            d = dict(zip(cab, fila))
            d["fila"], d["hoja"] = r, hoja
            if d["Nivel"] == 0:
                prod = d["Código"]
                prods[(hoja, prod)] = []
            prods[(hoja, prod)].append(d)
    ws = wb["MODELOS"]
    cab = [c.value for c in ws[2]]
    for fila in ws.iter_rows(min_row=3, values_only=True):
        if fila[0]:
            modelos[fila[0]] = dict(zip(cab, fila))
    return prods, modelos


def ubicar(ws):
    """Filas de las Proyecciones por sección y etiqueta."""
    sec, pos = None, {}
    for r, fila in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=1, values_only=True), 1):
        v = fila[0]
        if not isinstance(v, str):
            continue
        t = v.strip()
        if t.startswith(("PRODUCCION", "RECARGAS", "MATAFUEGOS", "DISTRIBUCI", "STOCK DE", "ALTA ANUAL", "A. ",
                         "B. ", "C. ")) and t != "RECARGAS TOT":
            sec = t
            continue
        pos[(sec, t)] = r
    return pos


def codigo(clave):
    """Código del producto (ABC_5kg, CO2_2kg…) a partir de la clave de DEMANDA."""
    return clave[3:] if clave.startswith("RK_") else clave.split("_", 2)[2]


def sec(ix, prefijo):
    return next(s for (s, _) in ix if s and s.startswith(prefijo))


# ------------------------------------------------------------------ libro
class Libro:
    def __init__(self, proyecciones, bom):
        self.src = load_workbook(proyecciones, data_only=True)
        self.prods, self.modelos = leer_bom(bom)
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.dem = {}        # (línea, código del producto) → fila de DEMANDA
        self.refs = {}       # Ref → justificación (para el Word)
        self.lin = {}        # línea → datos de su hoja

    def ref(self, rid, **kw):
        self.refs[rid] = kw
        return rid

    # -------- copia de las proyecciones
    def copia(self, origen, nombre):
        src = self.src[origen]
        ws = self.wb.create_sheet(nombre)
        for row in src.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                d = ws.cell(c.row, c.column, c.value)
                d.number_format = c.number_format
                d.font = _f(bool(c.font and c.font.b))
        for k, dim in src.column_dimensions.items():
            if dim.width:
                ws.column_dimensions[k].width = dim.width
        return ubicar(src)

    def copia_usados(self):
        """Columnas Año, Mat Usados y RECARGAS TOT (A, O y P) de las hojas de proyección que separan los equipos
        usados, y Año y recargas de las que no los separan (supermercados, automotor, edificios): arriba la suma por
        año y rubro, abajo las filas copiadas. 2025 es el año 0 de las Proyecciones y no se copia. El año se pasa a
        número (EDUCACION lo trae como texto y AUTOMOTOR como «Año N»)."""
        ws = self.wb.create_sheet("PROY_USADOS")
        _titulo(ws, "PROYECCIONES: equipos usados y recargas por rubro")
        ws.cell(2, 1, "Suma por año de «Mat Usados» y «RECARGAS TOT» de los rubros que separan los usados; "
                      "supermercados, automotor y edificios no tienen esa columna (sólo nuevos y recargas): el % de "
                      "usados se mide en los cuatro primeros y se aplica a todas las recargas. Debajo, las filas "
                      "copiadas; 2025 es el año 0 de las Proyecciones (sin usados ni recargas, «—») y no se copia.")

        def año_de(v):
            v = str(v).strip()
            if v.startswith("Año"):
                v = str(2025 + int(v.split()[1]))
            try:
                return int(float(v))
            except ValueError:
                return None

        nombres = ("Salud", "Educación", "Industrias y almacenes", "Hoteles")
        _cab(ws, 4, ["Año"] + [f"{m}: {x}" for m in nombres for x in ("usados", "recargas")] +
             ["Total usados", "Total recargas", "% usados", "Supermercados: recargas (no separa usados)",
              "Automotor: recargas (no separa usados)", "Edificios: recargas (no separa usados)",
              "Recargas de los 7 rubros", "Parte de las recargas con usados medidos"], [9] + [12] * 25)
        ws.row_dimensions[4].height = 45
        r0 = 5 + len(AÑOS) + 2
        self.usados, sin = [], []
        bloques = [(h, 0, 14, 15) for h in HOJAS_USADOS] + [(h, cy, None, cr) for h, cy, cr in HOJAS_SIN_USADOS]
        c0 = 1
        for h, cy, cu, cr in bloques:
            tit = ("Año", "Mat Usados", "RECARGAS TOT") if cu is not None else ("Año", "RECARGAS")
            ws.cell(r0, c0, "PROYECCION " + h).font = _f(True)
            for j, t in enumerate(tit):
                ws.cell(r0 + 1, c0 + j, t).font = _f(True)
            n = r0 + 1
            for row in self.src["PROYECCION " + h].iter_rows(min_row=2, max_col=16, values_only=True):
                año = año_de(row[cy]) if row[cy] is not None else None
                if año not in AÑOS:
                    continue
                n += 1
                ws.cell(n, c0, año)
                vals = (row[cu], row[cr]) if cu is not None else (row[cr],)
                for j, v in enumerate(vals, 1):
                    ws.cell(n, c0 + j, v)
            cols = [L(c0 + j) for j in range(len(tit))]
            if cu is not None:
                self.usados.append((cols[0], cols[1], cols[2], r0 + 2, n))
            else:
                sin.append((cols[0], cols[1], r0 + 2, n))
            c0 += len(tit) + 1

        def sumif(cy, cv, a, b, r):
            return f"SUMIF(${cy}${a}:${cy}${b},$A{r},${cv}${a}:${cv}${b})"

        self.r_usados = {}
        for j, a in enumerate(AÑOS):
            r = 5 + j
            ws.cell(r, 1, a)
            for k, t in enumerate(self.usados):
                for m in (1, 2):
                    ws.cell(r, 2 * k + 1 + m, "=" + sumif(t[0], t[m], t[3], t[4], r)).number_format = "#,##0"
            ws.cell(r, 10, f"=B{r}+D{r}+F{r}+H{r}").number_format = "#,##0"
            ws.cell(r, 11, f"=C{r}+E{r}+G{r}+I{r}").number_format = "#,##0"
            ws.cell(r, 12, f"=J{r}/K{r}").number_format = "0.00%"
            for k, t in enumerate(sin):
                ws.cell(r, 13 + k, "=" + sumif(t[0], t[1], t[2], t[3], r)).number_format = "#,##0"
            ws.cell(r, 16, f"=K{r}+M{r}+N{r}+O{r}").number_format = "#,##0"
            ws.cell(r, 17, f"=K{r}/P{r}").number_format = "0.0%"
            _fila(ws, r, 17)
            self.r_usados[a] = r
        ws.freeze_panes = "B5"

    # -------- demanda
    def demanda(self):
        ws = self.wb.create_sheet("DEMANDA", 0)
        _titulo(ws, "DEMANDA (unidades por año)")
        _cab(ws, 3, ["Ref.", "Línea", "Producto"] + AÑOS, [7, 12, 34] + [10] * 10)
        ix, rs = self.ix, self.rs
        s_fab, s_cil = sec(ix, "PRODUCCION DE MATAFUEGOS"), sec(ix, "PRODUCCION DE CILINDROS")
        s_rec, s_rev = sec(ix, "RECARGAS"), sec(ix, "MATAFUEGOS TERCIARIZADOS")
        s_dr, s_du = sec(ix, "DISTRIBUCIÓN DE RECARGAS"), sec(ix, "DISTRIBUCIÓN DE UNIDADES")
        s_st, s_al = sec(rs, "STOCK DE SUSTITUTOS"), sec(rs, "ALTA ANUAL")
        venta = rs[next(k for k in rs if k[1].startswith("VENTA de reserva"))]
        pool_st = {c: rs[(s_st, k)] for k, c in POOL.items()}
        pool_al = {c: rs[(s_al, k)] for k, c in POOL.items()}
        al_tot = rs[(s_al, "Total alta anual de sustitutos")]
        PI, PX = q("PROY_INDEXADO"), q("PROY_RESERVAS")
        r = 4
        n_ref = [0]

        def dref():
            n_ref[0] += 1
            return f"D{n_ref[0]:02d}"

        def fila(lin, prod, vals, fmt="#,##0", clave=None, ref=None, **just):
            nonlocal r
            rid = ref or dref()
            self.ref(rid, hoja="DEMANDA", dato=prod, **just)
            for j, v in enumerate([rid, lin, prod], 1):
                ws.cell(r, j, v)
            for j, v in enumerate(vals):
                ws.cell(r, 4 + j, v).number_format = fmt
            _fila(ws, r, 13)
            if clave:
                self.dem[clave] = r
            r += 1
            return r - 1

        # datos que usan las otras hojas
        zr = {z: ix[(s, z)] for (s, z) in ix if s == s_dr}
        z4 = [zr["CABA Sur"]] + [zr[k] for k in zr if k.startswith("PBA Sur")]
        sub_r = next(zr[k] for k in zr if k.startswith("SUBTOTAL"))
        r_pi = fila("RECARGAS", "Parte de la planta industrial", [
            "=(" + "+".join(f"{PI}!{L(2 + j)}{x}" for x in z4) + f")/{PI}!{L(2 + j)}{sub_r}" for j in range(10)],
            "0.0%", como="Distribución de recargas por zona del Indexado: las 4 zonas que atiende la planta industrial "
                         "(CABA Sur y PBA Sur 1.º, 2.º y 3.er cordón, según el Dimensionamiento V2) sobre el subtotal "
                         "CABA + PBA, año por año. Las demás recargas son de la planta comercial (otro sitio).",
            figs=["X02", "X13"])
        self.r_pi = r_pi
        caba_u = [ix[(s_du, z)] for (s, z) in ix if s == s_du and z.startswith("CABA")]
        sub_u = next(ix[(s, z)] for (s, z) in ix if s == s_du and z.startswith("SUBTOTAL"))
        self.r_caba_u = fila("", "% de unidades nuevas que van a CABA", [
            "=(" + "+".join(f"{PI}!{L(2 + j)}{x}" for x in caba_u) + f")/{PI}!{L(2 + j)}{sub_u}" for j in range(10)],
            "0.0%", como="Distribución de unidades nuevas del Indexado: CABA Norte, Este, Sur y Oeste sobre el "
                         "subtotal CABA + PBA. El resto va a PBA.", figs=["X02"])
        self.r_caba_r = fila("RECARGAS", "% de recargas de la planta industrial en CABA", [
            f"={PI}!{L(2 + j)}{z4[0]}/(" + "+".join(f"{PI}!{L(2 + j)}{x}" for x in z4) + ")" for j in range(10)],
            "0.0%", como="Distribución de recargas del Indexado en las 4 zonas de la planta industrial: CABA Sur sobre "
                         "CABA Sur + PBA Sur 1.º, 2.º y 3.er cordón. El resto va a PBA.", figs=["X02", "X13"])
        self.r_u = fila("RECARGAS", "% de recargas de equipos usados", [
            f"=PROY_USADOS!$L${self.r_usados[AÑOS[j]]}" for j in range(10)], "0.00%",
            como="Proyecciones: suma de «Mat Usados» / suma de «RECARGAS TOT» del año, en las hojas de proyección que "
                 "separan los usados (salud, educación, industrias y almacenes, hoteles; sumadas en PROY_USADOS). "
                 "Supermercados, automotor y edificios no los separan: se les aplica el mismo %. Son "
                 "los equipos que llegan accionados o sin presión.", figs=["X04"])
        self.u35 = f"DEMANDA!$M${self.r_u}"
        _fila(ws, r_pi, 13)

        def col(hoja, fil, j):
            return f"{hoja}!{L(2 + j)}{fil}"

        # MANUALES y RODANTES
        for lin, grupo in (("MANUALES", MAN), ("RODANTES", ROD)):
            for c, p, n in grupo:
                partes, txt, figs = [lambda j, c=c: col(PI, ix[(s_fab, INV[c])], j)], \
                    [f"Indexado, «Producción de matafuegos», fila {ix[(s_fab, INV[c])]}"], ["X01"]
                if c == "ABC_5kg":
                    partes.append(lambda j: col(PX, venta, j))
                    txt.append(f"+ venta de reserva (Reservas, fila {venta})")
                    figs.append("X03")
                if c in pool_al:
                    partes.append(lambda j, c=c: f"{col(PX, pool_al[c], j)}*$D${r_pi}")
                    txt.append(f"+ sustitutos nuevos de la planta industrial (Reservas, alta anual, fila "
                               f"{pool_al[c]}, × D01): un sustituto ABC es un matafuego que fabrica FLAMA")
                    figs.append("X03")
                fila(lin, f"Matafuego {n}", ["=" + "+".join(f(j) for f in partes) for j in range(10)],
                     clave=(lin, f"FL_MAT_{c}"), como=" ".join(txt), figs=figs)
            for c, p, n in grupo:
                fila(lin, f"Cilindro {n} vacío", [f"={col(PI, ix[(s_cil, 'Cilindros ' + INV[c].replace(' s/ruedas', ''))], j)}"
                                                  for j in range(10)], clave=(lin, f"FL_REC_{c}"),
                     como=f"Indexado, «Producción de cilindros» (cilindros que se venden vacíos).", figs=["X01"])
        # RECARGAS
        for c in REC:
            base = [lambda j, c=c: col(PI, ix[(s_rec, INV[c])], j)]
            txt = f"(Indexado, «Recargas», fila {ix[(s_rec, INV[c])]}"
            if c in pool_st:
                base.append(lambda j, c=c: col(PX, pool_st[c], j))
                txt += f" + recarga anual del pool de sustitutos, Reservas fila {pool_st[c]}"
            fila("RECARGAS", f"Recarga {NOMBRE[c]}", [f"=({'+'.join(f(j) for f in base)})*$D${r_pi}"
                                                       for j in range(10)], clave=("RECARGAS", f"RK_{c}"),
                 como=txt + ") × D01, la parte de la planta industrial.", figs=["X01", "X03", "X02"])
        fila("RECARGAS", "Sustitutos nuevos (kit)", [f"={col(PX, al_tot, j)}*$D${r_pi}" for j in range(10)],
             clave=("RECARGAS", "KIT"), como=f"Reservas, total del alta anual de sustitutos (fila {al_tot}) × D01. "
             "Cada sustituto nuevo lleva su kit de identificación; el equipo se fabrica (ABC) o se compra (los demás) "
             "en su línea.", figs=["X03"])
        # REVENDIDOS
        for c, n in REV:
            partes = [lambda j, c=c: col(PI, ix[(s_rev, INV[c])], j)]
            txt = f"Indexado, «Matafuegos terciarizados», fila {ix[(s_rev, INV[c])]}"
            if c in pool_al:
                partes.append(lambda j, c=c: f"{col(PX, pool_al[c], j)}*$D${r_pi}")
                txt += f" + sustitutos nuevos de la planta industrial (Reservas, fila {pool_al[c]}, × D01)"
            fila("REVENDIDOS", f"Matafuego {n}", ["=" + "+".join(f(j) for f in partes) for j in range(10)],
                 clave=("REVENDIDOS", f"FL_MAT_{c}"), como=txt + ".", figs=["X01", "X03"])
        ws.freeze_panes = "D4"
        _imprimir(ws, "3:3")

    # -------- estacionalidad (como el Dimensionamiento Técnico FLAMA)
    def estacionalidad(self):
        ws = self.wb.create_sheet("ESTACIONALIDAD", 1)
        _titulo(ws, "ESTACIONALIDAD: producción mensual del año 2035")
        _cab(ws, 3, ["Ref.", "Perfil"] + MESES + ["Total"], [7, 34] + [8] * 12 + [9])
        self.r_perf = []
        for k, (nom, v) in enumerate(PERFILES):
            r = 4 + k
            rid = self.ref(f"E{k + 1:02d}", hoja="ESTACIONALIDAD", dato=nom,
                           como="Dimensionamiento Técnico FLAMA, perfil de estacionalidad (% de la venta anual por mes)."
                           if k < 3 else "Dimensionamiento Técnico FLAMA: las recargas de las dos plantas se reparten "
                                         "parejo en los 12 meses (anual / 12).", figs=["F01"] if k < 3 else ["F03"])
            ws.cell(r, 1, rid)
            ws.cell(r, 2, nom)
            for j, x in enumerate(v):
                ws.cell(r, 3 + j, x).number_format = "0.0%"
            ws.cell(r, 15, f"=SUM(C{r}:N{r})").number_format = "0%"
            _fila(ws, r, 15)
            self.r_perf.append(r)
        r = 4 + len(PERFILES) + 1
        _cab(ws, r, ["Línea", "Producto"] + MESES + ["Total", "Mes pico del producto"])
        ws.column_dimensions["P"].width = 12
        r += 1
        self.r_est, tot = {}, {}
        resumen = r + 0
        filas_lin = {}
        for lin in LINEAS:
            filas_lin[lin] = []
            for (l2, clave), rd in sorted(self.dem.items(), key=lambda x: x[1]):
                if l2 != lin or clave == "KIT":
                    continue
                pr = self.r_perf[perfil(codigo(clave), lin)]
                ws.cell(r, 1, lin)
                ws.cell(r, 2, f"=DEMANDA!C{rd}")
                for j in range(12):
                    ws.cell(r, 3 + j, f"=DEMANDA!$M${rd}*{L(3 + j)}${pr}").number_format = "#,##0"
                ws.cell(r, 15, f"=SUM(C{r}:N{r})").number_format = "#,##0"
                self.r_est[(lin, clave)] = r
                filas_lin[lin].append(r)
                _fila(ws, r, 16)
                r += 1
            ws.cell(r, 1, lin)
            ws.cell(r, 2, f"Total {lin.lower()}")
            a, b = filas_lin[lin][0], filas_lin[lin][-1]
            for j in range(13):
                ws.cell(r, 3 + j, f"=SUM({L(3 + j)}{a}:{L(3 + j)}{b})").number_format = "#,##0"
            _fila(ws, r, 16, "gris")
            tot[lin] = r
            r += 1
        # resumen: mes pico y factor de pico por línea
        r += 1
        _cab(ws, r, ["Ref.", "Línea", "Mes pico", "N.º de mes", "Unidades del mes pico", "Promedio mensual",
                     "Factor de pico"])
        r += 1
        self.r_pico = {}
        for lin in LINEAS:
            t = tot[lin]
            rid = self.ref(f"E{len(PERFILES) + 1 + LINEAS.index(lin):02d}", hoja="ESTACIONALIDAD",
                           dato=f"Mes pico de {lin.lower()}",
                           como="El mes de mayor producción del año 2035 sumando todos los productos de la línea (como "
                                "el total mensual del Dimensionamiento Técnico FLAMA). Factor de pico = unidades del "
                                "mes pico / promedio mensual. Cada producto se compra con su propio mes pico (columna "
                                "P): el AFFF 50 l sigue el perfil industrial y en diciembre, pico de revendidos, "
                                "vende menos que su promedio.", figs=["F02"])
            ws.cell(r, 1, rid)
            ws.cell(r, 2, lin)
            ws.cell(r, 4, f"=MATCH(MAX(C{t}:N{t}),C{t}:N{t},0)")
            ws.cell(r, 3, f'=INDEX($C${resumen - 1}:$N${resumen - 1},D{r})')
            ws.cell(r, 5, f"=INDEX(C{t}:N{t},D{r})").number_format = "#,##0"
            ws.cell(r, 6, f"=O{t}/12").number_format = "#,##0"
            ws.cell(r, 7, f"=E{r}/F{r}").number_format = "0.000"
            _fila(ws, r, 7)
            self.r_pico[lin] = r
            r += 1
        for lin in LINEAS:
            for rr in filas_lin[lin]:
                ws.cell(rr, 16, f"=MAX(C{rr}:N{rr})").number_format = "#,##0"
        ws.freeze_panes = "C4"
        _imprimir(ws, "3:3")

    # -------- consumos por unidad
    def bom_ident(self, hoja, prod, pref):
        """Compras 1 a 1 del BOM de un producto: válvula, manguera, manómetro, junta, soporte, identificación, carro
        y cilindro. Devuelve {ítem: (cantidad, destino, fila del BOM)}."""
        filas = self.prods[(hoja, prod)]
        out = {}

        def add(it, cant, dest, f):
            q0 = out.get(it, (0, dest, []))
            out[it] = (q0[0] + cant, dest, q0[2] + [f["fila"]])

        mapa = {"Manómetro": ("MP-45", ""), "Etiqueta / placa de características": (f"MP-56/{pref}", ""),
                "Oblea de fabricación": ("MP-78", "PBA"), "Estampilla IRAM de conformidad": ("MP-79", ""),
                "Tarjeta de identificación AGC": ("MP-77", "CABA"), "Etiqueta de n° de serie GS1": ("MP-80", ""),
                "Faja de garantía válvula-cuello": ("MP-81", ""), "Precinto de fábrica": ("MP-51", ""),
                "Soporte de pared": (f"MP-52/{pref}", ""), "Soporte vehicular con cierre rápido": ("MP-87", ""),
                "Tornillo y tarugo de fijación": ("MP-86", ""),
                "Tapón protector de rosca del cuello": (TAPON.get(pref, "MP-62/M30"), ""),
                "Etiqueta de identificación del cilindro suelto": ("MP-88", ""),
                "Junta tórica de asiento": (JUNTA.get(pref, "MP-50/M30"), "")}
        for f in filas:
            cod, pz = f["Código"], f["Pieza"]
            if cod.startswith("FL_SUS_"):
                break
            if cod.endswith("-K2"):
                add(VALV[pref], 1, "", f)
            elif cod.endswith("-S3"):
                add(f"MP-48/{pref}" if pref in [p for _, p, _ in MAN] else f"MP-49/{pref}", 1, "", f)
            elif pz in mapa and f["Nivel"] == 2:
                it, dest = mapa[pz]
                add(it, f["Cant."] or 1, dest, f)
            elif pz == "Caño de pesca" and pref != "ABC1":
                add(f"MP-46/{pref}" if pref in [p for _, p, _ in MAN] else f"MP-47/{pref}", 1, "", f)
            elif pz == "Neumático macizo":
                add(RUEDA[pref], f["Cant."], "", f)
            elif pz.startswith("Arandelas de tope"):
                add("MP-70", 4, "", f)
            elif pz.startswith("Pasadores del eje"):
                add("MP-71", 2, "", f)
        return out

    def area(self, hoja, prod):
        """Área exterior del recipiente (m²): la mitad de la superficie de las dos caras del BOM."""
        s = 0.0
        for f in self.prods[(hoja, prod)]:
            if f["Nivel"] == 2 and f["Pieza"] in ("Cuerpo", "Cúpula", "Fondo", "Cilindro sin costura"):
                s += (f["Superficie (dm²)"] or 0) / 2 / 100
        return s

    def filas_fab(self, lin, grupo):
        """Filas de consumo de MANUALES o RODANTES: {(ítem, destino): {columna: valor}} con su justificación."""
        cols = [f"FL_MAT_{c}" for c, _, _ in grupo] + [f"FL_REC_{c}" for c, _, _ in grupo]
        n = len(grupo)
        rows = []

        def fila(it, cons, como, figs, dest=""):
            rows.append(dict(it=it, dest=dest, cons=cons, como=como, figs=figs))

        mod = self.modelos
        if lin == "MANUALES":
            for k, (c, p, nm) in enumerate(grupo):
                it, rend, fg = CUERPO[p]
                fila(it, {k: f"=1/{rend}", n + k: f"=1/{rend}"},
                     f"Un cuerpo por recipiente; de {'un caño de 6 m salen' if p == 'ABC1' else 'una hoja salen'} "
                     f"{rend} cuerpos.", [fg])
            for (p, pz), (it, an, d, e) in DISCO.items():
                k = [x[1] for x in grupo].index(p)
                ex = next((x for x in rows if x["it"] == it), None)
                v = f"{an}*({d}+{PUENTE})*{e}*7.85/1000000"
                if ex:
                    ex["cons"] = {kk: vv + "+" + v for kk, vv in ex["cons"].items()}
                    ex["como"] += f" Más el fondo: Ø{d} mm."
                    continue
                fila(it, {k: "=" + v, n + k: "=" + v},
                     f"{pz} del {p[3:]} kg: kg de fleje por disco = ancho {an} mm × (Ø del disco {d} mm + puente "
                     f"{PUENTE} mm) × espesor {num(e)} mm × 7,85 kg/dm³ (método de la planilla de MP).",
                     ["X14", "X08", "Q02"])
            fila("MP-18", {0: 1, n: 1}, "Un asiento de válvula por cúpula (se compra mecanizado).", ["P01"])
            fila("MP-19", {k: 1 for k in (1, 2, 3)} | {n + k: 1 for k in (1, 2, 3)},
                 "Un asiento de válvula por cúpula (se compra mecanizado).", ["P01"])
            fila("MP-23", {k: ALAMBRE_MAN[p] for k, (_, p, _) in enumerate(grupo)} |
                 {n + k: ALAMBRE_MAN[p] for k, (_, p, _) in enumerate(grupo)},
                 "kg de alambre por recipiente de la planilla de MP (largo de costura × alambre por metro). El proceso "
                 "de manuales no da otro valor.", ["X08"])
            suma = "+".join(f"{PICO_DOC_MAN[p]}*{GAS_MAN[p]}" for _, p, _ in grupo)
            fila("MP-25", {k: f"={GAS_MAN[p]}*{GAS_DOC}/({suma})" for k, (_, p, _) in enumerate(grupo)} |
                 {n + k: f"={GAS_MAN[p]}*{GAS_DOC}/({suma})" for k, (_, p, _) in enumerate(grupo)},
                 f"m³ por recipiente de la planilla de MP, llevados a los {GAS_DOC} m³ por día del proceso de manuales: "
                 f"con la producción del pico de ese proceso (1.216, 99, 446 y 27 recipientes) la planilla da 43,3 m³, "
                 f"así que cada valor se multiplica por 48 / 43,3.", ["P01", "X08"])
            fila("MP-28", {k: 0.08 for k in range(2 * n)}, "0,08 kg de granalla por cilindro (proceso de manuales: "
                 "«unos 133 kg por día en el pico»).", ["P04"])
            fila("MP-31", {k: f"={self.area('Fabricados', 'FL_MAT_' + c):.4f}/{PINT_REND}/{PINT_APROV}"
                           for k, (c, _, _) in enumerate(grupo)} |
                 {n + k: f"={self.area('Fabricados', 'FL_MAT_' + c):.4f}/{PINT_REND}/{PINT_APROV}"
                  for k, (c, _, _) in enumerate(grupo)},
                 "kg = área exterior (BOM: mitad de la superficie de cuerpo, cúpula y fondo) / 9,46 m² por kg / 0,95. "
                 "Rendimiento a 70 µm del poliéster WEG Politherm 66 (densidad 1,51 g/cm³); 70 µm es el espesor de la "
                 "planilla de MP (Sinteplast, 60 a 80 µm). La cabina recupera el polvo (proceso de manuales): 95 % de "
                 "aprovechamiento (Powder Coating Institute).", ["B02", "X08", "P04", "W2", "W3"])
            fila("MP-33", {k: mod[f"FL_MAT_{c}"]["Carga objetivo"] for k, (c, _, _) in enumerate(grupo)},
                 "Carga objetivo del BOM (nominal + mitad de la tolerancia). El proceso de manuales no agrega merma "
                 "(1.254 kg por día para 413 matafuegos en el pico).", ["B01", "P01"])
            fila("MP-34", {k: f"={mod[f'FL_MAT_{c}']['V recipiente (L)']}/3*14/1000" for k, (c, _, _) in
                           enumerate(grupo)},
                 "Cuenta del proceso de manuales: cámara de gas de un tercio del volumen interior a 14 bar (presión "
                 "de servicio del plano). Volumen interior del BOM.", ["P04", "B01"])
            figs_id = ["B02"]
        else:
            fila("MP-11", {0: "=1/9", 1: "=1/6", n: "=1/9", n + 1: "=1/6"}, "Un cuerpo por recipiente; de una hoja "
                 "de 3,2 mm salen 9 cuerpos de 25 kg o 6 de 50 kg (plan de corte del proceso de carros).", ["P08"])
            fila("MP-12", {2: "=2/9", 3: "=1/3", n + 2: "=2/9", n + 3: "=1/3"},
                 "Hoja mixta del proceso de carros: 3 cuerpos de 70 kg + 1 de 100 kg; los 100 kg que sobran, a 3 por "
                 "hoja. Hojas = n70 / 3 + (n100 − n70 / 3) / 3 = 2/9 × n70 + 1/3 × n100 (vale mientras n70 ≤ 3 × "
                 "n100, que se cumple todos los años).", ["P08"])
            for it in ("MP-15", "MP-16", "MP-17"):
                ks = [k for k, (_, p, _) in enumerate(grupo) if CASQ[p] == it]
                fila(it, {k: 2 for k in ks} | {n + k: 2 for k in ks}, "Dos casquetes por recipiente (cúpula y "
                     "fondo, comprados estampados).", ["P06"])
            for it in ("MP-20", "MP-21"):
                ks = [k for k, (_, p, _) in enumerate(grupo) if CUPLA[p] == it]
                fila(it, {k: 1 for k in ks} | {n + k: 1 for k in ks}, "Una cupla por recipiente.", ["P06"])
            fila("MP-66", {k: f"={MANIJA[p]}/6000/0.95" + (f"+{CANO_EJE[p][1]}/6000/0.95" if CANO_EJE[p][0] ==
                                                         "MP-66" else "") for k, (_, p, _) in enumerate(grupo)},
                 "Manija (1.440 a 1.590 mm) y, en el 25 kg, el caño del eje (354 mm), sobre barras de 6 m con 95 % "
                 "de aprovechamiento (proceso de carros, 8.1). Sólo el carro terminado.", ["P16", "P15"])
            fila("MP-67", {k: f"={CANO_EJE[p][1]}/6000/0.95" for k, (_, p, _) in enumerate(grupo)
                           if CANO_EJE[p][0] == "MP-67"},
                 "Caño del eje (354 o 434 mm) sobre barras de 6 m con 95 % de aprovechamiento (proceso de carros, "
                 "8.1).", ["P16", "P15"])
            for it in ("MP-68", "MP-69"):
                fila(it, {k: f"=1/INT({EJE[p][2]}/{EJE[p][1]})" for k, (_, p, _) in enumerate(grupo)
                          if EJE[p][0] == it},
                     "Ejes enteros por barra: INT(largo de la barra / largo del eje) (proceso de carros, 8.1).",
                     ["P16"])
            figs_id = ["B02"]
            fila("MP-24", {k: f"=({'+'.join(str(a[k]) for _, _, a in CARRO)})/1000" for k in range(n)} |
                 {n + k: f"=({'+'.join(str(a[k]) for _, _, a in CARRO[:4])})/1000" for k in range(n)},
                 "Suma de los gramos de alambre de los puestos 4 a 8 del proceso de carros; el cilindro vacío no "
                 "pasa por el puesto 8 (accesorios).", ["P10", "P11", "P12", "P13", "P14"])
            fila("MP-25", {k: f"=({'+'.join(str(g[k]) for _, g, _ in CARRO)})/1000" for k in range(n)} |
                 {n + k: f"=({'+'.join(str(g[k]) for _, g, _ in CARRO[:4])})/1000" for k in range(n)},
                 "Suma de los litros de gas de los puestos 4 a 8 del proceso de carros; el cilindro vacío no pasa "
                 "por el puesto 8.", ["P10", "P11", "P12", "P13", "P14"])
            fila("MP-92", {k: "=3/200" for k in range(n)}, "3 cortes por juego de accesorios; 200 cortes por disco "
                 "(proceso de carros, 8.1).", ["P17"])
            fila("MP-27", {k: "=1/40" + ("+7/80" if k >= 2 else "") for k in range(n)} |
                 {n + k: "=7/80" for k in (2, 3)},
                 "1 disco cada 40 juegos de triángulos (8.1) y, en 70 y 100 kg, 7 discos por mes para las placas de "
                 "refuerzo sobre 80 cuerpos por mes (puesto 4).", ["P17", "P18"])
            fila("MP-93", {k: "=2/400" + ("+7/4/80" if k >= 2 else "") for k in range(n)} |
                 {n + k: "=7/4/80" for k in (2, 3)},
                 "2 discos flap por mes para 400 juegos (8.1) y, en 70 y 100 kg, 1 disco cada 4 tandas de placas "
                 "(puesto 4).", ["P17", "P18"])
            fila("MP-94", {k: "=2/200" for k in range(n)}, "2 agujeros por juego; 200 agujeros por mecha (8.1).",
                 ["P17"])
            fila("MP-95", {k: "=1/400" for k in range(2 * n)}, "Un disparo de 1 ml por cupla; aerosol de 400 ml "
                 "(puesto 6).", ["P18"])
            for k, (c, p, nm) in enumerate(grupo):
                fila(f"MP-96/{p}", {k: 1, n + k: 1}, "Un servicio de pintura por recipiente: la pintura de carros es "
                     "tercerizada (proceso de carros, 8.1; cotización por carro).", ["P15", "Q09"])
            fila("MP-33B", {k: f"={mod[f'FL_MAT_{c}']['Carga objetivo']}*1.01" for k, (c, _, _) in
                            enumerate(grupo)},
                 "Carga del BOM + 1 % de merma (proceso de carros, puesto 1, tabla 3).", ["B01", "P06"])
            fila("MP-34", {k: f"={mod[f'FL_MAT_{c}']['V recipiente (L)']}/3*14/1000" for k, (c, _, _) in
                           enumerate(grupo)},
                 "El proceso de carros no da el consumo: se usa la cuenta del proceso de manuales (un tercio del "
                 "volumen interior a 14 bar, misma pinza Yukon) con el volumen interior del BOM.", ["P04", "B01"])
        # compras 1 a 1 del BOM (válvula, manguera, manómetro, junta, soportes, identificación, carro, cilindro)
        idx = {}
        for k, (c, p, nm) in enumerate(grupo):
            for hoja, prod, kk in (("Fabricados", f"FL_MAT_{c}", k), ("Cilindros", f"FL_REC_{c}", n + k)):
                for it, (cant, dest, fb) in self.bom_ident(hoja, prod, p).items():
                    idx.setdefault((it, dest), {})[kk] = cant
        orden = [it["cod"] for it in ITEMS]
        for (it, dest), cons in sorted(idx.items(), key=lambda x: orden.index(x[0][0])):
            fila(it, cons, ("Cantidad por unidad del BOM" + (f"; sólo lo que va a {dest} (% de DEMANDA)" if dest
                                                             else "") + "."), figs_id, dest)
        # cilindro vacío y embalaje
        if lin == "MANUALES":
            fila("MP-85", {n + k: 1 for k in range(n)}, "Un protocolo de ensayo acompaña a cada cilindro vacío "
                 "(proceso de manuales, embalaje).", ["P05"])
            for k, (c, p, nm) in enumerate(grupo):
                fila(f"MP-59/{p}", {k: 1, n + k: 1}, "Una caja por recipiente, simple hasta el 5 kg y doble en el "
                     "10 kg (proceso de manuales, embalaje).", ["P05"])
            fila("MP-89", {k: 1 for k in range(2 * n)}, "Una etiqueta por caja (proceso de manuales, embalaje).",
                 ["P05"])
            fila("MP-60", {k: f"=1/{POR_PALLET_MAN[p]}" for k, (_, p, _) in enumerate(grupo)} |
                 {n + k: f"=1/{POR_PALLET_MAN[p]}" for k, (_, p, _) in enumerate(grupo)},
                 "Un pallet cada 400, 167, 90 o 40 recipientes (proceso de manuales, embalaje).", ["P05"])
            fila("MP-90", {k: f"=1/{POR_PALLET_MAN[p]}" for k, (_, p, _) in enumerate(grupo)} |
                 {n + k: f"=1/{POR_PALLET_MAN[p]}" for k, (_, p, _) in enumerate(grupo)},
                 "Una etiqueta por pallet (proceso de manuales, embalaje).", ["P05"])
        else:
            fila("MP-60", {n + k: f"=1/{POR_PALLET_ROD[p]}" for k, (_, p, _) in enumerate(grupo)},
                 "Cilindros vacíos de carro por pallet: 24, 9, 6 y 6 (planilla de MP, Pallets PT).", ["X07"])
            fila("MP-90", {n + k: f"=1/{POR_PALLET_ROD[p]}" for k, (_, p, _) in enumerate(grupo)},
                 "Una etiqueta por pallet.", ["X07"])
        return cols, rows

    def filas_rec(self):
        cols = [f"RK_{c}" for c in REC] + ["KIT"]
        rows = []
        mod = self.modelos

        def fila(it, cons, como, figs, dest=""):
            rows.append(dict(it=it, dest=dest, cons=cons, como=como, figs=figs))

        ks = {c: k for k, c in enumerate(REC)}
        abc60 = ["ABC_1kg", "ABC_2.5kg", "ABC_5kg", "ABC_10kg"]
        abc90 = ["ABC_25kg", "ABC_50kg", "ABC_70kg", "ABC_100kg"]
        u = self.u35
        for it, grupo in (("MP-33", abc60), ("MP-33B", abc90)):
            fila(it, {ks[c]: f"={mod[f'FL_MAT_{c}']['Carga objetivo']}*({CAMBIO_POLVO}+{u})+(1-{CAMBIO_POLVO}-{u})*"
                             f"{mod[f'FL_MAT_{c}']['Reposición en mantenimiento (UM)']}" for c in grupo},
                 "El polvo se recupera. Va a descarte (carga nueva entera) el que cambia por el año de fabricación "
                 "(1 de cada 10 años: 10 %) y el de los equipos que llegan sin presión o accionados (% de usados, "
                 "DEMANDA); al resto se le repone lo que se pierde en el desarme (reposición del BOM).",
                 ["R02", "R03", "B01"])
        fila("MP-35", {ks["BC_5kg"]: mod["FL_MAT_BC_5kg"]["Carga objetivo"]},
             "El BC se recarga siempre con polvo nuevo (unos 13 por mes: no se justifica recuperarlo). Carga del BOM.",
             ["R07", "B01"])
        fila("MP-36", {ks["CLASED_9l"]: mod["FL_MAT_CLASED_9l"]["Carga objetivo"]},
             "Mismo criterio que el BC del proceso de recargas: unos 12 por mes en la planta industrial, polvo nuevo "
             "siempre. Carga del BOM.", ["R07", "B01"])
        fila("MP-37", {ks[c]: f"={mod[f'FL_MAT_{c}']['Carga objetivo']}*{P_DESCARG}" for c in ("CO2_2kg", "CO2_5kg")},
             "El CO₂ se recupera y se vuelve a cargar; nuevo sólo para el 10 % que llega descargado o con pérdida de "
             "masa (Dimensionamiento V2, Parametros Recargas).", ["R09", "X11"])
        fila("MP-38", {ks["HCFC-HFC_5kg"]: f"={mod['FL_MAT_HCFC-HFC_5kg']['Carga objetivo']}*{P_DESCARG}"},
             "El agente limpio se recupera en circuito cerrado; nuevo sólo para el 10 % que llega descargado o con "
             "pérdida (Dimensionamiento V2).", ["R10", "X11"])
        fila("MP-40", {ks[c]: f"={mod[f'FL_MAT_{c}']['Carga objetivo']}*{AFFF_CONC}" for c in ("AFFF_10l", "AFFF_50l")},
             "El agente se cambia siempre; premezcla al 3 %: litros de concentrado = volumen × 0,03.", ["R08"])
        fila("MP-41", {ks["SALESK_6l"]: mod["FL_MAT_SALESK_6l"]["Carga objetivo"]},
             "Recarga anual completa del agente líquido (Dimensionamiento V2, matriz: llenado de agua, AFFF y K al "
             "100 %). Carga del BOM.", ["X12", "B01"])
        n2 = {}
        for c in REC:
            m = mod[f"FL_MAT_{c}"]
            if c.startswith("ABC"):
                n2[ks[c]] = f"={m['V recipiente (L)']}/3*14/1000"
            elif c.startswith("CO2"):
                continue
            elif c == "HCFC-HFC_5kg":
                n2[ks[c]] = f"={round(m['Gas (L normales)'], 3)}/1000*{P_VACIA}"
            else:
                n2[ks[c]] = f"={round(m['Gas (L normales)'], 3)}/1000"
        fila("MP-34", n2, "Se presuriza con nitrógeno con la misma pinza de producción. ABC: cuenta del proceso de "
             "manuales (un tercio del volumen a 14 bar); los demás: gas del BOM en litros normales. Agente limpio: "
             "sólo el 35 % que se vacía (Dimensionamiento V2). El CO₂ no lleva gas impulsor.",
             ["R05", "P04", "B01", "X12"])
        fila("MP-50/OT", {ks[c]: (P_VACIA if c.startswith("CO2") or c == "HCFC-HFC_5kg" else 1) for c in REC},
             "El o-ring se cambia siempre que se desarma: todos los de polvo y líquidos; en CO₂ y agente limpio, el "
             "35 % que se vacía (Dimensionamiento V2, matriz).", ["R04", "X12"])
        fila("MP-55", {ks[c]: P_VACIA for c in ("CO2_2kg", "CO2_5kg")},
             "Se reemplaza en cada vaciado total: 35 % de los CO₂ (Dimensionamiento V2).", ["X11", "R09"])
        for it, txt, fg in (("MP-57", "Marbete del año: siempre.", ["R04"]),
                            ("MP-97", "Precinto con la identificación de FLAMA: siempre.", ["R06"]),
                            ("MP-72", "Etiqueta de servicio (lleva también el vencimiento de la prueba hidráulica): "
                                      "siempre.", ["R06"])):
            fila(it, {ks[c]: 1 for c in REC}, txt, fg)
        for it, dest in (("MP-75", "PBA"), ("MP-76", "PBA"), ("MP-77", "CABA")):
            fila(it, {ks[c]: 1 for c in REC}, f"Una por recarga que va a {dest} (BOM de recargas; % de DEMANDA).",
                 ["B03"], dest)
        es = {}
        for c in REC:
            hoja = "Fabricados" if c.startswith("ABC") else "Revendidos"
            es[ks[c]] = f"={P_PINTA}*{self.area(hoja, 'FL_MAT_' + c):.4f}/{ESMALTE_REND}"
        fila("MP-65", es, "15 % de los equipos llegan con la pintura deteriorada (Dimensionamiento V2) y se pintan "
             "completos con esmalte sintético rojo RAL 3000 (planilla de MP): litros = área exterior del BOM / 8,8 m² "
             "por litro (ficha Sinteplast Sintesmal, 40 µm, una mano).", ["X11", "X15", "R11", "W1", "B02"])
        k = len(REC)
        fila("MP-82", {k: 1}, "Faja amarilla del sustituto (kit del BOM).", ["B04"])
        fila("MP-83", {k: 1}, "Tarjeta de sustituto habilitado, si va a PBA (kit del BOM).", ["B04"], "PBA")
        fila("MP-84", {k: 1}, "Tarjeta de sustituto, si va a CABA (kit del BOM).", ["B04"], "CABA")
        fila("MP-85", {k: 1}, "Remito de préstamo (kit del BOM).", ["B04"])
        return cols, rows

    def filas_rev(self):
        cols = [f"FL_MAT_{c}" for c, _ in REV]
        rows = [dict(it=f"RV-{c}", dest="", cons={k: 1}, como="Se compra el equipo terminado.", figs=["X01"])
                for k, (c, _) in enumerate(REV)]
        rows.append(dict(it="MP-77", dest="CABA", cons={k: 1 for k in range(len(REV))},
                         como="La tarjeta AGC la agrega quien vende en CABA (BOM de revendidos).", figs=["B02"]))
        return cols, rows

    # -------- hoja de una línea
    def hoja_linea(self, lin, cols, rows, pref):
        ws = self.wb.create_sheet(lin)
        _titulo(ws, f"{lin}: consumo por unidad de producto y lo que hay que comprar")
        nc = len(cols)
        c0 = 6                                  # primera columna de producto
        cy = c0 + nc                            # primera columna de año
        cp = cy + 10                            # mes pico
        nombres = []
        for cl in cols:
            if cl == "KIT":
                nombres.append("Sustituto nuevo (kit)")
            else:
                c = codigo(cl)
                nombres.append(("Recarga " if cl.startswith("RK") else "Cilindro " if cl.startswith("FL_REC")
                                else "") + NOMBRE[c])
        _cab(ws, 3, ["Ref.", "Código", "Material", "UM", "Destino"] + nombres + [str(a) for a in AÑOS] +
             ["Mes pico 2035"], [7, 11, 40, 6, 8] + [10] * nc + [10] * 10 + [10])
        # demanda arriba
        for j, a in enumerate(AÑOS):
            r = 4 + j
            ws.cell(r, 3, f"Demanda {a}")
            for k, cl in enumerate(cols):
                ws.cell(r, c0 + k, f"=DEMANDA!{L(4 + j)}{self.dem[(lin, cl)]}").number_format = "#,##0"
            _fila(ws, r, cp, "gris")
        r = 14
        ws.cell(r, 3, "Producción del mes pico de cada producto, 2035")
        for k, cl in enumerate(cols):
            if cl == "KIT":
                ws.cell(r, c0 + k, f"=DEMANDA!M{self.dem[(lin, cl)]}/12").number_format = "#,##0"
            else:
                ws.cell(r, c0 + k, f"=ESTACIONALIDAD!P{self.r_est[(lin, cl)]}").number_format = "#,##0"
        _fila(ws, r, cp, "gris")
        r = 15
        ws.cell(r, 3, "CONSUMO POR UNIDAD (cantidad del material por cada producto)")
        _fila(ws, r, cp, "sec")
        r = 16
        out = []
        share = self.r_caba_r if lin == "RECARGAS" else self.r_caba_u
        for n_, x in enumerate(rows, 1):
            it = ITEM[x["it"]]
            rid = f"{pref}{n_:02d}"
            ws.cell(r, 1, rid)
            ws.cell(r, 2, it["cod"])
            ws.cell(r, 3, it["nom"])
            ws.cell(r, 4, it["um"])
            ws.cell(r, 5, x["dest"])
            for k, v in x["cons"].items():
                ws.cell(r, c0 + k, v).number_format = "0.0000"
            a, b = L(c0), L(c0 + nc - 1)

            def fac(colyr):
                if x["dest"] == "CABA":
                    return f"*DEMANDA!{colyr}${share}"
                if x["dest"] == "PBA":
                    return f"*(1-DEMANDA!{colyr}${share})"
                return ""

            for j in range(10):
                ws.cell(r, cy + j, f"=SUMPRODUCT(${a}{r}:${b}{r},${a}${4 + j}:${b}${4 + j}){fac(L(4 + j))}"
                        ).number_format = "#,##0.0"
            ws.cell(r, cp, f"=SUMPRODUCT(${a}{r}:${b}{r},${a}$14:${b}$14){fac('M')}").number_format = "#,##0.0"
            _fila(ws, r, cp)
            self.ref(rid, hoja=lin, codigo=it["cod"], dato=it["nom"], um=it["um"], dest=x["dest"], como=x["como"],
                     figs=x["figs"], cons={cols[k]: v for k, v in x["cons"].items()})
            out.append((it["cod"], r))
            r += 1
        self.lin[lin] = dict(cy=cy, cp=cp, r0=16, r1=r - 1, filas=out, cols=cols)
        ws.freeze_panes = ws.cell(16, c0)
        _imprimir(ws, "3:3")

    # -------- consolidado
    def consolidado(self):
        ws = self.wb.create_sheet("CONSOLIDADO")
        _titulo(ws, "CONSOLIDADO: lo que hay que comprar de cada material (las cuatro líneas)")
        _cab(ws, 3, ["Código", "Material", "UM"] + [f"{l.capitalize()} 2035" for l in LINEAS] + [str(a) for a in AÑOS]
             + ["Mes pico 2035"], [11, 44, 6] + [11] * 4 + [11] * 10 + [11])
        cods = []
        for lin in LINEAS:
            for cod, _ in self.lin[lin]["filas"]:
                if cod not in cods:
                    cods.append(cod)
        orden = [it["cod"] for it in ITEMS]
        cods.sort(key=orden.index)
        self.cons = {}
        r = 4

        def sumif(lin, colx):
            d = self.lin[lin]
            return (f"SUMIF({lin}!$B${d['r0']}:$B${d['r1']},$A{r},{lin}!{colx}${d['r0']}:{colx}${d['r1']})")

        for cod in cods:
            it = ITEM[cod]
            ws.cell(r, 1, cod)
            ws.cell(r, 2, it["nom"])
            ws.cell(r, 3, it["um"])
            for k, lin in enumerate(LINEAS):
                ws.cell(r, 4 + k, "=" + sumif(lin, L(self.lin[lin]["cy"] + 9))).number_format = "#,##0.0"
            for j in range(10):
                ws.cell(r, 8 + j, "=" + "+".join(sumif(lin, L(self.lin[lin]["cy"] + j)) for lin in LINEAS)
                        ).number_format = "#,##0.0"
            ws.cell(r, 18, "=" + "+".join(sumif(lin, L(self.lin[lin]["cp"])) for lin in LINEAS)
                    ).number_format = "#,##0.0"
            _fila(ws, r, 18)
            self.cons[cod] = r
            r += 1
        ws.freeze_panes = "D4"
        ws.auto_filter.ref = f"A3:R{r - 1}"
        _imprimir(ws, "3:3")

    # -------- política de stock
    def politica(self):
        ws = self.wb.create_sheet("POLITICA_STOCK")
        _titulo(ws, "POLÍTICA DE STOCK DEL AÑO 2035")
        grupos = [(1, 5, ""), (6, 8, "CONSUMO 2035 (velocidad)"), (9, 11, "PLAZO DE ENTREGA (días hábiles)"),
                  (12, 14, "CUÁNDO PEDIR (nivel de stock)"), (15, 20, "CUÁNTO PEDIR"),
                  (21, 23, "LUGAR EN EL ALMACÉN")]
        for a, b, txt in grupos:
            ws.merge_cells(start_row=3, start_column=a, end_row=3, end_column=b)
            c = ws.cell(3, a, txt)
            c.font = _f(True, BLANCO)
            c.fill = PatternFill("solid", fgColor=GRIS_O)
            c.alignment = Alignment(horizontal="center")
        _cab(ws, 4, ["Ref.", "Código", "Material", "UM", "Proveedor",
                     "Consumo del año", "Consumo por día, promedio\n(año / 240)",
                     "Consumo por día, mes pico\n(mes pico / 20)",
                     "Mínimo", "Máximo", "De dónde sale el plazo",
                     "Consumo mientras llega el pedido\n(promedio × plazo medio)",
                     "Stock de seguridad\n(pico × plazo máximo − consumo mientras llega)",
                     "Punto de pedido: cuando quedan esto, se pide\n(consumo mientras llega + seguridad)",
                     "Unidad de compra", "Contenido de la unidad de compra", "Cada cuántos días se pide\n(planilla de MP)",
                     "Cantidad por pedido (unidades de compra)", "Pedidos por año",
                     "Stock máximo: lo que hay al llegar el pedido\n(seguridad + pedido)",
                     "Almacén", "Cuánto entra en un lugar (pallet, estante o cuna)", "Lugares que ocupa el stock máximo"],
             [7, 11, 38, 6, 18, 11, 12, 12, 8, 8, 70, 14, 16, 16, 22, 11, 11, 11, 9, 15, 30, 12, 11])
        ws.row_dimensions[4].height = 66
        r = 5
        n_ = 0
        zonas = {}
        for cod, rc in self.cons.items():
            it = ITEM[cod]
            if it["alm"] == Z_SERV:
                continue
            n_ += 1
            rid = f"P{n_:02d}"
            lmin, lmax, lsrc, lfig = plazo(cod)
            vals = [rid, cod, it["nom"], it["um"], it["prov"], f"=CONSOLIDADO!Q{rc}", f"=F{r}/240",
                    f"=CONSOLIDADO!R{rc}/20", lmin, lmax, lsrc, f"=G{r}*(I{r}+J{r})/2", f"=H{r}*J{r}-L{r}",
                    f"=L{r}+M{r}", it["uc"], it["cont"], it["t"], f"=MAX(1,CEILING(H{r}*Q{r}/P{r},1))",
                    f"=F{r}/(R{r}*P{r})", f"=M{r}+R{r}*P{r}", it["alm"], it["umpos"], f"=T{r}/V{r}"]
            for j, v in enumerate(vals, 1):
                c = ws.cell(r, j, v)
                if j in (6, 12, 13, 14, 20):
                    c.number_format = "#,##0.0"
                elif j in (7, 8):
                    c.number_format = "#,##0.00"
                elif j in (16, 22):
                    c.number_format = "#,##0.##"
                elif j in (19, 23):
                    c.number_format = "0.0"
            links = [u for u in LINKS.get(lfig, ()) if u]
            if links:
                c = ws.cell(r, 11, ("Aproximado: " if lsrc.startswith("Aproximado") else "") +
                            f"{NOMBRE_FUENTE.get(lfig, lsrc.split(':')[0])}: {links[0]}")
                c.hyperlink = links[0]
                c.font = Font(name=AR, size=10, color=NEGRO, underline="single")
                c.comment = Comment(lsrc + (f". Otro link: {links[1]}" if len(links) > 1 else ""), "FLAMA",
                                    height=140, width=320)
            _fila(ws, r, 23)
            self.ref(rid, hoja="POLITICA_STOCK", codigo=cod, dato=it["nom"], uc=it["uc"], fuc=it["fuc"],
                     plan=it["plan"], lt=(lmin, lmax), flt=lsrc, flt_fig=lfig, t=it["t"], umpos=it["umpos"],
                     fpos=it["fpos"])
            zonas.setdefault(it["alm"], []).append(r)
            r += 1
        fin = r - 1
        r += 2
        _cab(ws, r, ["", "Almacén", "", "", "", "Lugares que ocupa el stock máximo", "Lugares que previó el proceso",
                     "¿Alcanza?"])
        self.r_alm0 = r + 1
        for z in sorted(zonas):
            r += 1
            ws.cell(r, 2, z)
            ws.cell(r, 6, f'=ROUNDUP(SUMIF($U$5:$U${fin},B{r},$W$5:$W${fin}),0)')
            if z in CAPACIDAD:
                ws.cell(r, 7, CAPACIDAD[z])
                ws.cell(r, 8, f'=IF(F{r}<=G{r},"Alcanza","No alcanza")')
            _fila(ws, r, 8)
        self.ref("P00", hoja="POLITICA_STOCK", dato="Fórmulas de la política",
                 como="Consumo mientras llega el pedido = consumo promedio por día × plazo medio. Stock de seguridad = "
                      "consumo del mes pico × plazo máximo − consumo mientras llega el pedido (cubre el mes pico y la "
                      "demora máxima). Punto de pedido = consumo mientras llega + stock de seguridad. Cantidad por "
                      "pedido = consumo del mes pico durante los días entre pedidos, en unidades de compra enteras. "
                      "Stock máximo = seguridad + pedido. Lugares = stock máximo / lo que entra en un lugar.",
                 figs=["X07", "Q01", "Q03", "Q04"])
        ws.freeze_panes = "D5"
        ws.auto_filter.ref = f"A4:W{fin}"
        _imprimir(ws, "3:4")

    def armar(self, ruta):
        self.ix = self.copia("Indexado Dim Técnico", "PROY_INDEXADO")
        self.rs = self.copia("Reservas y Sustitutos", "PROY_RESERVAS")
        self.copia_usados()
        self.demanda()
        self.estacionalidad()
        for lin, grupo, pref in (("MANUALES", MAN, "M"), ("RODANTES", ROD, "R")):
            cols, rows = self.filas_fab(lin, grupo)
            self.hoja_linea(lin, cols, rows, pref)
        cols, rows = self.filas_rec()
        self.hoja_linea("RECARGAS", cols, rows, "K")
        cols, rows = self.filas_rev()
        self.hoja_linea("REVENDIDOS", cols, rows, "V")
        self.consolidado()
        self.politica()
        orden = ["DEMANDA", "ESTACIONALIDAD", "MANUALES", "RODANTES", "RECARGAS", "REVENDIDOS", "CONSOLIDADO",
                 "POLITICA_STOCK", "PROY_INDEXADO", "PROY_RESERVAS", "PROY_USADOS"]
        self.wb._sheets = [self.wb[n] for n in orden]
        for ws in self.wb.worksheets:
            ws.sheet_view.zoomScale = 90
        self.wb.active = 0
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.save(ruta)
        with open(os.path.splitext(ruta)[0] + "_refs.json", "w", encoding="utf-8") as fh:
            json.dump(dict(refs=self.refs, lineas={k: dict(cols=v["cols"], cy=v["cy"], cp=v["cp"], r0=v["r0"],
                                                           r1=v["r1"]) for k, v in self.lin.items()},
                           dem={f"{a}|{b}": r for (a, b), r in self.dem.items()}, est={f"{a}|{b}": r for (a, b), r
                                                                                       in self.r_est.items()},
                           pico=self.r_pico, cons=self.cons, alm0=self.r_alm0, items=ITEM),
                      fh, ensure_ascii=False, indent=1)


def excel(proyecciones, ruta, bom=None):
    bom = bom or os.path.join(os.path.dirname(os.path.dirname(ruta)), "bom", "FLAMA_BOM.xlsx")
    Libro(proyecciones, bom).armar(ruta)
    recalcular(ruta)
    return ruta
