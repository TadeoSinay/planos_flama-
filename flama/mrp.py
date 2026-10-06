"""MRP de materias primas (requerimiento bruto anual 2026-2035) a partir de las proyecciones de FLAMA y del BOM.

Las dos primeras hojas son copia (valores) de «Indexado Dim Técnico» y «Reservas y Sustitutos» del Excel de
proyecciones; todo lo demás son fórmulas que leen esas dos hojas, de modo que si se pegan valores nuevos el MRP se
recalcula. Los coeficientes por unidad salen del mismo BOM (bom.maestro_datos: sólo lo que se compra).

Uso: python generar.py --mrp <Proyecciones.xlsx>   ->   salida/mrp/MRP_MP.xlsx
"""

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

from . import bom as B
from . import planos as P
from . import recipientes as RC
from . import sustituto as SU
from .catalogo import MODELOS

IX, RS, PA, DE = "Indexado Dim Técnico", "Reservas y Sustitutos", "Parametros", "Demanda"
q = lambda s: f"'{s}'"
YEARS = "BCDEFGHIJK"            # columnas de años en Indexado / Reservas / Parametros / Demanda (2026..2035)
ZONAL = "LMNOPQRSTU"            # columnas del % zonal por año en el Indexado
YC = [L(6 + i) for i in range(10)]          # F..O: requerimiento total por año en las hojas de MP
COEF0 = 37                                  # AK: primera columna de coeficientes / demanda por producto

AR = "Arial"
F_N, F_B = Font(name=AR, size=10), Font(name=AR, size=10, bold=True)
F_T, F_I = Font(name=AR, size=12, bold=True), Font(name=AR, size=9, italic=True)
F_W = Font(name=AR, size=10, bold=True, color="FFFFFF")
F_LINK = Font(name=AR, size=10, color="008000")
CAB = PatternFill("solid", fgColor="7F1D1D")
AMAR = PatternFill("solid", fgColor="FFFF00")
EST = {"E": PatternFill("solid", fgColor="FFF2CC"), "A": PatternFill("solid", fgColor="F8CBAD"),
       "X": PatternFill("solid", fgColor="FF8B8B")}

# ---------------- productos y filas de demanda
FAB = [("FL_MAT_ABC_1kg", 3, None), ("FL_MAT_ABC_2.5kg", 4, None), ("FL_MAT_ABC_5kg", 5, ("30", "69")),
       ("FL_MAT_ABC_10kg", 6, (None, "70")), ("FL_MAT_ABC_25kg", 7, None), ("FL_MAT_ABC_50kg", 8, (None, "79")),
       ("FL_MAT_ABC_70kg", 9, None), ("FL_MAT_ABC_100kg", 10, (None, "80"))]
REV = [("FL_MAT_AGUA_10l", 48, "71"), ("FL_MAT_AFFF_10l", 49, "72"), ("FL_MAT_AFFF_50l", 50, "73"),
       ("FL_MAT_BC_5kg", 51, "74"), ("FL_MAT_SALESK_6l", 52, "75"), ("FL_MAT_CO2_2kg", 53, None),
       ("FL_MAT_CO2_5kg", 54, "76"), ("FL_MAT_HCFC-HFC_5kg", 55, "77"), ("FL_MAT_CLASED_9l", 56, "78")]
CIL = [(f"FL_REC_ABC_{c}", 15 + i) for i, c in enumerate(["1kg", "2.5kg", "5kg", "10kg", "25kg", "50kg", "70kg",
                                                            "100kg"])]
# kit de sustituto por pool (fila de alta en «Reservas y Sustitutos»)
SUS = [("FL_SUS_ABC_5kg", 69), ("FL_SUS_ABC_10kg", 70), ("FL_SUS_AGUA_10l", 71), ("FL_SUS_AFFF_10l", 72),
       ("FL_SUS_AFFF_50l", 73), ("FL_SUS_BC_5kg", 74), ("FL_SUS_SALESK_6l", 75), ("FL_SUS_CO2_5kg", 76),
       ("FL_SUS_HCFC-HFC_5kg", 77), ("FL_SUS_CLASED_9l", 78), ("FL_SUS_ABC_50kg", 79), ("FL_SUS_ABC_100kg", 80)]
# recargas por tipo (fila del Indexado) + recarga anual del stock de sustitutos del pool (fila en Reservas)
REC = [("FL_MAT_ABC_1kg", 27, None), ("FL_MAT_ABC_2.5kg", 28, None), ("FL_MAT_ABC_5kg", 29, "54"),
       ("FL_MAT_ABC_10kg", 30, "55"), ("FL_MAT_AGUA_10l", 31, "56"), ("FL_MAT_AFFF_10l", 32, "57"),
       ("FL_MAT_AFFF_50l", 33, "58"), ("FL_MAT_BC_5kg", 34, "59"), ("FL_MAT_SALESK_6l", 35, "60"),
       ("FL_MAT_CO2_2kg", 36, None), ("FL_MAT_CO2_5kg", 37, "61"), ("FL_MAT_HCFC-HFC_5kg", 38, "62"),
       ("FL_MAT_CLASED_9l", 39, "63"), ("FL_MAT_ABC_25kg", 40, None), ("FL_MAT_ABC_50kg", 41, "64"),
       ("FL_MAT_ABC_70kg", 42, None), ("FL_MAT_ABC_100kg", 43, "65")]
COMERCIAL = (61, 62, 64, 68, 69, 70)        # zonas de la planta comercial (Indexado V60)
CABA_REC, CABA_NUE = (61, 62, 63, 64), (79, 80, 81, 82)


def _copiar(ws_src, ws):
    for row in ws_src.iter_rows():
        for c in row:
            if c.value is None:
                continue
            d = ws.cell(c.row, c.column, c.value)
            d.number_format = c.number_format
            d.font = Font(name=AR, size=10, bold=bool(c.font and c.font.b))
    for k, dim in ws_src.column_dimensions.items():
        if dim.width:
            ws.column_dimensions[k].width = dim.width


def _cab(ws, r, textos, c0=1):
    for j, t in enumerate(textos, c0):
        c = ws.cell(r, j, t)
        c.font, c.fill = F_W, CAB
        c.alignment = Alignment(wrap_text=True, vertical="center")


def _parametros(wb):
    ws = wb.create_sheet(PA)
    ws["A1"] = "Parámetros del MRP"
    ws["A1"].font = F_T
    ws.column_dimensions["A"].width = 70
    ws["A3"], ws["B3"] = "Fracción de recargas con agente 100 % nuevo (equipo accionado)", 1
    ws["B3"].fill, ws["B3"].number_format = AMAR, "0%"
    ws["C3"] = ("IRAM 3517-2:2020 9.9.1.4: el polvo de un extintor accionado no se reutiliza. 100 % = cota superior. "
                "El resto de las recargas usa el caso B del BOM (recuperado + reposición de mermas). Ajustar con el dato "
                "real de la planta de recargas.")
    ws["C3"].font = F_I
    r = 6
    ws.cell(r, 1, "Año").font = F_B
    for i, c in enumerate(YEARS):
        ws[f"{c}{r}"] = f"={q(IX)}!{c}26"
        ws[f"{c}{r}"].font = F_B
    filas = [(7, "% destino CABA — unidades nuevas (zonal de revendidos, filas 79-82)", CABA_NUE),
             (8, "% destino CABA — recargas (zonal de recargas, filas 61-64)", CABA_REC),
             (9, "% recargas planta COMERCIAL (zonas de Indexado V60: CABA N/E/O, PBA Norte 1°-3°)", COMERCIAL)]
    for rr, lab, zonas in filas:
        ws.cell(rr, 1, lab)
        for i, c in enumerate(YEARS):
            ws[f"{c}{rr}"] = "=" + "+".join(f"{q(IX)}!{ZONAL[i]}{z}" for z in zonas)
            ws[f"{c}{rr}"].number_format = "0.0%"
    ws.cell(10, 1, "% recargas planta INDUSTRIAL (CABA Sur, PBA Sur 1°-3°, PBA Oeste 1°-3°, Interior)")
    for c in YEARS:
        ws[f"{c}10"] = f"=1-{c}9"
        ws[f"{c}10"].number_format = "0.0%"
    ws["A12"] = ("Destino: lo que exige una norma de PBA (oblea Res. 522/07, tarjeta DPS Res. 349/07) se aplica a la "
                 "proporción PBA; lo de CABA (tarjeta AGC, Ord. 40.473) a la proporción CABA. Para unidades nuevas se "
                 "usa la distribución zonal de revendidos (única de unidades nuevas en el Indexado).")
    ws["A12"].font = F_I
    for row in ws.iter_rows(min_row=3, max_row=12):
        for c in row:
            if c.font is None or c.font.name != AR:
                c.font = Font(name=AR, size=10, bold=c.font.b if c.font else False, italic=c.font.i if c.font else False)
    return ws


def _demanda(wb):
    """Hoja Demanda: unidades por producto y año. Devuelve {grupo: [(código, fila)]}."""
    ws = wb.create_sheet(DE)
    ws["A1"] = "Demanda por producto (unidades/año) — fórmulas sobre las dos primeras hojas"
    ws["A1"].font = F_T
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["L"].width = 80
    ws["A3"] = "Producto"
    ws["A3"].font = F_B
    for i, c in enumerate(YEARS):
        ws[f"{c}3"] = f"={q(IX)}!{c}26"
        ws[f"{c}3"].font = F_B
    ws["L3"] = "Origen"
    ws["L3"].font = F_B
    out, r = {}, 4

    def bloque(titulo, filas, grupo):
        nonlocal r
        r += 1
        ws.cell(r, 1, titulo).font = F_B
        out[grupo] = []
        for cod, f_ix, extra, origen in filas:
            r += 1
            ws.cell(r, 1, cod)
            for c in YEARS:
                partes = [f"{q(IX)}!{c}{f_ix}"] if f_ix else []
                partes += [f"{q(RS)}!{c}{e}" for e in extra if e]
                ws[f"{c}{r}"] = "=" + "+".join(partes)
                ws[f"{c}{r}"].number_format = "#,##0"
            ws.cell(r, 12, origen).font = F_I
            out[grupo].append((cod, r))
        r += 1

    bloque("MATAFUEGOS FABRICADOS (unidades nuevas + venta de reserva + alta de sustitutos)",
           [(c, f, ex or (), f"Indexado fila {f}" + ("".join(f" + Reservas fila {e}" for e in (ex or ()) if e)))
            for c, f, ex in FAB], "FAB")
    bloque("MATAFUEGOS REVENDIDOS (compra de unidades nuevas + alta de sustitutos)",
           [(c, f, (e,), f"Indexado fila {f}" + (f" + Reservas fila {e}" if e else "")) for c, f, e in REV], "REV")
    bloque("CILINDROS SUELTOS (venta)", [(c, f, (), f"Indexado fila {f}") for c, f in CIL], "CIL")
    bloque("KITS DE SUSTITUTO (alta anual de sustitutos por pool)",
           [(c, None, (str(f),), f"Reservas fila {f}") for c, f in SUS], "SUS")
    bloque("RECARGAS (Indexado + mantenimiento anual del stock de sustitutos)",
           [(c, f, (e,), f"Indexado fila {f}" + (f" + Reservas fila {e}" if e else "")) for c, f, e in REC], "REC")
    ws.freeze_panes = "B4"
    return out


def _hoja_mp(wb, nombre, titulo, nota, productos, coef, meta, dem_filas, fila_dest, recargas=False):
    """Hoja de requerimiento: ítems × años (F..O) = SUMPRODUCT(coeficientes × demanda del año) × destino."""
    ws = wb.create_sheet(nombre)
    ws["A1"], ws["A2"] = titulo, nota
    ws["A1"].font, ws["A2"].font = F_T, F_I
    n = len(productos)
    cols = [L(COEF0 + j) for j in range(n)]
    # demanda por año (filas 5..14) alineada con las columnas de coeficientes
    ws.cell(4, COEF0 - 1, "Demanda (u) →").font = F_B
    for j, (cod, _) in enumerate(productos):
        c = ws.cell(4, COEF0 + j, cod.replace("FL_MAT_", "").replace("FL_REC_", "REC ").replace("FL_SUS_", "SUS "))
        c.font, c.fill = F_W, CAB
        c.alignment = Alignment(text_rotation=90)
    ws.row_dimensions[4].height = 90
    for i, yc in enumerate(YEARS):
        ws.cell(5 + i, COEF0 - 1, f"={q(DE)}!{yc}3").font = F_B
        for j, (cod, fila) in enumerate(productos):
            ws.cell(5 + i, COEF0 + j, f"={q(DE)}!{yc}{fila}").number_format = "#,##0"
    # encabezado de la tabla de ítems
    h = 16
    cab = ["Ítem", "UM", "Rubro", "Proveedor principal", "Destino"]
    _cab(ws, h, cab)
    for i, yc in enumerate(YEARS):
        c = ws.cell(h, 6 + i, f"={q(DE)}!{yc}3")
        c.font, c.fill = F_W, CAB
    for j, (cod, _) in enumerate(productos):
        c = ws.cell(h, COEF0 + j, "coef. " + cod.replace("FL_MAT_", "").replace("FL_REC_", "REC ")
                    .replace("FL_SUS_", "SUS "))
        c.font, c.fill = F_W, CAB
        c.alignment = Alignment(text_rotation=90, wrap_text=True)
    ws.row_dimensions[h].height = 90
    for k, w in zip("ABCDE", (62, 10, 22, 30, 9)):
        ws.column_dimensions[k].width = w
    for c in YC:
        ws.column_dimensions[c].width = 12
    if recargas:
        ws.cell(h - 1, 6, "TOTAL").font = F_B
        ws.cell(h - 1, 17, "PLANTA COMERCIAL").font = F_B
        ws.cell(h - 1, 27, "PLANTA INDUSTRIAL").font = F_B
        for blk, c0 in ((1, 16), (2, 26)):
            for i, yc in enumerate(YEARS):
                c = ws.cell(h, c0 + 1 + i, f"={q(DE)}!{yc}3")
                c.font, c.fill = F_W, CAB
    items = [k for k in sorted(coef, key=lambda k: (B.RUBROS.index(meta[k]["rubro"]), k[0]))
             if any(coef[k].get(cod) for cod, _ in productos)]
    r = h
    for k in items:
        r += 1
        d = meta[k]
        for j, v in enumerate([k[0], k[1], d["rubro"], d["prov"] or None, d["dest"] or "Todos"], 1):
            ws.cell(r, j, v).font = F_N
        if EST.get(d["est"]):
            for j in range(1, 6):
                ws.cell(r, j).fill = EST[d["est"]]
        for j, (cod, _) in enumerate(productos):
            v = coef[k].get(cod)
            if v not in (None, 0, ""):
                c = ws.cell(r, COEF0 + j, v)
                c.font = Font(name=AR, size=10, color="0000FF") if not isinstance(v, str) else F_N
                c.number_format = "0.0000"
        rng = f"${cols[0]}{r}:${cols[-1]}{r}"
        for i in range(10):
            dem = f"${cols[0]}${5 + i}:${cols[-1]}${5 + i}"
            pc = f"{q(PA)}!{YEARS[i]}${fila_dest}"
            dest = f'IF($E{r}="CABA",{pc},IF($E{r}="PBA",1-{pc},1))'
            c = ws.cell(r, 6 + i, f"=SUMPRODUCT({rng},{dem})*{dest}")
            c.number_format = "#,##0.00"
            if recargas:
                ws.cell(r, 17 + i, f"={YC[i]}{r}*{q(PA)}!{YEARS[i]}$9").number_format = "#,##0.00"
                ws.cell(r, 27 + i, f"={YC[i]}{r}*{q(PA)}!{YEARS[i]}$10").number_format = "#,##0.00"
    ws.freeze_panes = ws.cell(h + 1, 6)
    ws.auto_filter.ref = f"A{h}:E{r}"
    return ws, items


def _insumos(wb, term, cil, dem):
    ws = wb.create_sheet("MP_INSUMOS_PROCESO")
    ws["A1"] = "Insumos de proceso (no van por unidad: se planifican por su inductor)"
    ws["A1"].font = F_T
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["L"].width = 70
    ws["A3"] = "Minutos de arco MAG por unidad (del BOM: gas C2 / 14 L/min)"
    ws["A3"].font = F_B
    _cab(ws, 4, ["Producto", "min/u"] + [""] * 0)
    for i, yc in enumerate(YEARS):
        c = ws.cell(4, 3 + i, f"={q(DE)}!{yc}3")
        c.font, c.fill = F_W, CAB
    r = 4
    filas_min = []
    prod = {c: f for c, f in dem["FAB"]} | {c: f for c, f in dem["CIL"]}
    for m, f in list(term) + list(cil):
        cod = m.codigo if (m, f) in term else RC.codigo_rec(m)
        if cod not in prod:
            continue
        mins = next((x["cant"] / B.CAUDAL_GAS for x in f if x["codigo"].endswith("-C2") and x["cant"]), None)
        if not mins:
            continue
        r += 1
        ws.cell(r, 1, cod)
        ws.cell(r, 2, round(mins, 3)).font = Font(name=AR, size=10, color="0000FF")
        for i, yc in enumerate(YEARS):
            ws.cell(r, 3 + i, f"=$B{r}*{q(DE)}!{yc}{prod[cod]}").number_format = "#,##0"
        filas_min.append(r)
    r += 1
    ws.cell(r, 1, "Total horas de arco / año").font = F_B
    for i in range(10):
        c = L(3 + i)
        ws.cell(r, 3 + i, f"=SUM({c}{filas_min[0]}:{c}{filas_min[-1]})/60").number_format = "#,##0"
    tot = r
    r += 2
    _cab(ws, r, ["Insumo", "Planilla / norma", "Proveedor", "Consumo por inductor (completar)", "UM"])
    r0 = r
    ins = [("Puntas de contacto CuCrZr Ø0,9 / Ø1,2", "Planilla MP-26", "ESAB / Binzel vía distribuidor",
            "puntas por hora de arco"),
           ("Toberas MAG", "Planilla MP-26", "ESAB / Binzel vía distribuidor", "toberas por hora de arco"),
           ("Discos de amolado y corte Ø115", "Planilla MP-27", "Norton / Pferd / 3M", "discos por hora de arco"),
           ("Tapones de silicona para enmascarar roscas", "Planilla MP-32", "Distribuidor de enmascarado",
            "tapones por hora de arco (proxy de piezas)"),
           ("Solución detectora de fugas", "IRAM 3517-2:2020 9.4.10", "A relevar", "L por hora de arco (proxy)"),
           ("Esmalte rojo de retoque (recargas)", "Planilla MP-65", "Sinteplast / Tersuave", "L por hora de arco")]
    for nom, ref, prov, um in ins:
        r += 1
        for j, v in enumerate([nom, ref, prov, None, um], 1):
            ws.cell(r, j, v)
        ws.cell(r, 4).fill = AMAR
    r += 2
    ws.cell(r, 1, "Requerimiento anual = horas de arco × consumo por inductor (0 hasta completar la columna D)").font = F_I
    r += 1
    _cab(ws, r, ["Insumo"])
    for i, yc in enumerate(YEARS):
        c = ws.cell(r, 2 + i, f"={q(DE)}!{yc}3")
        c.font, c.fill = F_W, CAB
    for k in range(len(ins)):
        r += 1
        ws.cell(r, 1, f"=A{r0 + 1 + k}")
        for i in range(10):
            ws.cell(r, 2 + i, f"={L(3 + i)}${tot}*$D${r0 + 1 + k}").number_format = "#,##0.0"
    return ws


def _consolidado(wb, hojas):
    ws = wb.create_sheet("MRP_CONSOLIDADO")
    ws["A1"] = "MRP consolidado de materias primas e insumos (requerimiento bruto anual, todas las categorías)"
    ws["A1"].font = F_T
    ws["A2"] = "Suma por ítem y UM de las hojas MP_* (fórmulas SUMIFS). El detalle por categoría está en cada hoja."
    ws["A2"].font = F_I
    h = 4
    _cab(ws, h, ["Ítem", "UM", "Rubro", "Proveedor principal", "Categorías"])
    for i, yc in enumerate(YEARS):
        c = ws.cell(h, 6 + i, f"={q(DE)}!{yc}3")
        c.font, c.fill = F_W, CAB
    for k, w in zip("ABCDE", (62, 10, 22, 30, 40)):
        ws.column_dimensions[k].width = w
    todos = {}
    for nombre, items, meta in hojas:
        for k in items:
            todos.setdefault(k, (meta[k], []))[1].append(nombre)
    r = h
    for k in sorted(todos, key=lambda k: (B.RUBROS.index(todos[k][0]["rubro"]), k[0])):
        r += 1
        d, cats = todos[k]
        for j, v in enumerate([k[0], k[1], d["rubro"], d["prov"] or None, ", ".join(cats)], 1):
            ws.cell(r, j, v).font = F_N
        for i, yc in enumerate(YC):
            f = "+".join(f"SUMIFS({q(n)}!{yc}:{yc},{q(n)}!$A:$A,$A{r},{q(n)}!$B:$B,$B{r})" for n in cats)
            ws.cell(r, 6 + i, "=" + f).number_format = "#,##0.00"
    ws.freeze_panes = "F5"
    ws.auto_filter.ref = f"A{h}:E{r}"
    return ws


def excel(proyecciones, ruta):
    src = load_workbook(proyecciones, data_only=True)
    term = [(m, B.bom_producto(m)) for m in MODELOS]
    cil = [(m, B.bom_producto(m, cilindro=True)) for m in RC.modelos_abc()]
    columnas, datos, meta = B.maestro_datos(term, cil)
    wb = Workbook()
    ws = wb.active
    ws.title = IX
    _copiar(src[IX], ws)
    _copiar(src[RS], wb.create_sheet(RS))
    _parametros(wb)
    dem = _demanda(wb)

    def coefs(codigos):
        return {k: {c: v[c] for c in codigos if v.get(c)} for k, v in datos.items()}

    hojas = []
    # fabricados
    prods = dem["FAB"]
    co = coefs([c for c, _ in prods])
    _, it = _hoja_mp(wb, "MP_MATAFUEGOS_FABRICADOS", "Requerimiento de MP — matafuegos fabricados (ABC 1 a 100 kg)",
                     "Demanda = unidades nuevas + venta de reserva ABC 5 kg + alta de sustitutos. Coeficientes del BOM "
                     "(azul). Destino CABA/PBA según % zonal (Parametros fila 7).", prods, co, meta, None, 7)
    hojas.append(("MP_MATAFUEGOS_FABRICADOS", it, meta))
    # revendidos
    prods = dem["REV"]
    co = coefs([c for c, _ in prods])
    _, it = _hoja_mp(wb, "MP_MATAFUEGOS_REVENDIDOS", "Requerimiento — matafuegos revendidos (compra del equipo "
                     "terminado)", "Sólo se compra el equipo terminado y, si va a CABA, la tarjeta AGC.", prods, co,
                     meta, None, 7)
    hojas.append(("MP_MATAFUEGOS_REVENDIDOS", it, meta))
    # cilindros
    prods = dem["CIL"]
    co = coefs([c for c, _ in prods])
    _, it = _hoja_mp(wb, "MP_CILINDROS", "Requerimiento de MP — cilindros sueltos FL_REC (venta)",
                     "Recipiente, tapón, etiqueta, protocolo y embalaje.", prods, co, meta, None, 7)
    hojas.append(("MP_CILINDROS", it, meta))
    # sustitutos: kit de identificación (el extintor base ya está en fabricados / revendidos)
    prods = dem["SUS"]
    co = coefs([c for c, _ in prods])
    _, it = _hoja_mp(wb, "MP_SUSTITUTOS", "Requerimiento — kits de extintor sustituto (IRAM 3517-2:2020 9.4.5)",
                     "Faja amarilla, tarjetas DPS / AGC y remito. El extintor base se suma a la demanda de fabricados o "
                     "revendidos.", prods, co, meta, None, 7)
    hojas.append(("MP_SUSTITUTOS", it, meta))
    # recargas: coeficiente = p × caso A + (1 − p) × caso B, con p en Parametros!B3
    prods = [(c, f) for c, f in dem["REC"]]
    co = {}
    for k, v in datos.items():
        fila = {}
        for cod, _ in prods:
            m = next(x for x in MODELOS if x.codigo == cod)
            base = P.codigo_pieza(m, 0)[:-3]
            a, b = v.get(f"RK-{base}-A", 0), v.get(f"RK-{base}-B", 0)
            if a or b:
                fila[cod] = f"={q(PA)}!$B$3*{round(a, 6)}+(1-{q(PA)}!$B$3)*{round(b, 6)}"
        if fila:
            co[k] = fila
    _, it = _hoja_mp(wb, "MP_RECARGAS", "Requerimiento — recargas (total y dividido por planta comercial / industrial)",
                     "Coeficiente = p × caso A (descargado, agente nuevo) + (1 − p) × caso B (mantenimiento: recuperado "
                     "+ reposición de mermas), p en Parametros!B3. Planta comercial / industrial según el % zonal de "
                     "recargas del año (Parametros filas 9 y 10).", prods, co, meta, None, 8, recargas=True)
    hojas.append(("MP_RECARGAS", it, meta))
    _insumos(wb, term, cil, dem)
    _consolidado(wb, hojas)
    for w in wb.worksheets:
        w.sheet_properties.tabColor = {"MRP_CONSOLIDADO": "7F1D1D", PA: "FFFF00", DE: "1F4E79"}.get(w.title)
    wb.save(ruta)
    return ruta
