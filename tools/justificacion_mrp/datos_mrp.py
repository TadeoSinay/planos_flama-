"""Datos del Word «Justificación_MRP»: los valores salen del FLAMA_MRP.xlsx recalculado y la justificación de cada Ref.
del FLAMA_MRP_refs.json que escribe el generador (así el Word y el Excel salen de las mismas fórmulas). Agrega el
contraste con los números de los procesos.

Uso: python datos_mrp.py <FLAMA_MRP.xlsx recalculado> <salida.json>
"""
import json
import os
import sys

from openpyxl import load_workbook

AQUI = os.path.dirname(os.path.abspath(__file__))
XLSX, OUT = sys.argv[1], sys.argv[2]
META = json.load(open(os.path.splitext(XLSX)[0] + "_refs.json", encoding="utf-8"))
wb = load_workbook(XLSX, data_only=True)
REFS = META["refs"]
LINEAS = ("MANUALES", "RODANTES", "RECARGAS", "REVENDIDOS")


def r(x, n=4):
    return round(x, n) if isinstance(x, (int, float)) else x


# ------------------------------------------------------------------ demanda
ws = wb["DEMANDA"]
demanda = []
for f in ws.iter_rows(min_row=4, values_only=True):
    if f[0]:
        demanda.append(dict(ref=f[0], linea=f[1] or "", producto=f[2], v26=r(f[3]), v30=r(f[7]), v35=r(f[12]),
                            **{k: REFS[f[0]].get(k) for k in ("como", "figs")}))

# ------------------------------------------------------------------ estacionalidad
ws = wb["ESTACIONALIDAD"]
perfiles, picos = [], {}
for f in ws.iter_rows(min_row=4, values_only=True):
    if f[0] and str(f[0]).startswith("E") and isinstance(f[2], (int, float)):
        perfiles.append(dict(ref=f[0], nombre=f[1], meses=[r(x, 3) for x in f[2:14]],
                             **{k: REFS[f[0]].get(k) for k in ("como", "figs")}))
    elif f[0] and str(f[0]).startswith("E") and f[1] in LINEAS:
        picos[f[1]] = dict(ref=f[0], mes=f[2], unidades=r(f[4], 0), promedio=r(f[5], 0), factor=r(f[6], 3),
                           **{k: REFS[f[0]].get(k) for k in ("como", "figs")})

# ------------------------------------------------------------------ líneas
lineas = {}
for lin in LINEAS:
    d = META["lineas"][lin]
    ws = wb[lin]
    nombres = [ws.cell(3, 6 + k).value for k in range(len(d["cols"]))]
    dem35 = [ws.cell(13, 6 + k).value for k in range(len(d["cols"]))]
    filas = []
    for rr in range(d["r0"], d["r1"] + 1):
        ref = ws.cell(rr, 1).value
        cons = {nombres[k]: r(ws.cell(rr, 6 + k).value, 5) for k in range(len(d["cols"]))
                if isinstance(ws.cell(rr, 6 + k).value, (int, float))}
        filas.append(dict(ref=ref, codigo=ws.cell(rr, 2).value, material=ws.cell(rr, 3).value, um=ws.cell(rr, 4).value,
                          destino=ws.cell(rr, 5).value or "", cons=cons, v35=r(ws.cell(rr, d["cy"] + 9).value, 1),
                          pico=r(ws.cell(rr, d["cp"]).value, 1), como=REFS[ref]["como"], figs=REFS[ref]["figs"]))
    lineas[lin] = dict(productos=nombres, demanda35=[r(x, 0) for x in dem35], filas=filas)

# ------------------------------------------------------------------ consolidado y política
ws = wb["CONSOLIDADO"]
cons = {}
for f in ws.iter_rows(min_row=4, values_only=True):
    if f[0]:
        cons[f[0]] = dict(material=f[1], um=f[2], lineas=[r(x, 1) for x in f[3:7]], v26=r(f[7], 1), v35=r(f[16], 1),
                          pico=r(f[17], 1))
ws = wb["POLITICA_STOCK"]
pol, alm = [], []
for f in ws.iter_rows(min_row=5, values_only=True):
    if f[0] and str(f[0]).startswith("P"):
        x = REFS[f[0]]
        pol.append(dict(ref=f[0], codigo=f[1], material=f[2], um=f[3], proveedor=f[4], v35=r(f[5], 1),
                        dprom=r(f[6], 2), dpico=r(f[7], 2), ltmin=f[8], ltmax=f[9], fuente=x["flt"], enlt=r(f[11], 1),
                        ss=r(f[12], 1), pp=r(f[13], 1), uc=f[14], cont=f[15], t=f[16], pedido=f[17],
                        pedidos=r(f[18], 1), smax=r(f[19], 1), almacen=f[20], umpos=f[21], lugares=r(f[22], 2),
                        fig=x["flt_fig"], fuc=x["fuc"]))
    elif f[1] and not f[0] and f[1] != "Almacén" and isinstance(f[5], (int, float)):
        alm.append(dict(almacen=f[1], pos=f[5], cap=f[6], res=f[7] or ""))


def dia(cod, lin=None):
    """Consumo por día en el mes pico 2035 (mes pico / 20), de una línea o del consolidado."""
    if lin:
        x = next(f for f in lineas[lin]["filas"] if f["codigo"] == cod)
        return x["pico"] / 20
    return cons[cod]["pico"] / 20


pm = picos["MANUALES"]
validacion = [
    ("Caños Ø76,2 por día", dia("MP-01"), 53, "Proceso de manuales, corte de tubo (P03). La recepción dice 97 (P01), "
     "pero 1.216 tramos / 23 = 53."),
    ("Hojas de 1,6 mm (5 kg) por día", dia("MP-03"), 37, "Proceso de manuales, recepción (P01)"),
    ("Granalla, kg por día", dia("MP-28"), 133, "Proceso de manuales, granallado (P04)"),
    ("Nitrógeno de manuales, m³ por día", dia("MP-34", "MANUALES"), 7.1, "Proceso de manuales, presurización (P04)"),
    ("Gas MAG de manuales, m³ por día", dia("MP-25", "MANUALES"), 48, "Proceso de manuales, recepción (P01)"),
    ("Polvo ABC de manuales, kg por día", dia("MP-33", "MANUALES"), 1254, "Proceso de manuales, carga de polvo"),
    ("Recipientes de manuales por día", pm["unidades"] / 20, 1665, "Proceso de manuales (tabla de tiempos)"),
    ("Factor de pico de manuales", pm["factor"], 1665 / 1184, "Proceso de manuales: 1.665 / 1.184 por día"),
    ("Hojas de 4,75 mm por año", cons["MP-12"]["v35"], 0.7 * 240, "Proceso de carros: 0,7 hojas por día × 240"),
    ("Recipientes de carros por día", picos["RODANTES"]["unidades"] / 20, 20, "Proceso de carros: 20 por día en el pico "
     "(supone promedio = 70 % del pico, P09)"),
    ("Polvo ABC 90 de carros, kg por día", dia("MP-33B", "RODANTES"), 995, "Proceso de carros, tabla 3 (P06)"),
    ("Recargas de BC por mes", lineas["RECARGAS"]["demanda35"][7] / 12, 13, "Proceso de recargas BC (R07)"),
]
validacion = [dict(dato=a, mrp=r(b, 2), doc=r(c, 2), dif=r(100 * (b - c) / c, 1), fuente=d) for a, b, c, d in validacion]

# ------------------------------------------------------------------ gráfico del nivel de stock (MP-03)
def sierra(x, ruta):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    dp, lt = x["dprom"], (x["ltmin"] + x["ltmax"]) / 2
    q = x["pedido"] * x["cont"]
    t, s, ts, ss_ = 0.0, x["smax"], [0.0], [x["smax"]]
    pedidos, llegadas = [], []
    while t < 2.6 * q / dp:
        t_pp = t + (s - x["pp"]) / dp
        pedidos.append(t_pp)
        t_ll = t_pp + lt
        s_ll = x["pp"] - dp * lt
        ts += [t_pp, t_ll, t_ll]
        ss_ += [x["pp"], s_ll, s_ll + q]
        llegadas.append(t_ll)
        t, s = t_ll, s_ll + q
    ts, ss_ = ts[:-1], ss_[:-1]          # termina en la llegada, sin el salto final
    coma = lambda v, d=1: f"{v:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    fig, ax = plt.subplots(figsize=(8.2, 3.4), dpi=150)
    ax.plot(ts, ss_, color="black", lw=1.6)
    for y, txt, ls in ((x["smax"], f"Stock máximo {coma(x['smax'], 0)}", "-"),
                       (x["pp"], f"Punto de pedido {coma(x['pp'], 0)}", "--"),
                       (x["ss"], f"Stock de seguridad {coma(x['ss'], 0)}", ":")):
        ax.axhline(y, color="#7F7F7F", lw=0.9, ls=ls)
        ax.text(ts[-1] * 1.005, y, txt, va="center", fontsize=8)
    tp, tl = pedidos[0], llegadas[0]
    ax.annotate("", xy=(tl, x["ss"] * 0.45), xytext=(tp, x["ss"] * 0.45), arrowprops=dict(arrowstyle="<->", lw=0.9))
    ax.text((tp + tl) / 2, x["ss"] * 0.38, f"plazo de entrega: {x['ltmin']} a {x['ltmax']} días", ha="center",
            va="top", fontsize=7.5)
    ax.text(tl + 0.6, (x["pp"] + x["ss"]) / 2, f"se consumen {coma(x['enlt'], 0)}\nmientras llega", fontsize=7,
            va="center", color="#404040")
    ax.annotate("se pide", xy=(tp, x["pp"]), xytext=(tp - 6, x["pp"] + 0.18 * x["smax"]), fontsize=7.5,
                arrowprops=dict(arrowstyle="->", lw=0.8))
    ax.text(tl + 1.2, x["ss"] + q * 0.78, f"llega el pedido\n(+{coma(q, 0)})", fontsize=7.5)
    ax.set_xlim(0, ts[-1])
    ax.set_ylim(0, x["smax"] * 1.12)
    ax.set_xlabel("días hábiles", fontsize=8)
    ax.set_ylabel(f"stock ({x['um']})", fontsize=8)
    ax.tick_params(labelsize=7.5)
    for k in ("top", "right"):
        ax.spines[k].set_visible(False)
    ax.set_title(f"{x['codigo']} {x['material']}: se consumen {coma(x['dprom'])} por día en promedio "
                 f"({coma(x['dpico'], 0)} en el mes pico)", fontsize=8.5, loc="left")
    fig.tight_layout()
    fig.savefig(ruta)
    plt.close(fig)
    from PIL import Image
    im = Image.open(ruta)
    return dict(id="G01", titulo=f"Nivel de stock de {x['codigo']} en 2035: cuándo se pide, cuánto se pide y el colchón",
                tipo="MRP", png=os.path.join("figs", "G01.png"), w=im.width, h=im.height, medio=False, marcas=[])


G01 = sierra(next(x for x in pol if x["codigo"] == "MP-03"), os.path.join(AQUI, "figs", "G01.png"))

D = dict(demanda=demanda, perfiles=perfiles, picos=picos, lineas=lineas, consolidado=cons, politica=pol,
         almacenes=alm, validacion=validacion, p00=REFS["P00"],
         capturas=json.load(open(os.path.join(AQUI, "capturas.json"), encoding="utf-8")) + [G01])
json.dump(D, open(OUT, "w"), ensure_ascii=False, indent=1, default=str)
print("ok", OUT, sum(len(v["filas"]) for v in lineas.values()), "consumos", len(pol), "en la política")
