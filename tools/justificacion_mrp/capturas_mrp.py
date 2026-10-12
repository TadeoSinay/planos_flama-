"""Capturas marcadas para el Word «Justificación_MRP»: procesos de FLAMA, cotizaciones, planillas fuente y hojas del
propio MRP. Recorta sólo la tabla o el párrafo citado y marca en rojo el dato que se toma.

Uso: python capturas_mrp.py <carpeta de fuentes> <FLAMA_MRP.xlsx recalculado> [IDs]
  <carpeta de fuentes> tiene: carros.pdf, manuales.pdf, recargas.pdf y wpp.pdf (los procesos de carros, manuales y
  recargas y las cotizaciones por WhatsApp pasados a PDF con LibreOffice), Proyecciones.xlsx,
  Dimensionamiento_TecnicoV2.xlsx, Dimensionamiento_Tecnico_FLAMA.xlsx, MP_Abastecimiento.xlsx,
  Materias Primas.xlsx y la carpeta PRESUPUESTOS de la investigación de proveedores. El BOM se toma de
  salida/bom/FLAMA_BOM.xlsx.
Escribe figs/*.png y capturas.json (lo lee word_mrp.js).
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

import pymupdf
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from PIL import Image, ImageChops, ImageDraw, ImageFont

AQUI = os.path.dirname(os.path.abspath(__file__))
SRC, MRP = sys.argv[1], sys.argv[2]
OUT = os.path.join(AQUI, "figs")
Z = 2.0
ROJO = (220, 0, 0)
CARROS, MANUALES, RECARGAS, WPP = (os.path.join(SRC, x) for x in ("carros.pdf", "manuales.pdf", "recargas.pdf",
                                                                    "wpp.pdf"))
BOM = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(MRP))), "bom", "FLAMA_BOM.xlsx")
PRES = os.path.join(SRC, "PRESUPUESTOS", "Cotizaciones recibidas")


def _fnt(t=22):
    try:
        return ImageFont.truetype("DejaVuSans-Bold.ttf", t)
    except OSError:
        return ImageFont.load_default()


def _marcar(im, cajas):
    d = ImageDraw.Draw(im)
    for n, (x0, y0, x1, y1) in cajas:
        d.rectangle((x0, y0, x1, y1), outline=ROJO, width=4)
        cx, cy = max(x0 - 2, 17), max(y0 - 2, 17)
        d.ellipse((cx - 16, cy - 16, cx + 16, cy + 16), fill=ROJO)
        d.text((cx, cy), str(n), fill=(255, 255, 255), font=_fnt(), anchor="mm")
    return im


def _buscar(pg, texto, desde=0.0):
    rs = [r for r in pg.search_for(texto) if r.y0 >= desde - 1]
    return rs[0] if rs else None


def _xs(bls, y0, y1, pg):
    bl = [x for x in bls if x[3] > y0 and x[1] < y1]
    if not bl:
        return 0, pg.rect.x1
    return max(min(x[0] for x in bl) - 8, 0), min(max(x[2] for x in bl) + 8, pg.rect.x1)


def zona(pdf, inicio, fin=None, n=0, marcas=(), alto=None, pag=None):
    """Recorte desde el párrafo que empieza con «inicio» (n-ésima aparición en el documento) hasta el que termina con
    «fin» (aunque esté en la página siguiente) o hasta «alto» puntos. marcas: [(número, texto[, ocurrencia])]."""
    doc = pymupdf.open(pdf)
    k = 0
    for i, pg in enumerate(doc):
        if pag and i + 1 != pag:
            continue
        for r in pg.search_for(inicio):
            if k < n:
                k += 1
                continue
            y0 = r.y0 - 6
            partes = []
            j, ytop = i, y0
            while True:
                pj = doc[j]
                bls = [x for x in pj.get_text("blocks") if x[1] < pj.rect.y1 - 40]
                if fin:
                    rf = _buscar(pj, fin, ytop if j == i else 0)
                    y1 = rf.y1 if rf else max(x[3] for x in bls) + 4
                else:
                    rf, y1 = True, min(ytop + alto, pj.rect.y1 - 30)
                y1 = max([y1] + [x[3] for x in bls if x[1] < y1 - 2 and x[3] > y1 and x[3] - x[1] < 220]) + 5
                if j > i:
                    ytop = min(x[1] for x in bls) - 6
                x0, x1 = _xs(bls, ytop, y1, pj)
                partes.append(dict(pdf=pdf, pag=j + 1, rect=(x0, ytop, x1, y1), marcas=[]))
                if rf:
                    break
                j += 1
                assert j < len(doc) and j - i < 3, f"no está «{fin}»"
            partes[0]["marcas"] = list(marcas)
            return partes
    raise AssertionError(f"no está «{inicio}» en {os.path.basename(pdf)}")


def render(listas):
    partes = [p for x in listas for p in (x if isinstance(x, list) else [x])]
    imgs, cajas, dy = [], [], 0
    pend = [m for p in partes for m in p["marcas"]]
    for p in partes:
        pg = pymupdf.open(p["pdf"])[p["pag"] - 1]
        clip = pymupdf.Rect(*p["rect"])
        pix = pg.get_pixmap(matrix=pymupdf.Matrix(Z, Z), clip=clip)
        imgs.append(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
        for m in list(pend):
            num, q = m[0], m[1]
            k = m[2] if len(m) > 2 else 0
            rs = [r for r in pg.search_for(q) if clip.intersects(r)]
            if len(rs) > k:
                r = rs[k]
                cajas.append((num, ((r.x0 - clip.x0) * Z - 5, (r.y0 - clip.y0) * Z + dy - 4,
                                    (r.x1 - clip.x0) * Z + 5, (r.y1 - clip.y0) * Z + dy + 4)))
                pend.remove(m)
        dy += pix.height + 14
    assert not pend, f"marcas sin encontrar: {pend}"
    im = Image.new("RGB", (max(i.width for i in imgs), dy - 14), (255, 255, 255))
    y = 0
    for i in imgs:
        im.paste(i, (0, y))
        y += i.height + 14
    return _marcar(im, cajas)


def recortar_blanco(im, margen=12):
    fondo = Image.new(im.mode, im.size, (255, 255, 255))
    bb = ImageChops.difference(im, fondo).getbbox()
    if not bb:
        return im
    return im.crop((max(bb[0] - margen, 0), max(bb[1] - margen, 0), min(bb[2] + margen, im.width),
                    min(bb[3] + margen, im.height)))


def planilla(ruta, hoja, rango, marcas=(), cols=None, data_only=True, ancho_max=None, filas=None, anchos=None):
    """Captura de un rango de una planilla: se copia (valores, formatos, negrita, relleno, anchos) a un libro
    temporal, se imprime con LibreOffice y se recorta. marcas: [(número, celda)] con borde rojo y número."""
    wb = load_workbook(ruta, data_only=data_only)
    ws = wb[hoja]
    c0, r0, c1, r1 = range_boundaries(rango)
    cols = cols or list(range(c0, c1 + 1))
    filas = filas or list(range(r0, r1 + 1))
    out = Workbook()
    wo = out.active
    fino = Side(style="thin", color="A6A6A6")
    rojo = Side(style="thick", color="DC0000")
    mapa = {}
    for i, r in enumerate(filas, 1):
        for j, c in enumerate(cols, 1):
            s = ws.cell(r, c)
            d = wo.cell(i, j, s.value)
            mapa[s.coordinate] = d.coordinate
            d.number_format = s.number_format
            f = s.font
            d.font = Font(name="Arial", size=9, bold=bool(f and f.b), color=f.color.rgb if f and f.color and
                          isinstance(f.color.rgb, str) and f.color.rgb not in ("FF000000",) else None)
            if s.fill and s.fill.fgColor and isinstance(s.fill.fgColor.rgb, str) and s.fill.fill_type == "solid":
                d.fill = PatternFill("solid", fgColor=s.fill.fgColor.rgb)
            d.alignment = Alignment(wrap_text=True, vertical="top")
            d.border = Border(left=fino, right=fino, top=fino, bottom=fino)
    for j, c in enumerate(cols, 1):
        w = ws.column_dimensions[get_column_letter(c)].width or 10
        wo.column_dimensions[get_column_letter(j)].width = anchos[j - 1] if anchos else (min(w, ancho_max) if ancho_max
                                                                                          else w)
    for k, cel in marcas:
        d = wo[mapa[cel]]
        d.border = Border(left=rojo, right=rojo, top=rojo, bottom=rojo)
        d.value = f"({k}) {d.value}" if not isinstance(d.value, (int, float)) else d.value
    wo.page_setup.orientation = "landscape"
    wo.page_setup.fitToWidth, wo.page_setup.fitToHeight = 1, 0
    wo.sheet_properties.pageSetUpPr.fitToPage = True
    wo.page_margins.left = wo.page_margins.right = wo.page_margins.top = wo.page_margins.bottom = 0.2
    with tempfile.TemporaryDirectory() as d:
        x = os.path.join(d, "c.xlsx")
        out.save(x)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", d, x], check=True,
                       capture_output=True, timeout=300)
        doc = pymupdf.open(os.path.join(d, "c.pdf"))
        ims = []
        for pg in doc:
            pix = pg.get_pixmap(matrix=pymupdf.Matrix(Z, Z))
            ims.append(recortar_blanco(Image.frombytes("RGB", (pix.width, pix.height), pix.samples)))
    im = Image.new("RGB", (max(i.width for i in ims), sum(i.height for i in ims)), (255, 255, 255))
    y = 0
    for i in ims:
        im.paste(i, (0, y))
        y += i.height
    return im


def imagen(ruta):
    return Image.open(ruta).convert("RGB")


def pdf_q(*partes):
    return os.path.join(PRES, *partes)


CAPS = []


def cap(fid, titulo, tipo, fn, marcas=(), medio=False):
    CAPS.append(dict(id=fid, titulo=titulo, tipo=tipo, fn=fn, marcas=list(marcas), medio=medio))


# ------------------------------------------------------------------ procesos (documentos de FLAMA)
cap("P01", "Proceso de manuales, recepción: formatos y presentación de la materia prima y consumos del pico", "Proceso",
    lambda: render([zona(MANUALES, "En cuanto a las hojas de chapa", "con cañería a los 13 puestos",
                         marcas=[(1, "1220 × 2440 × 1,25 mm"), (2, "paquete de 80 hojas"), (3, "50 y 20 hojas"),
                                 (4, "atados de 16"), (5, "97 tubos por día"), (6, "1,7 t por día en el pico"),
                                 (7, "48 m³ por día en el pico")])]),
    ["Formatos LAF del 2,5, 5 y 10 kg", "Paquete de 80 hojas (5 kg)", "Paquetes de 50 y 20 hojas (2,5 y 10 kg)",
     "Atados de 16 tubos", "97 tubos por día: no coincide con el corte de tubo (53, P03)",
     "Polvo de manuales: 1,7 t por día en el pico", "Gas de soldadura: 48 m³ por día en el pico"])
cap("P02", "Proceso de manuales, corte de chapa: cuerpos por hoja y aprovechamiento", "Proceso",
    lambda: render([zona(MANUALES, "Pieza (2ª fase)", "Controles internos", marcas=[(1, "27"), (2, "12", 0),
                                                                                      (3, "15", 1)])]),
    ["2,5 kg: 27 cuerpos por hoja", "5 kg: 12 por hoja", "10 kg: 15 por hoja"], medio=True)
cap("P03", "Proceso de manuales, corte de tubo: 23 tramos de 255 mm por caño", "Proceso",
    lambda: render([zona(MANUALES, "Actividad por caño", "Tiempo por día", marcas=[(1, "Cortar 23"), (2, "53")])]),
    ["23 tramos por caño de 6 m", "53 caños por día en el pico"], medio=True)
cap("P04", "Proceso de manuales: granalla por cilindro, salida de la pintura y nitrógeno", "Proceso",
    lambda: render([zona(MANUALES, "Entra también la granalla", alto=24, marcas=[(1, "0,08 kg por cilindro")]),
                    zona(MANUALES, "La línea de pintura en polvo Firesafer", alto=34,
                         marcas=[(2, "con recuperación de polvo")]),
                    zona(MANUALES, "Al final de la línea de pintura", alto=60,
                         marcas=[(3, "413 por día"), (4, "1.252")]),
                    zona(MANUALES, "Con la presión de servicio del plano", alto=36, marcas=[(5, "7,1 m³")])]),
    ["Granalla: 0,08 kg por cilindro", "Cabina con recuperación de polvo", "Terminados: 413 por día en el pico",
     "Cilindros vacíos: 1.252 por día", "Nitrógeno: 7,1 m³ por día en el pico"])
cap("P05", "Proceso de manuales, embalaje: una caja por recipiente y unidades por pallet", "Proceso",
    lambda: render([zona(MANUALES, "La BOM de cada plano define el embalaje", "Embalaje según la BOM",
                         marcas=[(1, "una caja de cartón corrugado por recipiente"), (2, "400"), (3, "167"),
                                 (4, "90"), (5, "40")])]),
    ["Caja por recipiente", "1 kg: 400 por pallet", "2,5 kg: 167", "5 kg: 90", "10 kg: 40"], medio=True)
cap("P06", "Proceso de carros, puesto 1, tabla 3: consumo por día, unidad de entrega y demora del proveedor",
    "Proceso",
    lambda: render([zona(CARROS, "Consumos y abastecimiento", "Tabla 3. Consumo por día",
                         marcas=[(1, "10 días háb. (Pradecon)"), (2, "60 días háb. (importado)"),
                                 (3, "28 días háb. (Eli-Met)"), (4, "7 días háb. (Polvex)"),
                                 (5, "Carga de cada carro + 1 % de merma")]),
                    zona(CARROS, "Los demás insumos se consumen", alto=46,
                         marcas=[(6, "20 días hábiles"), (7, "15; estructura")])]),
    ["Chapa Pradecon: 10 días hábiles", "Casquetes importados: 60 días hábiles", "Cuplas Eli-Met: 28 días hábiles",
     "Polvo Polvex: 7 días hábiles", "Merma de carga del polvo: 1 %",
     "Válvula de carro y manómetro: 20 días hábiles", "Manguera con lanza 15, ruedas 15, alambre 5, gases 1"])
cap("P07", "Proceso de carros, puesto 1: unidades de entrega (paquetes, pallets de casquetes, cajas de cuplas)",
    "Proceso",
    lambda: render([zona(CARROS, "Hoja LAC 3,2 mm", "Tabla 1. Bultos que llegan",
                         marcas=[(1, "Paquete de 20 hojas"), (2, "Paquete de 15 hojas"), (3, "Pallet de 400 u"),
                                 (4, "Pallet de 300 u"), (5, "Pallet de 150 u"), (6, "Caja de 100 u"),
                                 (7, "Caja de 40 u")])]),
    ["LAC 3,2: paquete de 20", "LAC 4,75: paquete de 15", "Casquete 10\": 400 por pallet", "12\": 300",
     "15\": 150", "Cupla M30: caja de 100", "Cupla 2½\": caja de 40"], medio=True)
cap("P08", "Proceso de carros, puesto 2: planes de corte de la hoja de 1500 × 3000 mm", "Proceso",
    lambda: render([zona(CARROS, "1.ª fase: columnas", "Tabla . Corte por modelo",
                         marcas=[(1, "84,1 %"), (2, "84,8 %"), (3, "3 + 1"), (4, "79,2 %"), (5, "73,3 %"),
                                 (6, "72,7 %")])]),
    ["25 kg: 9 cuerpos por hoja (84,1 %)", "50 kg: 6 por hoja (84,8 %)", "Hoja mixta: 3 de 70 kg + 1 de 100 kg",
     "Aprovechamiento de la mixta: 79,2 %", "70 kg sola: 4 por hoja", "100 kg sola: 3 por hoja"])
cap("P09", "Proceso de carros, puesto 2: «El promedio se toma como el 70 % del pico»", "Proceso",
    lambda: render([zona(CARROS, "Las tiras para placas se cortan de una de cada", alto=30,
                         marcas=[(1, "El promedio se toma como el 70 % del pico")])]),
    ["Supuesto del proceso de carros (factor de pico 1 / 0,7 = 1,43)"], medio=True)
M25 = ([(1, "25 kg"), (2, "50 kg")], ["Modelo de 25 kg (gas y alambre por unidad en la fila)",
                                      "Modelo de 50 kg; siguen 70 y 100 kg"])
for i, (ini, fin, tit, mm) in enumerate([
        ("Cada punto lleva unos 3 segundos", "Tabla. Consumo de gas y alambre del punteo", "puesto 4, punteo", M25),
        ("El consumo se estimó con el tiempo de arco de cada cuerpo", "Tabla. Consumo de gas y alambre de la soldadura "
         "longitudinal", "puesto 5, soldadura longitudinal", M25),
        ("El cordón es un filete", "Tabla. Consumo de gas y alambre de la soldadura de cupla",
         "puesto 6, soldadura de cupla", ([(1, "M30 × 1,5"), (2, "BSP")], ["Cupla M30 (carros de 25 y 50 kg)",
                                                                           "Cupla 2½\" (carros de 70 y 100 kg)"])),
        ("Para que sea comparable con la longitudinal", "Tabla. Consumo de gas y alambre del armado",
         "puesto 7, soldadura circunferencial", M25),
        ("El consumo se calculó con el tiempo de arco de cada carro", "Tabla. Consumo de gas y alambre de la soldadura "
         "de accesorios", "puesto 8, soldadura de accesorios", M25)]):
    cap(f"P1{i}", f"Proceso de carros, {tit}: gas y alambre por unidad", "Proceso",
        lambda ini=ini, fin=fin, mm=mm: render([zona(CARROS, ini, fin, marcas=mm[0])]), mm[1], medio=True)
cap("P15", "Proceso de carros, 8.1: lo que se compra (caño, barra trefilada, ruedas) y lo que sale de las orillas",
    "Proceso",
    lambda: render([zona(CARROS, "En este puesto se fabrican las piezas del carro", alto=70,
                         marcas=[(1, "las piezas de chapa salen de las orillas del corte"),
                                 (2, "luego del pintado que es tercerizado")]),
                    zona(CARROS, "Material", "Tabla.Lo que llega a la fabricación de accesorios", pag=67,
                         marcas=[(3, "Barras de 6 m (6,9 kg)"), (4, "Barras de 4 m (9,0 kg)")])]),
    ["Piezas de chapa de las orillas del corte (sin compra)", "Pintura de carros tercerizada",
     "Caño en barras de 6 m", "Barra trefilada 3/4\" en barras de 4 m"])
cap("P16", "Proceso de carros, 8.1: medidas de corte de manija, caño del eje y eje", "Proceso",
    lambda: render([zona(CARROS, "Recipiente (dato)", "Ruedas (2, compradas)",
                         marcas=[(1, "Corte 1.440"), (2, "1\" × 354"), (3, "3/4\" × 490"), (4, "1\" × 510"),
                                 (5, "1\" × 590")])]),
    ["Manija 25 kg: 1.440 mm de caño 1\" (1.520 y 1.590 en los demás)", "Caño del eje 25 kg: 354 mm",
     "Eje 25 kg: 490 mm de barra 3/4\"", "Eje 50 kg: 510 mm de barra 1\"", "Eje 70 y 100 kg: 590 mm"], medio=True)
cap("P17", "Proceso de carros, 8.1: consumos de discos y mechas", "Proceso",
    lambda: render([zona(CARROS, "Como no se suelda, no hay gas ni alambre", "Tabla. Consumos de la fabricación",
                         marcas=[(1, "200 cortes por disco de sensitiva"), (2, "40 juegos de triángulos"),
                                 (3, "200 agujeros por mecha")])]),
    ["200 cortes por disco de sensitiva", "40 juegos de triángulos por disco de 115 mm", "200 agujeros por mecha"],
    medio=True)
cap("P18", "Proceso de carros, puestos 4 y 6: discos de las placas y antiproyecciones", "Proceso",
    lambda: render([zona(CARROS, "El gas 92 % Ar + 8 % CO", "Cuánto pedir y cada cuánto", n=0,
                         marcas=[(1, "un disco de corte por tanda"), (2, "un disco flap cada 4 tandas")]),
                    zona(CARROS, "Antes de soldar cada cúpula se le da un disparo", alto=24,
                         marcas=[(3, "unos 1 ml")])]),
    ["Un disco de corte por tanda de placas (7 por mes en el pico)", "Un disco flap cada 4 tandas",
     "Antiproyecciones: 1 ml por cúpula"], medio=True)

# ------------------------------------------------------------------ cotizaciones
cap("Q01", "Cotización Pradecon (mail del 28/09/2026): plazo de entrega y peso de bobina de fleje", "Cotización",
    lambda: render([zona(pdf_q("Materia Prima", "Chapa de acero", "Bobina flejada", "Pradecon",
                               "cotizacion flejes PRADECON SA.pdf"), "El peso promedio por bobina", "proveedor.",
                         marcas=[(1, "1.8 y 2 tn"), (2, "son 7/10 dias el plazo de entrega")])]),
    ["Bobina de fleje de 1,8 a 2 t", "Plazo de entrega: 7 a 10 días"], medio=True)
cap("Q02", "Cotización Pacheco Chapas PR-42021: flejes en rollos (kg y cantidad de rollos)", "Cotización",
    lambda: render([zona(pdf_q("Materia Prima", "Chapa de acero", "Bobina flejada", "Pacheco Chapas",
                               "FLAMA SA -- COTIZACION N#U00b042021.pdf"), "Item", "OBSERVACIONES",
                         marcas=[(1, "6 ROLLOS"), (2, "4 ROLLOS")])]),
    ["1,25 × 200: 2.951 kg en 6 rollos (492 kg por rollo)", "1,6 × 250: 2.459 kg en 4 rollos (615 kg)"])
cap("Q03", "Cotizaciones de MP por WhatsApp: Air Liquide (plazo, baterías) y Eli-Met (plazo de las cuplas)",
    "Cotización",
    lambda: render([zona(WPP, "Gases AIR LIQUIDE", "Cupla soldable 2-1/2",
                         marcas=[(1, "a día + 1"), (2, "115,5m3"), (3, "plazo de entrega 35 a 40 días")])]),
    ["Air Liquide: entrega a día + 1", "Batería de 11 cilindros, 115,5 m³", "Eli-Met: 35 a 40 días"])
cap("Q04", "Cotización Comercial MID 27626: caño, entrega inmediata", "Cotización",
    lambda: render([zona(pdf_q("Materia Prima", "Tubo de acero con costura ", "FLAMA-COTI-27626.pdf"), "LOS PRECIOS",
                         "VALIDEZ", marcas=[(1, "ENTREGA INMEDIATA SALVO VENTAS")])]),
    ["Entrega inmediata (de stock)"], medio=True)
cap("Q07", "Lista de precios de agentes y válvulas: envases y válvulas HZ", "Cotización",
    lambda: render([zona(pdf_q("Materia Prima", "Agentes extintores", "Lista de precios Septiembre 26.pdf"),
                         "LISTA DE PRECIOS", "Valvula carro 50-70-100",
                         marcas=[(1, "Bolsa 25 kgs"), (2, "Bidon de 20 litros"), (3, "M-22 1kg c/tubo de pesca"),
                                 (4, "M-30 c/resorte tubo 7/8")])]),
    ["Polvo en bolsas de 25 kg", "AFFF en bidón de 20 L", "Válvula HZ M-22 (1 kg)", "Válvula HZ M-30 (2,5 a 10 kg)"],
    medio=True)
cap("Q09", "Cotización del servicio de pintura de carros (mensaje): precio por carro", "Cotización",
    lambda: imagen(pdf_q("Maquinaria", "Pintura electricolor", "Servicio de pintura carros", "pintura_carros.png")),
    ["USD 13,5 / 20,3 / 27,5 / 32 por carro; sin plazo de entrega"], medio=True)

# ------------------------------------------------------------------ proceso de recargas
cap("R01", "Proceso de recargas, mapa: cuándo corresponde cada operación", "Proceso",
    lambda: render([zona(RECARGAS, "En la siguiente tabla vemos", "Llega sin presión y sin polvo.",
                         marcas=[(1, "Siempre que el extintor tenga"), (2, "Cuando la PH está vencida o"),
                                 (3, "Solo si el extintor tiene óxido")])]),
    ["Descarga: siempre que tenga presión", "Prueba hidráulica: sólo si está vencida o hay daño menor",
     "Pintura: sólo si tiene óxido o pintura dañada"])
cap("R02", "Proceso de recargas, descarga: cuándo el polvo va a descarte y cuándo se recupera", "Proceso",
    lambda: render([zona(RECARGAS, "El polvo va a descarte", "una mezcla de polvos.", pag=6,
                         marcas=[(1, "mismo dígito que el año en curso"),
                                 (2, "manómetro en cero o el extintor llega accionado"),
                                 (3, "el polvo es recuperable")])]),
    ["Cambio de polvo por el año de fabricación (1 de cada 10 años)", "Sin presión o accionado: todo a descarte",
     "El resto: polvo recuperable"], medio=True)
cap("R03", "Proceso de recargas, carga de polvo ABC: polvo recuperado más polvo nuevo", "Proceso",
    lambda: render([zona(RECARGAS, "Recipiente vacío, seco por dentro", alto=40, pag=12,
                         marcas=[(1, "recuperado aprobado más nuevo certificado")]),
                    zona(RECARGAS, "Al empezar el turno, contrastar la balanza", alto=30, pag=13,
                         marcas=[(2, "reponer con polvo nuevo")])]),
    ["La tolva tiene polvo recuperado más polvo nuevo", "Se repone con polvo nuevo lo que falta"], medio=True)
cap("R04", "Proceso de recargas, ensamblaje: o-ring y marbete siempre; repuestos según el registro", "Proceso",
    lambda: render([zona(RECARGAS, "El manómetro se cambia si", "qué repuestos conviene tener en stock",
                         marcas=[(1, "El o-ring se cambia siempre"), (2, "marbete siempre se debe cambiar"),
                                 (3, "conviene tener en stock")])]),
    ["O-ring nuevo en cada recarga", "Marbete del año en cada recarga",
     "Los repuestos se cambian si están dañados; su stock sale del registro"], medio=True)
cap("R05", "Proceso de recargas, presurización: nitrógeno con la misma pinza de producción", "Proceso",
    lambda: render([zona(RECARGAS, "Pinza universal neumática Yukon", "Solo nitrógeno seco", pag=15,
                         marcas=[(1, "la misma de producción"), (2, "Solo nitrógeno seco")])]),
    ["Misma pinza Yukon que en fabricación", "Sólo nitrógeno seco"], medio=True)
cap("R06", "Proceso de recargas, precinto y etiqueta de servicio", "Proceso",
    lambda: render([zona(RECARGAS, "El precinto debe llevar", "la marca de FLAMA.",
                         marcas=[(1, "El precinto debe llevar"), (2, "vencimiento de la prueba hidráulica")])]),
    ["Precinto con la identificación de FLAMA", "La etiqueta de servicio lleva el vencimiento de la PH"],
    medio=True)
cap("R07", "Proceso de recargas BC: siempre con polvo nuevo (13 equipos por mes en la planta industrial)", "Proceso",
    lambda: render([zona(RECARGAS, "Una notación que se debe tener", "La disposición final del polvo BC",
                         marcas=[(1, "todas con polvo nuevo"), (2, "13 extintores por mes")])]),
    ["Recargas de BC con polvo nuevo", "13 por mes en la planta industrial: no se justifica recuperar"], medio=True)
cap("R08", "Proceso de recargas de agua y AFFF: el agente se cambia siempre; premezcla al 3 %", "Proceso",
    lambda: render([zona(RECARGAS, "El agente se cambia siempre", alto=34, marcas=[(1, "El agente se cambia siempre")]),
                    zona(RECARGAS, "La mezcla es al 3%", alto=20, marcas=[(2, "La mezcla es al 3%")])]),
    ["Agente nuevo en cada recarga", "AFFF: premezcla al 3 %"], medio=True)
cap("R09", "Proceso de recargas de CO₂: se recupera y se vuelve a cargar", "Proceso",
    lambda: render([zona(RECARGAS, "Solo se hace cuando corresponde PH", alto=30,
                         marcas=[(1, "Solo se hace cuando corresponde PH")]),
                    zona(RECARGAS, "Elegir de dónde se carga", alto=30, marcas=[(2, "recuperado si el extintor")])]),
    ["Sólo se vacía si va a PH o hay que abrir la válvula", "Se carga con el CO₂ recuperado"], medio=True)
cap("R10", "Proceso de recargas de halogenados: se recupera en circuito cerrado", "Proceso",
    lambda: render([zona(RECARGAS, "El agente nunca se ventea", "nitrógeno o con argón",
                         marcas=[(1, "El agente nunca se ventea"), (2, "nitrógeno o con argón")])]),
    ["El agente se recupera", "Presurización con nitrógeno o argón"], medio=True)
cap("R11", "Proceso de recargas, pintura: esmalte poliuretano, recipiente completo", "Proceso",
    lambda: render([zona(RECARGAS, "Esmalte poliuretano rojo", alto=16, marcas=[(1, "Esmalte poliuretano rojo")]),
                    zona(RECARGAS, "Se optó por esta opción", alto=34, marcas=[(2, "pintar siempre completo")])]),
    ["Esmalte poliuretano rojo", "Se pinta siempre completo"], medio=True)

# ------------------------------------------------------------------ planillas fuente
PROY = os.path.join(SRC, "Proyecciones.xlsx")
DIM = os.path.join(SRC, "Dimensionamiento_TecnicoV2.xlsx")
DIMF = os.path.join(SRC, "Dimensionamiento_Tecnico_FLAMA.xlsx")
ABAS = os.path.join(SRC, "MP_Abastecimiento.xlsx")
MPRIM = os.path.join(SRC, "Materias Primas.xlsx")


def col_de(ruta, hoja, fila, nombres):
    ws = load_workbook(ruta, read_only=True)[hoja]
    cab = [c.value for c in next(ws.iter_rows(min_row=fila, max_row=fila))]
    return [cab.index(n) + 1 for n in nombres]


def apilar(*ims):
    im = Image.new("RGB", (max(i.width for i in ims), sum(i.height for i in ims) + 14 * (len(ims) - 1)),
                   (255, 255, 255))
    y = 0
    for i in ims:
        im.paste(i, (0, y))
        y += i.height + 14
    return im


cap("F01", "Dimensionamiento Técnico FLAMA: perfiles de estacionalidad", "Planilla",
    lambda: planilla(DIMF, "DATOS-BASE", "A100:N104", [(1, "B102"), (2, "M103"), (3, "D104")],
                     anchos=[30] + [6] * 13),
    ["1 kg: pico de 12 % en enero y diciembre", "Manual comercial: pico de 11 % en diciembre",
     "Industrial sobre ruedas: 9 % en los meses más altos"])
cap("F02", "Dimensionamiento Técnico FLAMA: producción mensual de matafuegos del año 2035 (demanda × perfil)",
    "Planilla",
    lambda: planilla(DIMF, "DATOS-BASE", "A107:M117", [(1, "M117")], anchos=[22] + [7] * 12),
    ["Total del mes: de ahí sale el mes pico"])
cap("F03", "Dimensionamiento Técnico FLAMA: recargas de la planta industrial, anual / 12", "Planilla",
    lambda: planilla(DIMF, "DATOS-BASE", "A133:N153", [(1, "B135"), (2, "M135")],
                     filas=[133, 134, 135, 136, 137, 152, 153], anchos=[22] + [7] * 13),
    ["Enero = anual / 12", "Diciembre = anual / 12: sin estacionalidad"], medio=True)
cap("X01", "Proyecciones, «Indexado Dim Técnico»: demanda anual por producto (nuevos, cilindros, recargas, revendidos)",
    "Planilla",
    lambda: planilla(PROY, "Indexado Dim Técnico", "A1:K57", [(1, "A1"), (2, "A13"), (3, "A25"), (4, "A46")]),
    ["Matafuegos nuevos", "Cilindros vacíos", "Recargas (17 tipos)", "Revendidos (9 tipos)"])
cap("X02", "Proyecciones, «Indexado Dim Técnico»: distribución por zona (% de CABA)", "Planilla",
    lambda: planilla(PROY, "Indexado Dim Técnico", "A59:K93", [(1, "A63"), (2, "A65"), (2, "A66"), (2, "A67"),
                                                                (3, "A75"), (4, "A79"), (5, "A93")],
                     cols=[1, 2, 11], filas=list(range(59, 76)) + list(range(77, 94))),
    ["CABA Sur (zona de la planta industrial)", "PBA Sur 1.º, 2.º y 3.er cordón (zonas de la planta industrial)",
     "Subtotal de recargas CABA + PBA", "CABA (unidades nuevas)", "Subtotal de unidades nuevas"], medio=True)
cap("X03", "Proyecciones, «Reservas y Sustitutos»: venta de reserva, stock del pool y alta anual de sustitutos",
    "Planilla",
    lambda: planilla(PROY, "Reservas y Sustitutos", "A26:K82", [(1, "A30"), (2, "A53"), (3, "A68"), (4, "A81")],
                     filas=list(range(26, 31)) + list(range(53, 67)) + list(range(68, 83))),
    ["Venta de reserva: se suma al ABC 5 kg", "Stock del pool: se recarga cada año", "Alta anual de sustitutos",
     "Total del alta (kit de sustituto)"])
cap("X04", "Proyecciones, hoja «PROYECCION SALUD»: columnas Mat Usados y RECARGAS TOT", "Planilla",
    lambda: planilla(PROY, "PROYECCION SALUD", "A1:P8", [(1, "O1"), (2, "P1")], cols=[1, 3, 4, 5, 6, 8, 15, 16],
                     anchos=[6, 12, 8, 10, 20, 9, 12, 12]),
    ["Equipos usados (accionados)", "Recargas totales"])
cap("X11", "Dimensionamiento Técnico V2, Parametros Recargas: prueba hidráulica, descargados, vaciados y pintura",
    "Planilla",
    lambda: planilla(DIM, "Parametros Recargas", "A11:C22", [(1, "B12"), (2, "B14"), (3, "B15"), (4, "B22")],
                     filas=[11, 12, 13, 14, 15, 22], anchos=[44, 7, 90]),
    ["25 % con prueba hidráulica", "10 % de CO₂ y agente limpio llegan descargados o con pérdida",
     "35 % de CO₂ y agente limpio se vacían (25 % + 10 %)", "15 % llegan con la pintura deteriorada"])
cap("X12", "Dimensionamiento Técnico V2, Parametros Recargas: matriz de aplicación por familia", "Planilla",
    lambda: planilla(DIM, "Parametros Recargas", "A53:I79", [(1, "H62"), (2, "G70"), (3, "I74")],
                     filas=[53, 62, 70, 73, 74, 77], cols=[1, 4, 5, 6, 7, 8, 9], anchos=[40, 8, 8, 8, 8, 8, 8]),
    ["Desarme de CO₂: 35 %", "Agua, AFFF y K: recarga completa", "Agente limpio: presurización 35 %"])
cap("X13", "Dimensionamiento Técnico V2, Almacen PT: zonas de la planta industrial y pallets de producto terminado",
    "Planilla",
    lambda: apilar(planilla(DIM, "Almacen PT", "A15:C22", [(1, "B15"), (2, "B20"), (3, "B22")],
                            filas=[15, 20, 22], anchos=[46, 8, 60]),
                   planilla(DIM, "Almacen PT", "A28:B34", [(4, "B34")], anchos=[40, 14])),
    ["Un carro ocupa lo que 3 manuales", "90 matafuegos por pallet", "2 pallets por posición",
     "Las 4 zonas que atiende la planta industrial"])
cap("X14", "Planilla «Materias Primas», Lineamientos: diámetro del disco de cúpula y fondo", "Planilla",
    lambda: planilla(MPRIM, "Lineamientos PlanosFicha Tecnic", "A13:K21", [(1, "G14"), (2, "G15"), (3, "H16")],
                     anchos=[7, 7, 7, 7, 7, 9, 9, 7, 7, 7, 34]),
    ["1 kg, cúpula: disco de 95 mm", "1 kg, fondo: disco de 112 mm", "2,5 a 10 kg: disco de la planilla"])


def _abas(f0, f1):
    cols = col_de(ABAS, "MP Abastecimiento", 4, ["Código", "Materia prima / insumo", "Unidad de carga (UC)",
                                                 "Contenido por UC (unid. consumo)", "Lead time (días háb.)",
                                                 "Frecuencia de entrega (días háb.)", "Stock de seguridad (días háb.)",
                                                 "UC por posición"])
    return planilla(ABAS, "MP Abastecimiento", "A4:AK70", [(1, "O4"), (2, "P4"), (3, "U4")], cols=cols,
                    filas=[4] + list(range(f0, f1 + 1)), anchos=[7, 40, 24, 9, 8, 8, 8, 8])


cap("X07", "Planilla de MP («MP Abastecimiento»): unidad de carga, lead time, días entre pedidos, stock de seguridad y "
    "UC por posición (MP-01 a MP-36)", "Planilla", lambda: _abas(5, 39),
    ["Lead time (días hábiles)", "Días entre pedidos (T)", "Stock de seguridad (días)"])
cap("X07B", "Planilla de MP («MP Abastecimiento»): lo mismo de MP-37 a MP-65", "Planilla",
    lambda: _abas(40, 69), ["Lead time", "Días entre pedidos", "Stock de seguridad"])
cap("X08", "Planilla de MP, Geometria Acero: puente entre discos, espesor de pintura, alambre y gas por recipiente",
    "Planilla",
    lambda: planilla(ABAS, "Geometria Acero", "A13:F71", [(1, "C13"), (2, "C15"), (3, "E67"), (4, "E71")],
                     filas=[13, 15, 30, 67, 71], cols=[1, 2, 3, 4, 5, 6], anchos=[44, 7, 8, 8, 8, 8]),
    ["Puente entre discos: 3 mm", "Espesor de pintura: 70 µm (Sinteplast)", "Alambre por recipiente",
     "Gas por recipiente"])

# ------------------------------------------------------------------ BOM
def _bom(hoja, prod, nombres, hasta=None, desde=None):
    wb = load_workbook(BOM, read_only=True, data_only=True)
    ws = wb[hoja]
    filas, en = [1], False
    for i, f in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if f[0] == 0:
            en = f[1] == prod
        if en and (not desde or i >= desde) and (not hasta or not str(f[1]).startswith(hasta)):
            filas.append(i)
        elif en and hasta and str(f[1]).startswith(hasta):
            en = False
    return filas, col_de(BOM, hoja, 1, nombres)


def _bom_kit():
    filas, cols = _bom("Fabricados", "FL_MAT_ABC_5kg", ["Nivel", "Código", "Pieza", "Cant.", "UM"])
    wb = load_workbook(BOM, read_only=True, data_only=True)
    ws = wb["Fabricados"]
    k = [i for i in filas if str(ws.cell(i, 2).value).startswith(("FL_SUS", "SUS-"))]
    return planilla(BOM, "Fabricados", "A1:AZ400", [(1, f"C{k[1]}")], cols=cols, filas=[1] + k,
                    anchos=[5, 16, 44, 6, 5])


def _bom_fab():
    filas, cols = _bom("Fabricados", "FL_MAT_ABC_5kg", ["Nivel", "Código", "Pieza", "Cant.", "UM", "Superficie (dm²)"],
                       hasta="FL_SUS")
    wb = load_workbook(BOM, read_only=True, data_only=True)
    ws = wb["Fabricados"]
    m = {str(ws.cell(i, 3).value): i for i in filas}
    sup, cant = get_column_letter(cols[5]), get_column_letter(cols[3])
    return planilla(BOM, "Fabricados", "A1:AZ400", [(1, f"{sup}{m['Cuerpo']}"),
                                                    (2, f"C{m['Tarjeta de identificación AGC']}"),
                                                    (3, f"{cant}{m['Tornillo y tarugo de fijación']}")],
                    cols=cols, filas=filas, anchos=[5, 16, 40, 6, 5, 10])


def _bom_mod():
    cols = col_de(BOM, "MODELOS", 2, ["Código", "Carga objetivo", "V recipiente (L)", "Gas (L normales)",
                                      "Reposición en mantenimiento (UM)", "Prueba hidráulica cada (años)"])
    return planilla(BOM, "MODELOS", "A2:AP19", [(1, "L5"), (2, "O5")], cols=cols, anchos=[22, 9, 9, 9, 11, 9])


def _bom_rec():
    filas, cols = _bom("Recargas", "RK_ABC_5kg", ["Nivel", "Código", "Pieza", "Cant.", "UM"])
    wb = load_workbook(BOM, read_only=True, data_only=True)
    ws = wb["Recargas"]
    sel = [i for i in filas if i == 1 or ws.cell(i, 1).value in (0, 1) or "PBA" in str(ws.cell(i, 3).value) or
           "AGC" in str(ws.cell(i, 3).value)]
    return planilla(BOM, "Recargas", "A1:AZ2000", [], cols=cols, filas=sel[:14], anchos=[5, 16, 44, 6, 7])


cap("B01", "FLAMA_BOM.xlsx, MODELOS: carga objetivo, volumen interior, gas, reposición y prueba hidráulica", "BOM",
    _bom_mod, ["Carga objetivo (nominal + mitad de la tolerancia)", "Volumen interior del recipiente"])
cap("B02", "FLAMA_BOM.xlsx, Fabricados: ABC 5 kg pieza por pieza (cantidades y superficies)", "BOM", _bom_fab,
    ["Superficie de las dos caras del cuerpo (la exterior es la mitad)", "Tarjeta AGC: sólo destino CABA",
     "Dos tornillos con tarugo por soporte"])
cap("B03", "FLAMA_BOM.xlsx, Recargas: identificación por jurisdicción (ABC 5 kg)", "BOM", _bom_rec, [], medio=True)
cap("B04", "FLAMA_BOM.xlsx: kit del sustituto (ABC 5 kg)", "BOM", _bom_kit, ["Faja amarilla de sustituto"],
    medio=True)


# ------------------------------------------------------------------ hojas del MRP
def _filas(hoja, col, pred, extra=()):
    ws = load_workbook(MRP, read_only=True, data_only=True)[hoja]
    out = []
    for i, f in enumerate(ws.iter_rows(min_row=1, values_only=True), 1):
        if pred(f[col - 1] if len(f) >= col else None, f):
            out.append(i)
    return sorted(set(list(extra) + out))


def _cols_lin(hoja, n):
    """Columnas de una hoja de línea: Ref, código, material, destino, los n primeros productos, 2035 y mes pico."""
    ws = load_workbook(MRP, read_only=True)[hoja]
    cab = [c.value for c in next(ws.iter_rows(min_row=3, max_row=3))]
    return [1, 2, 3, 5] + list(range(6, 6 + n)) + [cab.index("2035") + 1, cab.index("Mes pico 2035") + 1]


cap("M03", "FLAMA_MRP.xlsx, hoja MANUALES: demanda arriba, consumo por unidad abajo y lo que hay que comprar",
    "MRP",
    lambda: planilla(MRP, "MANUALES", "A3:Z80", cols=_cols_lin("MANUALES", 8), anchos=[5, 9, 30, 5] + [7] * 8 + [8, 8],
                     filas=[3, 4, 13, 14, 15] + list(range(16, 32))), [])
cap("M07", "FLAMA_MRP.xlsx, hoja POLITICA_STOCK: consumo, plazo, cuándo pedir, cuánto pedir y lugar", "MRP",
    lambda: planilla(MRP, "POLITICA_STOCK", "A3:W200", cols=[2, 6, 7, 8, 9, 10, 12, 13, 14, 18, 20, 23],
                     anchos=[8, 9, 9, 9, 6, 6, 10, 10, 10, 8, 9, 8],
                     filas=[4] + _filas("POLITICA_STOCK", 2, lambda v, f: v in (
                         "MP-01", "MP-03", "MP-07", "MP-12", "MP-15", "MP-20", "MP-25", "MP-33", "MP-35", "MP-43"))),
    [])
cap("M08", "FLAMA_MRP.xlsx, hoja POLITICA_STOCK: lugares por almacén", "MRP",
    lambda: planilla(MRP, "POLITICA_STOCK", "A3:H200", cols=[2, 6, 7, 8], anchos=[40, 12, 12, 10],
                     filas=_filas("POLITICA_STOCK", 2, lambda v, f: v == "Almacén" or (
                         f[0] is None and isinstance(v, str) and len(f) > 5 and isinstance(f[5], (int, float))))),
    [])

if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    solo = set(sys.argv[3:])
    meta = []
    previo = {}
    js = os.path.join(AQUI, "capturas.json")
    if solo and os.path.exists(js):
        previo = {x["id"]: x for x in json.load(open(js))}
    for c in CAPS:
        if solo and c["id"] not in solo:
            if c["id"] in previo:
                meta.append(previo[c["id"]])
            continue
        im = c["fn"]()
        ruta = os.path.join(OUT, c["id"] + ".png")
        im.save(ruta, optimize=True)
        meta.append(dict(id=c["id"], titulo=c["titulo"], tipo=c["tipo"], png=os.path.join("figs", c["id"] + ".png"),
                         w=im.width, h=im.height, medio=c["medio"],
                         marcas=[dict(n=i + 1, etiqueta=m) for i, m in enumerate(c["marcas"])]))
        print(c["id"], im.size)
    json.dump(meta, open(js, "w"), ensure_ascii=False, indent=1)
