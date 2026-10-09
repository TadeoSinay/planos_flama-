"""Excel «Justificacion de medidas planos»: crudo pieza a pieza + figura del Word de donde sale cada valor."""
import collections
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SP = "/tmp/claude-0/-home-user-TP-optimizacion/15292aaf-e550-5dee-a83e-77fd9177a6ef/scratchpad"
REPO = "/home/user/planos_flama-"
SHA = "ecef27be9ab1983f1efbb3e894818ff574e64af7"
GH = f"https://github.com/TadeoSinay/planos_flama-/blob/{SHA}/flama/"
OUTDIR = f"{REPO}/salida/justificacion"
os.makedirs(OUTDIR, exist_ok=True)
DOCX = "Justificacion de medidas planos.docx"

rows = json.load(open(f"{SP}/jw/crudo.json"))
figs = json.load(open(f"{SP}/jw/figs.json"))

# ------------------------------------------------------------------ correcciones de origen (sin documento)
for r in rows:
    r["src"] = [k for k in r["src"] if k not in ("N:2533", "N:3517", "N:3533")]
    if r["dim"] == "Rango y sector verde":
        r["tipo"], r["estado"] = "Diseño FLAMA (supuesto)", "E"
        r["calc"] = ("Rango ≈ 2,5 × presión de servicio del catálogo (criterio de diseño); sello IRAM 3533 "
                     "exigido en la compra")
    if r["pieza"] == "Precinto de fábrica":
        r["calc"] = "Valor de diseño; IRAM 3517-2 9.4.13 exige el precinto con identificación, no fija la medida"

# ------------------------------------------------------------------ índice clave -> (figura, marcas)
idx = collections.defaultdict(list)
fig_by = {f["id"]: f for f in figs}
for f in figs:
    for m in f["marcas"]:
        for k in m["claves"]:
            idx[k].append((f["id"], m["n"]))
FR_FIG = {"1kg": "F01", "2.5kg": "F02", "5kg": "F03", "10kg": "F04", "25kg": "F05", "50kg": "F06"}


def ubicar(keys):
    """-> (figuras, marcas, documento, ubicación) como texto."""
    out = collections.OrderedDict()
    sin = []
    for k in keys:
        if k == "CALC":
            continue
        hits = idx.get(k)
        if not hits and k.startswith("FR:"):
            fid = FR_FIG[k.split(":")[1]]
            out.setdefault(fid, []).append("nota")
            continue
        if not hits:
            sin.append(k)
            continue
        for fid, n in hits:
            out.setdefault(fid, [])
            if n not in out[fid]:
                out[fid].append(n)
    figs_t = ", ".join(out)
    marcas_t = "; ".join(f"{fid}: " + ", ".join(str(n) for n in ns) for fid, ns in out.items())
    docs = " | ".join(dict.fromkeys(fig_by[fid]["fuente"] for fid in out))
    ubic = " | ".join(dict.fromkeys(fig_by[fid]["ubic"] for fid in out))
    return figs_t, marcas_t, docs, ubic, sin


lineas = {}


def L(archivo, patron):
    key = (archivo, patron)
    if key not in lineas:
        for i, s in enumerate(open(f"{REPO}/flama/{archivo}", encoding="utf-8"), 1):
            if patron in s:
                lineas[key] = i
                break
        else:
            raise KeyError(key)
    return lineas[key]


# ------------------------------------------------------------------ estilos
AR = "Arial"
F_T = Font(name=AR, size=14, bold=True)
F_B = Font(name=AR, size=10, bold=True)
F_N = Font(name=AR, size=10)
F_W = Font(name=AR, size=10, bold=True, color="FFFFFF")
F_L = Font(name=AR, size=10, color="0563C1", underline="single")
F_I = Font(name=AR, size=9, italic=True)
CAB = PatternFill("solid", fgColor="1F3864")
AMA = PatternFill("solid", fgColor="FFF2CC")
NAR = PatternFill("solid", fgColor="F8CBAD")
GRI = PatternFill("solid", fgColor="F2F2F2")
TH = Side(style="thin", color="BFBFBF")
BOR = Border(left=TH, right=TH, top=TH, bottom=TH)
WR = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
# ------------------------------------------------------------------ Leyenda
ws = wb.active
ws.title = "Leyenda"
ws["A1"] = "Justificación de medidas planos — FLAMA S.A."
ws["A1"].font = F_T
txt = [
    "Este libro tiene los valores CRUDOS de los planos: cada dimensión de cada pieza de los 17 planos FL_MAT "
    "(incluye las piezas del recipiente FL_REC), tal como salen del generador de planos.",
    f"De dónde sale cada valor está en el Word «{DOCX}»: las columnas «Fig. Word» y «Marca» indican la figura "
    "(captura del documento fuente) y el número rojo marcado sobre el dato tomado.",
    "«Cómo se obtuvo» explica la cuenta cuando el valor no se lee directo del documento; «Código» abre la línea "
    "exacta del generador en GitHub (commit fijo, no cambia).",
]
for i, t in enumerate(txt, 3):
    ws.cell(i, 1, t).font = F_N
    ws.cell(i, 1).alignment = WR
    ws.row_dimensions[i].height = 30
ws.column_dimensions["A"].width = 120
r = 8
ws.cell(r, 1, "Hojas").font = F_B
for i, (h, d) in enumerate([("Crudo", "Una fila por dimensión: modelo, plano, posición, pieza, valor, origen, figura y marca del Word, "
                                     "cálculo y link al código."),
                            ("Fuentes", "Las 47 figuras del Word: documento, página o celda, link y qué marca cada número."),
                            ("Resumen", "Cantidad de dimensiones por modelo y por tipo de origen (fórmulas sobre «Crudo»).")], r + 1):
    ws.cell(i, 1, f"{h}: {d}").font = F_N
r = 13
ws.cell(r, 1, "Colores de la columna «Estado» (mismo criterio que el BOM)").font = F_B
for i, (fill, t) in enumerate([(None, "Sin color = V: tomado de un documento (plano Fadesa, catálogo, norma, planilla MP, ficha) "
                                      "o calculado sólo con datos documentados."),
                               (AMA, "Amarillo = E: valor de diseño FLAMA o estimado, sin documento externo (ver la figura del "
                                     "código)."),
                               (NAR, "Naranja = A: a validar (geometría derivada del 70 kg, medidas relevadas en fotos de "
                                     "mercado).")], r + 1):
    c = ws.cell(i, 1, t)
    c.font = F_N
    if fill:
        c.fill = fill
ws.cell(18, 1, "Tipos de origen").font = F_B
cnt = collections.Counter(r_["tipo"] for r_ in rows)
for i, (t, n) in enumerate(sorted(cnt.items(), key=lambda x: -x[1]), 19):
    ws.cell(i, 1, f"{t}: {n} dimensiones").font = F_N
ws.cell(19 + len(cnt) + 1, 1, "Fuente de los valores: generador de planos FLAMA (repositorio planos_flama-, commit "
        f"{SHA[:7]}), mismo estado que los planos y despieces publicados.").font = F_I

# ------------------------------------------------------------------ Crudo
ws = wb.create_sheet("Crudo")
cab = ["Modelo", "Plano", "Pos", "Pieza", "Cant", "Material", "Dimensión", "Valor", "Unidad", "Origen", "Fig. Word",
       "Marca", "Documento fuente", "Ubicación en el documento", "Cómo se obtuvo", "Código (línea exacta)", "Estado"]
anchos = [22, 22, 5, 30, 5, 30, 40, 16, 7, 24, 12, 16, 50, 40, 60, 26, 7]
for j, (h, w) in enumerate(zip(cab, anchos), 1):
    c = ws.cell(1, j, h)
    c.font, c.fill, c.alignment, c.border = F_W, CAB, Alignment(wrap_text=True, vertical="center"), BOR
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[1].height = 30
uso = collections.defaultdict(list)          # figura -> filas que la usan (para el Word)
faltan = []
for i, r_ in enumerate(rows, 2):
    figs_t, marcas_t, docs, ubic, sin = ubicar(r_["src"])
    if sin:
        faltan.append((r_["modelo"], r_["pieza"], r_["dim"], sin))
    if not figs_t:
        # sin documento externo: la fuente es el código (y la figura del código si existe)
        figs_t, marcas_t, docs, ubic = "—", "—", "Generador de planos FLAMA (ver «Código»)", "—"
    arch, pat = r_["code"]
    ln = L(arch, pat)
    try:
        val = float(r_["valor"]) if not isinstance(r_["valor"], str) else r_["valor"]
    except (TypeError, ValueError):
        val = r_["valor"]
    plano = r_["modelo"] if r_["pos"] in ("0",) else (r_["modelo"] + " / " + r_["modelo"].replace("FL_MAT_", "FL_REC_")
                                                      if r_["modelo"] in ("FL_MAT_ABC_1kg", "FL_MAT_ABC_2.5kg", "FL_MAT_ABC_5kg",
                                                                          "FL_MAT_ABC_10kg", "FL_MAT_ABC_25kg", "FL_MAT_ABC_50kg",
                                                                          "FL_MAT_ABC_70kg", "FL_MAT_ABC_100kg")
                                                      and r_["pieza"] in ("Recipiente", "Cuerpo (virola)", "Cúpula", "Fondo",
                                                                          "Cuello roscado", "Varilla de refuerzo interior")
                                                      else r_["modelo"])
    vals = [r_["modelo"], plano, r_["pos"], r_["pieza"], r_["cant"], r_["material"], r_["dim"], val, r_["unidad"],
            r_["tipo"], figs_t, marcas_t, docs, ubic, r_["calc"], f"{arch} línea {ln}", r_["estado"]]
    for j, v in enumerate(vals, 1):
        c = ws.cell(i, j, v)
        c.font, c.border = F_N, BOR
        c.alignment = Alignment(wrap_text=j in (4, 6, 7, 13, 14, 15), vertical="top")
    lk = ws.cell(i, 16)
    lk.hyperlink = f"{GH}{arch}#L{ln}"
    lk.font = F_L
    fill = {"E": AMA, "A": NAR}.get(r_["estado"])
    if fill:
        for j in range(1, len(cab) + 1):
            ws.cell(i, j).fill = fill
    for fid in [x.strip() for x in figs_t.split(",") if x.strip() not in ("—", "")]:
        uso[fid].append(dict(modelo=r_["modelo"], pieza=r_["pieza"], dim=r_["dim"], valor=r_["valor"], unidad=r_["unidad"],
                             fila=i, marcas=marcas_t))
n_crudo = len(rows) + 1
ws.freeze_panes = "H2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{n_crudo}"

# ------------------------------------------------------------------ Fuentes
ws = wb.create_sheet("Fuentes")
cab = ["Fig. Word", "Título", "Tipo de fuente", "Documento", "Ubicación", "Link", "Marcas (n.º: qué se tomó)", "Nota",
       "Dimensiones del Excel que la usan"]
anchos = [10, 50, 22, 50, 40, 30, 80, 60, 14]
for j, (h, w) in enumerate(zip(cab, anchos), 1):
    c = ws.cell(1, j, h)
    c.font, c.fill, c.alignment, c.border = F_W, CAB, Alignment(wrap_text=True, vertical="center"), BOR
    ws.column_dimensions[get_column_letter(j)].width = w
for i, f in enumerate(figs, 2):
    marcas = "\n".join(f"{m['n']}: {m['etiqueta']}" for m in f["marcas"])
    vals = [f["id"], f["titulo"], f["tipo"], f["fuente"], f["ubic"], f["link"] or "—", marcas, f["nota"] or ""]
    for j, v in enumerate(vals, 1):
        c = ws.cell(i, j, v)
        c.font, c.border, c.alignment = F_N, BOR, WR
    if f["link"]:
        ws.cell(i, 6).hyperlink = f["link"]
        ws.cell(i, 6).font = F_L
        ws.cell(i, 6).value = "Abrir en GitHub"
    c = ws.cell(i, 9, f'=COUNTIF(Crudo!$K$2:$K${n_crudo},"*{f["id"]}*")')
    c.font, c.border = F_N, BOR
ws.freeze_panes = "B2"

# ------------------------------------------------------------------ Resumen (fórmulas)
ws = wb.create_sheet("Resumen")
ws["A1"] = "Dimensiones por modelo y tipo de origen (cuenta sobre la hoja «Crudo»)"
ws["A1"].font = F_T
tipos = sorted(cnt, key=lambda t: -cnt[t])
modelos = list(dict.fromkeys(r_["modelo"] for r_ in rows))
ws.cell(3, 1, "Modelo").font = F_W
ws.cell(3, 1).fill = CAB
for j, t in enumerate(tipos, 2):
    c = ws.cell(3, j, t)
    c.font, c.fill, c.alignment = F_W, CAB, Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[get_column_letter(j)].width = 15
jt = len(tipos) + 2
for j, t in enumerate(["Total", "% con documento (V)", "Amarillo (E)", "Naranja (A)"], jt):
    c = ws.cell(3, j, t)
    c.font, c.fill, c.alignment = F_W, CAB, Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[get_column_letter(j)].width = 14
ws.row_dimensions[3].height = 45
ws.column_dimensions["A"].width = 26
R = f"Crudo!$A$2:$A${n_crudo}"
T = f"Crudo!$J$2:$J${n_crudo}"
E = f"Crudo!$Q$2:$Q${n_crudo}"
for i, mo in enumerate(modelos, 4):
    ws.cell(i, 1, mo).font = F_N
    for j in range(2, jt):
        cl = get_column_letter(j)
        ws.cell(i, j, f"=COUNTIFS({R},$A{i},{T},{cl}$3)").font = F_N
    ws.cell(i, jt, f"=SUM(B{i}:{get_column_letter(jt - 1)}{i})").font = F_B
    c = ws.cell(i, jt + 1, f'=IF({get_column_letter(jt)}{i}=0,0,COUNTIFS({R},$A{i},{E},"V")/{get_column_letter(jt)}{i})')
    c.font, c.number_format = F_N, "0%"
    ws.cell(i, jt + 2, f'=COUNTIFS({R},$A{i},{E},"E")').font = F_N
    ws.cell(i, jt + 3, f'=COUNTIFS({R},$A{i},{E},"A")').font = F_N
it = 4 + len(modelos)
ws.cell(it, 1, "Total").font = F_B
for j in range(2, jt + 4):
    cl = get_column_letter(j)
    if j == jt + 1:
        c = ws.cell(it, j, f'=IF({get_column_letter(jt)}{it}=0,0,COUNTIF({E},"V")/{get_column_letter(jt)}{it})')
        c.number_format = "0%"
    else:
        c = ws.cell(it, j, f"=SUM({cl}4:{cl}{it - 1})")
    c.font = F_B
    c.fill = GRI

from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
out = f"{OUTDIR}/Justificacion de medidas planos.xlsx"
wb.save(out)
json.dump(uso, open(f"{SP}/jw/uso.json", "w"), ensure_ascii=False)
print(out, len(rows), "filas;", "claves sin figura:", len(faltan))
for x in faltan[:15]:
    print("  ", x)
