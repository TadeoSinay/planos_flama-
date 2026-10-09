"""Excel «Justificacion de medidas planos.xlsx» compacto: una matriz por familia (pieza × modelo: valor + fuente)."""
import json
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

SP = "/tmp/claude-0/-home-user-TP-optimizacion/15292aaf-e550-5dee-a83e-77fd9177a6ef/scratchpad"
OUT = sys.argv[1] if len(sys.argv) > 1 else "/home/user/planos_flama-/salida/justificacion/Justificacion de medidas planos.xlsx"
C = json.load(open(f"{SP}/jw/compacto.json"))

AZUL, BLANCO = "1F3864", "FFFFFF"
EST = {"V": None, "E": "FFF2CC", "A": "F8CBAD"}
fino = Side(style="thin", color="BFBFBF")
borde = Border(left=fino, right=fino, top=fino, bottom=fino)
cab_fill = PatternFill("solid", fgColor=AZUL)
sub_fill = PatternFill("solid", fgColor="D9E1F2")
grupo_fill = PatternFill("solid", fgColor="F2F2F2")
F = lambda **k: Font(name="Arial", size=k.pop("size", 9), **k)  # noqa: E731
wrap = Alignment(wrap_text=True, vertical="top")
centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()

# ------------------------------------------------------------------ Cómo leer
ws = wb.active
ws.title = "Cómo leer"
ws["A1"] = "Justificación de medidas planos — FLAMA S.A."
ws["A1"].font = F(size=14, bold=True, color=AZUL)
lineas = [
    "Valores crudos de cada pieza de los 17 planos FL_MAT (las piezas del recipiente son las de los planos FL_REC).",
    "Hojas «Manuales ABC», «Rodantes ABC» y «Revendidos»: una fila por pieza y dimensión; por modelo, el valor del plano y su fuente.",
    "Fuente «F10·3» = figura F10 del Word, marca roja n° 3. «D07» = valor de diseño FLAMA (hoja «Fuentes», con link a la línea del código).",
    "«cálculo» = sale de otros valores de la misma fila o de la pieza (la cuenta está en la columna «Criterio»; si cambia por modelo, en el comentario de la celda).",
    "Hacé clic en la fuente para ir a su fila en «Fuentes». El Word «Justificacion de medidas planos.docx» tiene la captura con la marca.",
]
for i, t in enumerate(lineas, 3):
    ws.cell(i, 1, "• " + t).font = F(size=10)
r = 3 + len(lineas) + 1
ws.cell(r, 1, "Colores (mismo criterio que el BOM)").font = F(size=10, bold=True)
for j, (k, txt) in enumerate((("V", "Sin color: tomado de un documento (plano Fadesa, catálogo, norma, planilla MP) o calculado sólo con esos datos"),
                              ("E", "Amarillo: valor de diseño FLAMA, sin documento (confirmar con el proveedor o el prototipo)"),
                              ("A", "Naranja: a validar (derivado o relevado en fotos)")), r + 1):
    c = ws.cell(j, 1, txt)
    c.font = F(size=10)
    if EST[k]:
        c.fill = PatternFill("solid", fgColor=EST[k])
n, e = C["n"], C["estados"]
r += 5
ws.cell(r, 1, f"Total: {n} dimensiones — con documento {e.get('V', 0)} ({round(100 * e.get('V', 0) / n)} %), "
              f"diseño FLAMA {e.get('E', 0)} ({round(100 * e.get('E', 0) / n)} %), a validar {e.get('A', 0)} "
              f"({round(100 * e.get('A', 0) / n)} %).").font = F(size=10, bold=True)
r += 2
ws.cell(r, 1, "Lo que falta validar (naranja)").font = F(size=11, bold=True, color=AZUL)
r += 1
for j, h in enumerate(("Pieza: dimensión", "Modelos", "Criterio usado / qué confirmar")):
    c = ws.cell(r, 1 + j, h)
    c.font, c.fill, c.alignment, c.border = F(bold=True, color=BLANCO), cab_fill, centro, borde
for p in C["pendientes"]:
    r += 1
    for j, v in enumerate((p["que"], p["modelos"], p["criterio"])):
        c = ws.cell(r, 1 + j, v)
        c.font, c.alignment, c.border = F(), wrap, borde
        c.fill = PatternFill("solid", fgColor=EST["A"])
ws.column_dimensions["A"].width = 60
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 70
ws.sheet_view.showGridLines = False
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ------------------------------------------------------------------ Fuentes (se arma primero el índice de filas)
fila_fuente = {}
FR0 = 4
rr = FR0
for f in C["figs"]:
    fila_fuente[f["id"]] = rr
    rr += 1
rr += 2
for d in C["disenos"]:
    fila_fuente[d["id"]] = rr
    rr += 1

# ------------------------------------------------------------------ matrices por familia
for h in C["hojas"]:
    ws = wb.create_sheet(h["nombre"])
    nm = len(h["modelos"])
    ult = 3 + 2 * nm + 1  # columna Criterio
    ws.cell(1, 1, f"{h['nombre']} — valor de cada dimensión en el plano y su fuente").font = F(size=12, bold=True, color=AZUL)
    for j, t in enumerate(("Pieza", "Dimensión", "Unidad")):
        ws.merge_cells(start_row=2, start_column=1 + j, end_row=3, end_column=1 + j)
        c = ws.cell(2, 1 + j, t)
        c.font, c.fill, c.alignment = F(bold=True, color=BLANCO), cab_fill, centro
    for i, mod in enumerate(h["modelos"]):
        col = 4 + 2 * i
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 1)
        c = ws.cell(2, col, mod)
        c.font, c.fill, c.alignment = F(bold=True, color=BLANCO), cab_fill, centro
        for k, t in enumerate(("Valor", "Fuente")):
            c = ws.cell(3, col + k, t)
            c.font, c.fill, c.alignment, c.border = F(bold=True), sub_fill, centro, borde
    ws.merge_cells(start_row=2, start_column=ult, end_row=3, end_column=ult)
    c = ws.cell(2, ult, "Criterio (cómo se obtiene)")
    c.font, c.fill, c.alignment = F(bold=True, color=BLANCO), cab_fill, centro
    for cc in range(1, ult + 1):
        ws.cell(2, cc).border = borde
        ws.cell(3, cc).border = borde
    r = 4
    previa = None
    for fl in h["filas"]:
        nueva = fl["pieza"] != previa
        a = ws.cell(r, 1, fl["pieza"] if nueva else "")
        a.font = F(bold=True)
        ws.cell(r, 2, fl["dim"]).font = F()
        ws.cell(r, 3, "" if fl["unidad"] == "-" else fl["unidad"]).font = F()
        for i, cel in enumerate(fl["celdas"]):
            col = 4 + 2 * i
            v, s = ws.cell(r, col), ws.cell(r, col + 1)
            if cel:
                v.value = cel["v"]
                s.value = cel["f"]
                if isinstance(cel["v"], float):
                    v.number_format = "0.0##"
                if EST[cel["e"]]:
                    v.fill = s.fill = PatternFill("solid", fgColor=EST[cel["e"]])
                if cel["dest"] in fila_fuente:
                    s.hyperlink = f"#'Fuentes'!A{fila_fuente[cel['dest']]}"
                    s.font = F(color="1F4E79", underline="single", size=8)
                else:
                    s.font = F(size=8)
                if fl["varia"]:
                    v.comment = Comment(cel["calc"], "FLAMA", width=320, height=90)
            else:
                v.value = "—"
                v.font = F(color="A6A6A6")
            v.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            s.alignment = Alignment(vertical="top", wrap_text=True)
        k = ws.cell(r, ult, fl["criterio"])
        k.font = F(size=8, color="404040")
        k.alignment = wrap
        for cc in range(1, ult + 1):
            x = ws.cell(r, cc)
            x.border = Border(left=fino, right=fino, bottom=fino, top=Side(style="thin", color="595959") if nueva else fino)
            if cc <= 3 and nueva:
                x.fill = grupo_fill
            if cc in (1, 2, ult):
                x.alignment = wrap
        previa = fl["pieza"]
        r += 1
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 6
    for i in range(nm):
        ws.column_dimensions[L(4 + 2 * i)].width = 11
        ws.column_dimensions[L(5 + 2 * i)].width = 10 if nm <= 4 else 8.5
    ws.column_dimensions[L(ult)].width = 70
    ws.freeze_panes = "D4"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f"A3:{L(ult)}{r - 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "2:3"

# ------------------------------------------------------------------ Fuentes
ws = wb.create_sheet("Fuentes")
ws["A1"] = "Fuentes — capturas del Word (F) y valores de diseño FLAMA (D)"
ws["A1"].font = F(size=12, bold=True, color=AZUL)
cab = ("Id", "Documento", "Dónde", "Datos tomados (n° de marca roja: valor)", "Link")
for j, t in enumerate(cab):
    c = ws.cell(3, 1 + j, t)
    c.font, c.fill, c.alignment, c.border = F(bold=True, color=BLANCO), cab_fill, centro, borde
for f in C["figs"]:
    r = fila_fuente[f["id"]]
    marcas = "; ".join(f"{m['n']}: {m['etiqueta']}" for m in f["marcas"])
    vals = (f["id"], f["titulo"], f["ubic"], marcas + (f"  — Nota: {f['nota']}" if f.get("nota") else ""), f["link"] or "")
    for j, v in enumerate(vals):
        c = ws.cell(r, 1 + j, v)
        c.font, c.alignment, c.border = F(bold=j == 0), wrap, borde
    if f["link"]:
        ws.cell(r, 5).hyperlink = f["link"]
        ws.cell(r, 5).font = F(color="1F4E79", underline="single", size=8)
r0 = fila_fuente[C["disenos"][0]["id"]] - 1
for j, t in enumerate(("Id", "Código del generador", "Qué define", "Criterio", "Modelos")):
    c = ws.cell(r0, 1 + j, t)
    c.font, c.fill, c.alignment, c.border = F(bold=True, color=BLANCO), cab_fill, centro, borde
ws.cell(r0 - 1, 1, f"Valores de diseño FLAMA (amarillo/naranja en las matrices). Links al commit {C['sha'][:7]}.").font = F(size=10, bold=True, color=AZUL)
for d in C["disenos"]:
    r = fila_fuente[d["id"]]
    vals = (d["id"], f"flama/{d['archivo']} línea {d['linea']}", d["que"], d["criterio"], d["modelos"])
    for j, v in enumerate(vals):
        c = ws.cell(r, 1 + j, v)
        c.font, c.alignment, c.border = F(bold=j == 0), wrap, borde
        c.fill = PatternFill("solid", fgColor=EST[d["estado"]])
    ws.cell(r, 2).hyperlink = d["link"]
    ws.cell(r, 2).font = F(color="1F4E79", underline="single")
for col, w in zip("ABCDE", (6, 38, 30, 90, 30)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "B4"
ws.sheet_view.showGridLines = False
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.print_title_rows = "3:3"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(OUT)
print("ok", OUT, [(s.title, s.max_row) for s in wb.worksheets])
