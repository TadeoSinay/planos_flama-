"""Capturas marcadas para «Justificación de medidas planos».

Cada figura: imagen de la fuente (página de PDF, foto o recorte) con rectángulos rojos numerados sobre
el dato que se tomó. Coordenadas de las marcas en px a 100 dpi para los PDF (salvo `unidad`), y en px de la
imagen para fotos/recortes. Salida: figs/*.png + figs.json (manifiesto que usan el Word y el Excel).
"""
import json
import os
import re

import pymupdf
from PIL import Image, ImageDraw, ImageFont

SP = "/tmp/claude-0/-home-user-TP-optimizacion/15292aaf-e550-5dee-a83e-77fd9177a6ef/scratchpad"
OUT = f"{SP}/jw/figs"
os.makedirs(OUT, exist_ok=True)
FAD = f"{SP}/fadesa/PLANOS FADESA"
CAT = f"{SP}/cat"
NOR = f"{SP}/nor/IRAM COMPLETO"
RES = f"{SP}/nor/ORDENANZAS-RESOLUCIONES"
GEO = f"{SP}/ar/6bc5703f-GEORGIA_especificaciones/GEORGIA especificaciones"
MEL = f"{SP}/ar/4d41f45c-MELISAM_especirficaciones/MELISAM especirficaciones"
REPO = "/home/user/planos_flama-"
SHA = "ecef27be9ab1983f1efbb3e894818ff574e64af7"
GH = f"https://github.com/TadeoSinay/planos_flama-/blob/{SHA}/flama/"
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
MONO = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"
RED = (220, 0, 0)
DPI = 150

figs = []


def pdf_img(path, page, dpi=DPI):
    p = pymupdf.open(path)[page]
    pix = p.get_pixmap(dpi=dpi)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def marcar(img, marks, k=1.0, crop=None, der=False):
    """marks: [(n, (x0, y0, x1, y1))] en coordenadas base; k = factor base -> img."""
    im = img.convert("RGB")
    d = ImageDraw.Draw(im)
    f = ImageFont.truetype(FONT, max(14, int(im.width / 55)))
    r = max(11, int(im.width / 90))
    usados = []
    W, H = im.size
    for n, (x0, y0, x1, y1) in marks:
        X0, Y0, X1, Y1 = x0 * k, y0 * k, x1 * k, y1 * k
        d.rectangle([X0, Y0, X1, Y1], outline=RED, width=max(3, int(im.width / 400)))
        ym = (Y0 + Y1) / 2
        cands = [(X0 - r - 4, ym), (X1 + r + 4, ym), (X0 - r - 2, Y0 - r - 2), (X1 + r + 2, Y0 - r - 2),
                 (X0 - r - 2, Y1 + r + 2), (X1 + r + 2, Y1 + r + 2), (X0 - 3 * r - 6, ym), (X1 + 3 * r + 6, ym)]
        if der:
            cands = [(X1 + r + 4, ym), (X1 + r + 4, Y0 + r), (X1 + r + 4, Y1 - r)] + cands
        elegido = None
        for cx, cy in cands:
            cx, cy = min(max(r + 1, cx), W - r - 1), min(max(r + 1, cy), H - r - 1)
            if all((cx - ux) ** 2 + (cy - uy) ** 2 >= (2 * r + 3) ** 2 for ux, uy in usados):
                elegido = (cx, cy)
                break
        cx, cy = elegido or (min(max(r + 1, cands[0][0]), W - r - 1), min(max(r + 1, cands[0][1]), H - r - 1))
        usados.append((cx, cy))
        d.ellipse([cx - r, cy - r, cx + r, cy + r], fill=RED)
        d.text((cx, cy), str(n), fill="white", font=f, anchor="mm")
    if crop:
        c = [v * k for v in crop]
        im = im.crop(c)
    return im


def limpio(x):
    return x.replace("#U00d8", "Ø").replace("#U00e1", "á").replace("te#U00eccnica", "técnica")


def fig(fid, titulo, fuente, ubic, img, marks, k=1.0, crop=None, link="", tipo="", nota="", der=False):
    """marks: [(n, caja, etiqueta, [claves])]."""
    im = marcar(img, [(m[0], m[1]) for m in marks], k, crop, der)
    titulo, fuente = limpio(titulo), limpio(fuente)
    path = f"{OUT}/{fid}.png"
    im.save(path)
    figs.append(dict(id=fid, titulo=titulo, fuente=fuente, ubic=ubic, png=path, w=im.width, h=im.height, link=link,
                     tipo=tipo, nota=nota, marcas=[dict(n=m[0], etiqueta=m[2], claves=m[3]) for m in marks]))


K = DPI / 100.0  # cajas definidas a 100 dpi

# =========================================================== planos Fadesa de recipiente
def rec(f, archivo, codigo, tam, marks, crop=None, nota=""):
    fig(f, f"Plano Fadesa «{archivo[:-4]}» (código {codigo})", f"Plano de recipiente Fadesa, archivo «{archivo}»",
        "Página 1 (única), cotas y notas del plano", pdf_img(f"{FAD}/{archivo}", 0),
        [(n, b, e, [f"FR:{tam}:{c}" for c in cl]) for n, b, e, cl in marks], K, crop, tipo="Plano Fadesa", nota=nota)


rec("F01", "1kg 3pulg R1.pdf", "G809", "1kg", [
    (1, (255, 53, 295, 72), "Ø29 — Ø exterior del cuello", ["cuello_d"]),
    (2, (163, 110, 182, 165), "12,5 ± 0,2 — altura de cúpula", ["hd"]),
    (3, (360, 350, 380, 405), "298 ± 0,5 — altura total del recipiente", ["total"]),
    (4, (163, 620, 182, 660), "22 ± 0,2 — altura del fondo", ["hf"]),
    (5, (225, 680, 322, 700), "Ø76,2 (3\") ± 0,5 — Ø exterior", ["D"]),
    (6, (493, 100, 540, 117), "M22x1,5 — rosca del cuello", ["rosca"]),
    (7, (476, 203, 524, 220), "Ø26 ± 0,5 — asiento de la válvula", ["asiento"]),
    (8, (508, 723, 640, 738), "Espesor chapa 1,25 — cuerpo y fondo", ["t", "tf"]),
    (9, (508, 740, 645, 755), "Espesor cúpula 0,9", ["td"]),
    (10, (518, 755, 635, 770), "Masa 0,76 kg — masa del recipiente", ["masa"]),
    (11, (505, 770, 650, 786), "Volumen 1,18 dm³", ["vol"]),
    (12, (536, 445, 614, 523), "Detalle 2 (1:1): bordón del caño bajo la cúpula — labio ≈ 5,5, escalón ≈ 4,5", ["det2"]),
    (13, (653, 445, 730, 523), "Detalle 3 (1:1): bordón inferior del caño, fondo por fuera", ["det3"]),
], nota="El plano no acota la altura del cuello por separado: se tomó 12,5 mm para completar la altura total de 298.")
rec("F02", "Recipiente_2.5kg_R2_sin_logo.pdf", "G731", "2.5kg", [
    (1, (245, 52, 282, 70), "Ø37 — Ø exterior del cuello", ["cuello_d"]),
    (2, (108, 62, 125, 88), "14 — altura del cuello", ["cuello_h"]),
    (3, (106, 140, 125, 198), "39,5 ± 0,3 — altura de cúpula", ["hd"]),
    (4, (424, 365, 442, 425), "313,5 ± 0,5 — altura total", ["total"]),
    (5, (92, 618, 110, 660), "16 ± 0,3 — pollera del fondo cóncavo", ["zf_borde"]),
    (6, (297, 655, 315, 682), "5,5 — flecha del fondo cóncavo", ["zf_centro"]),
    (7, (235, 722, 295, 740), "Ø124 ± 0,5 — Ø exterior", ["D"]),
    (8, (553, 467, 600, 481), "M30x1,5 — rosca del cuello", ["rosca"]),
    (9, (540, 542, 590, 556), "Ø35 ± 0,5 — asiento de la válvula", ["asiento"]),
    (10, (148, 792, 302, 804), "Espesor de chapa 1,25 (cuerpo, cúpula y fondo)", ["t", "td", "tf"]),
    (11, (165, 804, 290, 816), "Masa 1,48 kg", ["masa"]),
    (12, (170, 816, 282, 830), "Volumen 3,5 dm³", ["vol"]),
])
rec("F03", "Recipiente_5kg_R2_sin_logo.pdf", "G732", "5kg", [
    (1, (262, 42, 312, 60), "Ø37 — Ø exterior del cuello", ["cuello_d"]),
    (2, (136, 15, 155, 60), "15 ± 0,2 — altura del cuello", ["cuello_h"]),
    (3, (136, 115, 155, 175), "54,5 ± 0,3 — altura de cúpula", ["hd"]),
    (4, (446, 360, 465, 428), "389,5 ± 0,5 — altura total", ["total"]),
    (5, (112, 615, 131, 668), "21,5 ± 0,2 — pollera del fondo cóncavo", ["zf_borde"]),
    (6, (329, 635, 348, 685), "5,5 ± 0,1 — flecha del fondo cóncavo", ["zf_centro"]),
    (7, (255, 736, 335, 754), "Ø153 ± 0,5 — Ø exterior", ["D"]),
    (8, (548, 423, 602, 440), "M30x1,5 — rosca del cuello", ["rosca"]),
    (9, (548, 495, 605, 512), "Ø35 ± 0,5 — asiento de la válvula", ["asiento"]),
    (10, (183, 801, 341, 816), "Espesor de chapa 1,6 (cuerpo, cúpula y fondo)", ["t", "td", "tf"]),
    (11, (205, 815, 318, 830), "Masa 2,78 kg", ["masa"]),
    (12, (205, 829, 320, 844), "Volumen 6 dm³", ["vol"]),
    (13, (616, 636, 698, 718), "Detalle 2 (1:1): bordón del cuerpo, cúpula por fuera — labio ≈ 5,8, escalón ≈ 5,2", ["det2"]),
    (14, (515, 745, 597, 827), "Detalle 3 (1:1): fondo encastrado con pestaña ≥ 10, filete por debajo", ["det3"]),
])
pt = 100 / 72  # el 10 kg tiene capa de texto: cajas tomadas en pt con search
rec("F04", "Recipiente_10kg_R2_sin_logo.pdf", "G733", "10kg", [
    (1, tuple(v * pt for v in (186, 26, 214, 44)), "Ø36,8 — Ø exterior del cuello", ["cuello_d"]),
    (2, tuple(v * pt for v in (81, 20, 99, 33)), "14 — altura del cuello", ["cuello_h"]),
    (3, tuple(v * pt for v in (81, 83, 99, 122)), "62,5 ± 0,3 — altura de cúpula", ["hd"]),
    (4, tuple(v * pt for v in (306, 304, 324, 348)), "562,5 ± 0,8 — altura total", ["total"]),
    (5, tuple(v * pt for v in (71, 525, 89, 563)), "16,5 ± 0,2 — pollera del fondo cóncavo", ["zf_borde"]),
    (6, tuple(v * pt for v in (225, 608, 243, 625)), "4,8 — flecha del fondo cóncavo", ["zf_centro"]),
    (7, tuple(v * pt for v in (176, 606, 224, 623)), "Ø181,5 ± 0,5 — Ø exterior", ["D"]),
    (8, tuple(v * pt for v in (395, 279, 431, 296)), "M30x1,5 — rosca del cuello", ["rosca"]),
    (9, tuple(v * pt for v in (393, 342, 432, 360)), "Ø35 + 0,3 — asiento de la válvula", ["asiento"]),
    (10, tuple(v * pt for v in (289, 623, 406, 635)), "Espesor de cúpula 1,6", ["td"]),
    (11, tuple(v * pt for v in (291, 634, 405, 645)), "Espesor de cuerpo 2", ["t"]),
    (12, tuple(v * pt for v in (293, 644, 403, 655)), "Espesor de fondo 2", ["tf"]),
    (13, tuple(v * pt for v in (311, 653, 384, 664)), "Masa 5,48 kg", ["masa"]),
    (14, tuple(v * pt for v in (304, 663, 392, 675)), "Volumen 12,5 dm³", ["vol"]),
])
rec("F05", "Recipiente 25kg R3.pdf", "G690", "25kg", [
    (1, (287, 60, 322, 77), "Ø37 — Ø exterior del cuello", ["cuello_d"]),
    (2, (129, 94, 146, 118), "35 — altura del cuello", ["cuello_h"]),
    (3, (129, 140, 146, 185), "87 ± 0,3 — altura del cabezal superior", ["hd"]),
    (4, (129, 415, 146, 475), "493 ± 0,8 — largo del cuerpo entre cabezales", ["hc"]),
    (5, (451, 398, 470, 458), "696,5 ± 0,8 — altura total del recipiente", ["total"]),
    (6, (129, 700, 146, 745), "87 ± 0,3 — altura del cabezal inferior", ["hd"]),
    (7, (255, 780, 352, 796), "Ø276,5 ± 0,5 — Ø exterior", ["D"]),
    (8, (561, 440, 610, 455), "M30x1,5 — rosca del cuello", ["rosca"]),
    (9, (561, 508, 612, 522), "Ø35 + 0,5 — asiento de la válvula", ["asiento"]),
    (10, (430, 851, 552, 865), "Espesor de chapa 3,2", ["t", "td", "tf"]),
    (11, (440, 864, 545, 878), "Masa 15,92 kg", ["masa"]),
    (12, (430, 878, 552, 892), "Volumen 34,4 dm³", ["vol"]),
    (13, (627, 647, 687, 707), "Detalle 2 (1:1): casquete con borde reducido y tope dentro del cuerpo", ["det2"]),
    (14, (545, 752, 605, 812), "Detalle 3 (1:1): ídem en el casquete inferior", ["det3"]),
], nota="La altura total 696,5 no es igual a la suma de sus cotas parciales (87 + 493 + 87 + 35 = 702); el plano FLAMA usa las cotas parciales.")
rec("F06", "Recipiente 50kg R2.pdf", "G689", "50kg", [
    (1, (285, 72, 318, 88), "Ø88 — Ø exterior del cuello (cupla)", ["cuello_d"]),
    (2, (127, 108, 145, 132), "37 — altura del cuello", ["cuello_h"]),
    (3, (127, 145, 145, 200), "105 ± 0,5 — altura del cabezal superior", ["hd"]),
    (4, (127, 460, 145, 490), "640 ± 2 — largo del cuerpo entre cabezales", ["hc"]),
    (5, (450, 440, 468, 484), "892 ± 0,5 — altura total del recipiente", ["total"]),
    (6, (127, 735, 145, 785), "105 ± 0,5 — altura del cabezal inferior", ["hd"]),
    (7, (265, 810, 340, 830), "Ø320 +3/−1 — Ø exterior", ["D"]),
    (8, (545, 427, 625, 447), "RBSP 2 ½\" x 11 h — rosca del cuello", ["rosca"]),
    (9, (556, 502, 608, 518), "Ø80 ± 0,5 — asiento de la válvula", ["asiento"]),
    (10, (414, 846, 486, 860), "Espesor 3,2", ["t", "td", "tf"]),
    (11, (398, 861, 522, 876), "Masa 23,3 kg", ["masa"]),
    (12, (390, 877, 535, 893), "Volumen 61,8 dm³", ["vol"]),
], nota="La altura total 892 no es igual a la suma de sus cotas parciales (105 + 640 + 105 + 37 = 887); el plano FLAMA usa las cotas parciales. "
        "Estas cotas también se usan para el 70 kg (proporción del cabezal y volumen por kg) y para el cuello de 70 y 100 kg.")

# =========================================================== planos Fadesa de extintor (masa total y otras cotas)
def ext(f, archivo, codigo, tam, marks, nota=""):
    fig(f, f"Plano Fadesa «{archivo[:-4]}» (código {codigo})", f"Plano de conjunto Fadesa, archivo «{archivo}»",
        "Página 1 (única), notas y cotas generales", pdf_img(f"{FAD}/{archivo}", 0),
        [(n, b, e, cl) for n, b, e, cl in marks], K, tipo="Plano Fadesa", nota=nota)


def pt(x0, x1, y0, y1, pad=1.5):
    """caja en pt del PDF (medida con las líneas del dibujo) -> px a 100 dpi"""
    k = 100 / 72
    return ((x0 - pad) * k, (y0 - pad) * k, (x1 + pad) * k, (y1 + pad) * k)


def piezas_fe(tam, d, valv):
    """marcas de las piezas medidas a escala; d: comp -> (caja pt, texto)"""
    out = []
    for comp, (caja, txt) in d.items():
        out.append((caja, txt, [f"FE:{tam}:{comp}"]))
    return out


def ext2(f, archivo, codigo, tam, generales, piezas, nota=""):
    marks = [(n, b, e, cl) for n, (b, e, cl) in enumerate(generales + piezas, 1)]
    fig(f, f"Plano Fadesa «{archivo[:-4]}» (código {codigo})", f"Plano de conjunto Fadesa, archivo «{archivo}»",
        "Página 1 (única): notas, cotas generales y piezas del despiece medidas a escala", pdf_img(f"{FAD}/{archivo}", 0),
        marks, K, tipo="Plano Fadesa", nota=nota)


NOTA_ESC = ("Piezas medidas sobre el dibujo vectorial a la escala del rótulo (verificada con el Ø del recipiente acotado en "
            "su plano de recipiente). Precisión ≈ ±0,5 mm.")
ext2("F07", "Extintor 1kg #U00d876 (V#U00e1lvula HZ) R1.pdf", "A114/117/122/123/138", "1 kg", [
    ((110, 902, 262, 920), "Masa total 1,84 kg (polvo) — referencia de verificación de masa", ["FE:1 kg:masa"]),
    ((370, 1055, 472, 1073), "Válvula HZ", ["FE:HZ"])], piezas_fe("1 kg", {
        "valvula": (pt(298.0, 422.3, 62.9, 131.2), "Válvula F510 (esc. 1:2,5): cuerpo 31 × 12 + horquilla 26,5; palanca 67; manija fija 49 + 27 (bajando 34,5); ancho de la válvula 109; espiga M22 Ø21 × 8,6"),
        "manometro": (pt(218.8, 254.6, 63.8, 104.4), "Manómetro F645: Ø36"),
        "oring": (pt(304.4, 332.5, 131.0, 135.0), "O-ring F656: Ø24,8 × 3,5"),
        "vastago": (pt(314.6, 322.4, 169.3, 190.7), "Vástago F915: Ø5 × 21, asiento Ø9"),
        "resorte": (pt(312.9, 324.1, 221.6, 244.2), "Resorte F139: Ø9,8 × 19,9"),
        "cano": (pt(308.1, 328.6, 266.6, 304.7), "Tubo de pesca F003/617: cabeza Ø18,1, tubo Ø14,4"),
    }, "F510"), nota=NOTA_ESC)
ext2("F08", "Extintor_2.5kg_Valvula_HZ_R1_sin_logo.pdf", "A101/139/151/128/135", "2,5 kg", [
    ((116, 902, 287, 920), "Masa total 4,62 kg (polvo)", ["FE:2,5 kg:masa"]),
    ((350, 1055, 456, 1073), "Válvula HZ", ["FE:HZ"])], piezas_fe("2,5 kg", {
        "valvula": (pt(375.6, 495.6, 69.2, 122.5), "Válvula F192 (esc. 1:3): igual a la del 5 y 10 kg"),
        "racor": (pt(302.5, 333.3, 93.1, 110.9), "Racor: tuerca Ø18,8 × 5,9, casquillo Ø16,9 × 13,5, rosca Ø13,1 × 8,9"),
        "manometro": (pt(513.6, 543.4, 117.4, 151.1), "Manómetro F645: Ø36 (35,7)"),
        "oring": (pt(378.6, 409.5, 150.6, 154.0), "O-ring F646: Ø32,8 × 3,6"),
        "vastago": (pt(387.4, 400.8, 181.7, 223.8), "Vástago F916: Ø6 × 30,6, asiento Ø14"),
        "resorte": (pt(385.2, 401.5, 251.9, 279.1), "Resorte F653: Ø17,3 × 28,8"),
        "cano": (pt(384.0, 404.1, 307.2, 320.9), "Caño de pesca F682: cabeza Ø21,3"),
        "manguera": (pt(237.5, 257.5, 200.0, 240.0), "Manguera: Ø17,4"),
        "tobera": (pt(238.2, 256.9, 306.8, 379.1), "Portatobera Ø18,8 × 13,4 + tobera Ø19,9 × 60,1"),
    }, "F192"), nota=NOTA_ESC)
ext2("F09", "Extintor_5kg_Valvula_HZ_R1_sin_logo.pdf", "A102/116/140/163/152/111/129/136", "5 kg", [
    ((106, 902, 258, 920), "Masa total 8,4 kg (polvo)", ["FE:5 kg:masa"]),
    ((370, 1055, 472, 1073), "Válvula HZ", ["FE:HZ"])], piezas_fe("5 kg", {
        "valvula": (pt(415.8, 505.8, 85.1, 125.0), "Válvula F192 (esc. 1:4): palanca 127 de largo, 41,5 de alto sobre el cuerpo; cuerpo 38 × 12,8 + horquilla 32,2; espiga M30 Ø29,8 × 10,8; cuello Ø27,2 × 4,1; manija fija 50 + 36"),
        "racor": (pt(322.5, 345.7, 101.5, 114.8), "Racor: tuerca Ø18,8 × 5,9, casquillo Ø16,9 × 13,5, rosca Ø13,1 × 8,9"),
        "manometro": (pt(365.2, 387.7, 128.8, 154.1), "Manómetro F645: Ø36 (35,7)"),
        "oring": (pt(418.0, 441.2, 153.1, 155.6), "O-ring F646: Ø32,7 × 3,6"),
        "vastago": (pt(424.6, 434.6, 183.4, 215.0), "Vástago F916: Ø6 × 30,8, asiento Ø14"),
        "resorte": (pt(423.0, 435.2, 243.1, 263.5), "Resorte F653: Ø17,3 × 28,8"),
        "cano": (pt(422.0, 437.1, 291.7, 301.9), "Caño de pesca F682: cabeza Ø21,3"),
        "manguera": (pt(273.5, 289.5, 220.0, 260.0), "Manguera: Ø17,4"),
        "tobera": (pt(274.3, 288.4, 332.2, 386.3), "Portatobera Ø19 × 13,3 + tobera Ø19,9 × 60"),
    }, "F192"), nota=NOTA_ESC + " Esta válvula F192 es la que se usa también en los manuales revendidos (BC, HCFC, agua, AFFF, Sales K, clase D, CO₂ como referencia).")
ext2("F10", "Extintor_10kg_HZ_R1_sin_logo.pdf", "A103/141/164/153/112/130/137", "10 kg", [
    ((113, 902, 270, 920), "Masa total 16,4 kg (polvo)", ["FE:10 kg:masa"]),
    ((370, 1055, 473, 1073), "Válvula HZ", ["FE:HZ"])], piezas_fe("10 kg", {
        "valvula": (pt(402.9, 493.0, 60.1, 100.1), "Válvula F192 (esc. 1:4): 127 × 41,5; espiga M30 Ø29,8 × 10,8"),
        "racor": (pt(294.6, 317.8, 74.4, 87.7), "Racor: tuerca Ø18,8, casquillo Ø16,9, rosca Ø13,1"),
        "manometro": (pt(333.6, 355.9, 119.0, 144.4), "Manómetro F645: Ø36 (35,7) × 31,5 con vástago"),
        "oring": (pt(405.2, 428.4, 119.8, 122.3), "O-ring F646: Ø32,7 × 3,6"),
        "vastago": (pt(411.8, 421.8, 141.6, 173.2), "Vástago F916: Ø6 × 30,6, asiento Ø14 × 14"),
        "resorte": (pt(410.1, 422.4, 193.0, 213.4), "Resorte F653: Ø17,3 × 28,8"),
        "cano": (pt(409.0, 424.5, 233.0, 257.3), "Caño de pesca F682: cabeza Ø21,3 × 14,4"),
        "manguera": (pt(245.5, 261.5, 200.0, 240.0), "Manguera: Ø17,4 (líneas a 12,1 pt)"),
        "tobera": (pt(246.4, 260.5, 340.2, 394.4), "Portatobera Ø18,8 × 13,4 + tobera Ø19,9 × 60,1"),
        "suncho": (pt(332.4, 352.8, 537.0, 547.2), "Suncho portamanguera F674: banda 14,4, sobresale 28,8"),
    }, "F192"), nota=NOTA_ESC)
ext2("F11", "Extintor rodante 25kg R1.pdf", "B202/204/217/225/229/235", "25 kg", [
    ((110, 902, 226, 920), "Masa total 53,2 kg — referencia de masa (calibración de llanta y sunchos)", ["FE:25 kg:masa"]),
    ((730, 450, 748, 480), "1140 — altura total (coincide con el catálogo)", ["FE:25 kg:H"]),
    ((376, 720, 392, 760), "Ø300 — rueda (coincide con el catálogo)", ["FE:25 kg:rueda"])], piezas_fe("25 kg", {
        "valvula": (pt(347.4, 407.4, 103.3, 147.8), "Válvula F192 (esc. 1:6): la misma de los manuales (cuello M30)"),
        "manometro": (pt(427.8, 445.2, 127.4, 149.4), "Manómetro F645 + cubremanómetro G711: Ø47"),
        "oring": (pt(390.4, 405.9, 157.3, 158.9), "O-ring F646: Ø32,8"),
        "vastago": (pt(394.9, 401.6, 168.1, 189.1), "Vástago F916: Ø6 × 30,6"),
        "resorte": (pt(393.7, 402.0, 198.1, 211.8), "Resorte F653: Ø17,5 × 29"),
        "cano": (pt(393.2, 403.3, 218.8, 225.6), "Caño de pesca F226: cabeza Ø21,3"),
    }, "F192"), nota=NOTA_ESC)
ext2("F12", "Extintor rodante 50kg R1.pdf", "B206/216/222/230/232/236/247", "50 kg", [
    ((110, 902, 216, 920), "Masa total 94 kg", ["FE:50 kg:masa"]),
    ((405, 440, 422, 475), "1210 — altura total (coincide con el catálogo)", ["FE:50 kg:H"]),
    ((437, 700, 453, 740), "Ø350 — rueda (coincide con el catálogo)", ["FE:50 kg:rueda"])], piezas_fe("50 kg", {
        "valvula": (pt(418.0, 451.8, 125.8, 172.3), "Válvula G763 (esc. 1:7): brida Ø83 × 12 + rosca 2½\" Ø75 × 17,6; cuerpo 49 × 67"),
        "manometro": (pt(391.0, 405.9, 133.4, 156.4), "Manómetro F644 + cubremanómetro G711: Ø47"),
        "rueda": (pt(505.0, 530.0, 438.0, 586.0), "Rueda Ø350: ancho de banda 48,9 (líneas a 19,8 pt)"),
    }, "G763"), nota=NOTA_ESC + " Esta válvula y estas ruedas se usan también en el 70 kg y en el AFFF 50 l.")
ext2("F13", "Extintor rodante 100kg R1.pdf", "B208/214/237/238/239/243/244/245", "100 kg", [
    ((114, 897, 235, 914), "Masa total 187,6 kg", ["FE:100 kg:masa"]),
    ((580, 839, 616, 853), "Ø390 — Ø exterior del recipiente de 100 kg", ["FE:100 kg:D"]),
    ((392, 450, 410, 485), "1460 — altura total del plano Fadesa (el catálogo dice 1400)", ["FE:100 kg:H"])], piezas_fe("100 kg", {
        "valvula": (pt(415.8, 445.3, 107.8, 148.6), "Válvula G763 (esc. 1:8): igual a la del 50 kg"),
        "rueda": (pt(326.0, 357.0, 458.0, 601.0), "Rueda Ø400: ancho de banda 76 (líneas a 26,9 pt); trocha 500"),
    }, "G763"), nota="Del plano de 100 kg se tomó el Ø390, la masa total y el ancho de rueda; el plano no acota el "
        "recipiente: cuerpo y cabezales del 100 kg se calcularon (figura del código _g_rodante). " + NOTA_ESC)

# =========================================================== catálogo Fadesa 3 (tablas)
CATF = "Catálogo Fadesa 3, archivo «Cat_logo_Fadesa_3.pdf» (escaneo)"


def cat(f, t, pag, cols, nota=""):
    marks = []
    n = 1
    for cod, nombre, a, b, rod in cols:
        campos = ["Peso", "Altura", "Ancho", "Profundidad"] + (["Rueda", "Manga"] if rod else [])
        marks.append((n, a, f"{nombre}: peso cargado, altura, ancho, profundidad" + (", Ø rueda, manga" if rod else ""),
                      [f"CAT:{cod}:{c}" for c in campos]))
        marks.append((n + 1, b, f"{nombre}: presión de servicio y de ensayo (MPa)", [f"CAT:{cod}:Ps", f"CAT:{cod}:Pe"]))
        n += 2
    fig(f, f"Catálogo Fadesa 3, pág. {pag}: tabla «Especificaciones»", CATF, f"Página {pag}, tabla Especificaciones",
        Image.open(f"{CAT}/{t}.png"), marks, 1.0, tipo="Catálogo Fadesa", nota=nota)


cat("F14", "t04", 4, [
    ("FL_MAT_ABC_1kg", "1 kg Ø3", (553, 95, 652, 245), (553, 330, 652, 395), False),
    ("FL_MAT_ABC_2.5kg", "2,5 kg", (848, 92, 948, 242), (848, 327, 948, 392), False),
    ("FL_MAT_ABC_5kg", "5 kg", (998, 90, 1102, 238), (998, 323, 1102, 388), False),
    ("FL_MAT_ABC_10kg", "10 kg", (1148, 92, 1252, 242), (1148, 326, 1252, 392), False)])
cat("F15", "t05", 5, [
    ("FL_MAT_ABC_25kg", "25 kg", (518, 148, 612, 358), (518, 452, 612, 512), True),
    ("FL_MAT_ABC_50kg", "50 kg", (638, 148, 738, 358), (638, 452, 738, 512), True),
    ("FL_MAT_ABC_70kg", "70 kg", (883, 148, 987, 358), (883, 452, 987, 512), True),
    ("FL_MAT_ABC_100kg", "100 kg", (1128, 148, 1228, 358), (1128, 452, 1228, 512), True)])
cat("F16", "t06", 6, [("FL_MAT_BC_5kg", "BC 5 kg", (765, 105, 875, 255), (765, 350, 875, 410), False)])
cat("F17", "t08", 8, [("FL_MAT_AGUA_10l", "Agua 10 dm³", (835, 210, 955, 352), (835, 452, 955, 508), False)],
    nota="La altura figura como 620/650: el plano usa 650.")
cat("F18", "t10", 10, [("FL_MAT_AFFF_10l", "AFFF 10 dm³", (880, 168, 995, 315), (880, 410, 995, 468), False)],
    nota="La altura figura como 620/650: el plano usa 650.")
cat("F19", "t11", 11, [("FL_MAT_AFFF_50l", "AFFF 50 dm³", (610, 190, 705, 392), (610, 488, 705, 542), True)])
cat("F20", "t13", 13, [("FL_MAT_HCFC-HFC_5kg", "HCFC 5 kg", (1018, 125, 1122, 272), (1018, 362, 1122, 418), False)],
    nota="La pág. 14 (HFC 236fa) repite los mismos valores para 5 kg.")
cat("F21", "t15", 15, [("FL_MAT_SALESK_6l", "Sales K 6 dm³", (655, 225, 762, 372), (655, 462, 762, 520), False)])
cat("F22", "t16", 16, [("FL_MAT_CO2_2kg", "CO₂ 2 kg", (582, 95, 638, 235), (582, 395, 638, 445), False),
                       ("FL_MAT_CO2_5kg", "CO₂ 5 kg", (727, 95, 782, 235), (727, 395, 782, 445), False)])
cat("F23", "t17", 17, [("FL_MAT_CLASED_9l", "Clase D (se usa la columna 10 kg)", (1060, 255, 1168, 400),
                        (1060, 492, 1168, 550), False)],
    nota="El catálogo sólo lista clase D de 5 y 10 kg; el 9 l de la lista FLAMA usa la columna de 10 kg (a confirmar).")

# =========================================================== fichas de mercado
g = pymupdf.open(f"{GEO}/8_Ficha-te#U00eccnica-ABC90.pdf")[0]
pg = 100 / 72
fig("F24", "Ficha técnica Georgia ABC 90 — 70 kg", "Ficha técnica publicada por Matafuegos Georgia "
    "(www.matafuegosgeorgia.com), archivo «8_Ficha-técnica-ABC90.pdf»", "Página 1, recuadro de datos",
    pdf_img(f"{GEO}/8_Ficha-te#U00eccnica-ABC90.pdf", 0),
    [(1, tuple(v * pg for v in (43, 669, 146, 699)), "Capacidad nominal 70 kg", ["GE:70kg:cap"]),
     (2, tuple(v * pg for v in (42, 706, 150, 738)), "Peso con carga 140 kg — referencia de masa del 70 kg", ["GE:70kg:masa"])],
    K, crop=(0, 600, 600, 842 * pg), tipo="Ficha técnica de mercado",
    nota="No hay plano Fadesa de 70 kg: la masa de referencia es la de esta ficha.")
fig("F25", "Ficha técnica Melisam — acetato de potasio 6 L", "Ficha técnica publicada por Melisam (www.melisam.com), "
    "archivo «Acetato-de-Potasio.pdf»", "Página 2, tabla de especificaciones, fila «Peso bruto»",
    pdf_img(f"{MEL}/Acetato-de-Potasio.pdf", 1),
    [(1, tuple(v * pg for v in (40, 514, 352, 531)), "Peso bruto 9,5 kg (6 L) — referencia de masa del Sales K 6 l", ["ME:6l:masa"])],
    K, crop=(0, 600, 827, 800), tipo="Ficha técnica de mercado")

# =========================================================== normas y resoluciones (sólo la cláusula usada)
fig("F26", "IRAM 3550:1981, cláusula 4.1.3.2 (espesor mínimo del recipiente)", "Norma IRAM 3550 (Dic. 1981), "
    "archivo «IRAM-3550.pdf»", "Pág. 9 de la norma (hoja 11 del PDF), apartado 4.1.3.2",
    pdf_img(f"{NOR}/1008147370-IRAM-3550.pdf", 10), [(1, (82, 612, 792, 681), "4.1.3.2 — mínimo 2,9 mm hasta Ø320 y "
    "4,5 mm por encima → chapa 4,75 en 70 y 100 kg", ["N:3550"])], K, crop=(40, 80, 800, 700), tipo="Norma / resolución")
fig("F27", "IRAM 3534:1983, apartado 2.2.4.3 d) (longitud de la leyenda: arco de 108°)", "Norma IRAM 3534 (Ago. 1983), "
    "archivo «IRAM-3534.pdf»", "Pág. 8 de la norma (hoja 8 del PDF), apartado 2.2.4.3 d)",
    pdf_img(f"{NOR}/1008147342-IRAM-3534.pdf", 7), [(1, (92, 333, 768, 360), "2.2.4.3 d) — máximo: arco de 108° → ancho del "
    "panel central de la etiqueta", ["N:3534"])], K, crop=(40, 80, 790, 372), tipo="Norma / resolución")
fig("F28", "IRAM 3525:1983, apartado 3.2.1 b) (espesor mínimo del recipiente inoxidable)", "Norma IRAM 3525 (Dic. 1983), "
    "archivo «IRAM-3525.pdf»", "Pág. 5 de la norma (hoja 5 del PDF), apartado 3.2.1 b)",
    pdf_img(f"{NOR}/959734356-IRAM-3525.pdf", 4), [(1, (155, 408, 698, 447), "3.2.1 b) — acero inoxidable de espesor no menor "
    "de 0,63 mm → se adoptó 0,8", ["N:3525"])], K, crop=(100, 120, 710, 455), tipo="Norma / resolución")
r5 = pymupdf.open(f"{RES}/RESOLUCION 522 07.pdf")
pr = 100 / 72
fig("F29", "Res. OPDS 522/07, Anexo 2 (oblea de fabricación de extintores de más de 1 kg)", "Resolución 522/07 "
    "(Provincia de Buenos Aires), archivo «RESOLUCION 522 07.pdf»", "Hoja 6 del PDF, Anexo 2, apartado «Oblea»",
    pdf_img(f"{RES}/RESOLUCION 522 07.pdf", 5),
    [(1, tuple(v * pr for v in (83, 352, 490, 368)), "Anexo 2: oblea de fabricación de extintores de más de 1 kg", ["N:522"]),
     (2, tuple(v * pr for v in (450, 425, 530, 441)), "«formato de 46…»", ["N:522"]),
     (3, tuple(v * pr for v in (83, 441, 203, 455)), "«…milímetros de diámetro» → oblea Ø46", ["N:522"])],
    K, crop=(60, 440, 790, 680), tipo="Norma / resolución",
    nota="El Anexo 1 (hoja 4 del PDF) fija el mismo formato de 46 mm para la oblea de 1 kg.")
fig("F30", "IRAM Anexo R — estampilla de conformidad de extintor nuevo", "IRAM, «Anexo R — Protección contra "
    "incendios» (DC-PG-129), archivo «Anexo_R-Proteccion_contra_Incendios.pdf»", "Hoja 3 del PDF, «Conformidad con "
    "norma IRAM para la certificación de extintores nuevos»", pdf_img(f"{RES}/Anexo_R-Proteccion_contra_Incendios.pdf", 2),
    [(1, tuple(v * 100 / 70 for v in (52, 268, 400, 290)), "Estampilla para la certificación de extintores nuevos", ["N:AnexoR"]),
     (2, tuple(v * 100 / 70 for v in (58, 358, 292, 517)), "Estampilla: proporción ≈ 3:2 → se adoptó 60 × 40 mm (sin cota en "
      "el documento; a confirmar con IRAM)", ["N:AnexoR"])], K, crop=(40, 330, 790, 760), tipo="Norma / resolución")

# =========================================================== relevamiento de mercado (fotos)
def foto(f, archivo, titulo, marks, nota=""):
    im = Image.open(f"{SP}/fotos/{archivo}")
    im = im.resize((600, int(im.height * 600 / im.width)))
    fig(f, titulo, f"Relevamiento de mercado FLAMA 2026, foto «{archivo}»", "Foto completa", im,
        [(n, b, e, cl) for n, b, e, cl in marks], 1.0, tipo="Relevamiento de mercado", nota=nota)


foto("F31", "f01.jpg", "Relevamiento: tarjeta AGC (CABA) en un extintor de 10 kg",
     [(1, (100, 295, 505, 500), "Tarjeta de identificación AGC autoadhesiva con QR → 140 × 55 mm", ["RV:tarjeta"])],
     nota="Medida estimada por proporción con el Ø del recipiente de 10 kg (181,5 mm); confirmar con la AGC.")
foto("F32", "f08.jpg", "Relevamiento: etiqueta envolvente y estampilla IRAM (línea Georgia)",
     [(1, (150, 222, 462, 525), "Etiqueta envolvente: panel central + alas (≈ 216° en total)", ["RV:etiqueta"]),
      (2, (250, 540, 470, 670), "Estampilla IRAM con n° vertical, logo y QR → 60 × 40 mm", ["RV:estampilla"])],
     nota="Medidas estimadas por proporción con el Ø del recipiente; confirmar.")
foto("F33", "f13.jpg", "Relevamiento: etiqueta de n° de serie GS1 (línea Melisam)",
     [(1, (230, 360, 358, 447), "Etiqueta GS1 con QR «(01)779…» → 45 × 25 mm", ["RV:serie"])],
     nota="Medida estimada por proporción; confirmar.")
foto("F34", "f19.jpg", "Relevamiento: faja de garantía rayada (línea de etiquetado Georgia)",
     [(1, (60, 508, 268, 672), "Faja de garantía rayada roja/blanca → 30 × 40 mm", ["RV:faja"])],
     nota="Medida estimada por proporción; confirmar.")
foto("F35", "f15.jpg", "Relevamiento: oblea circular PBA (DPS) aplicada",
     [(1, (80, 212, 467, 612), "Oblea circular de la Provincia de Buenos Aires (ver Res. 522/07: Ø46)", ["N:522"])])


# =========================================================== planilla MP FLAMA (render de celdas reales)
def planilla(f, hoja, filas, cols, marcas, titulo, nota=""):
    from openpyxl import load_workbook
    wb = load_workbook("/home/user/planos_industrial_flama/referencias/MP_Abastecimiento_Almacenamiento_FLAMA_OpcionD.xlsx",
                       data_only=True)
    ws = wb[hoja]
    fnt = ImageFont.truetype(MONO, 15)
    fb = ImageFont.truetype(FONT, 15)
    wcol = {c: (330 if c == "A" else 92) for c in cols}
    W = 46 + sum(wcol.values()); H = 30 + 28 * len(filas)
    im = Image.new("RGB", (W, H), "white"); d = ImageDraw.Draw(im)
    x = 46; xs = {}
    d.rectangle([0, 0, W, 30], fill=(230, 230, 230))
    for c in cols:
        xs[c] = x; d.rectangle([x, 0, x + wcol[c], 30], outline=(160, 160, 160)); d.text((x + wcol[c] / 2, 15), c, fill="black", font=fb, anchor="mm")
        x += wcol[c]
    ys = {}
    for i, r in enumerate(filas):
        y = 30 + 28 * i; ys[r] = y
        d.rectangle([0, y, 46, y + 28], fill=(230, 230, 230), outline=(160, 160, 160)); d.text((23, y + 14), str(r), fill="black", font=fb, anchor="mm")
        for c in cols:
            v = ws[f"{c}{r}"].value
            s = "" if v is None else (f"{v:g}" if isinstance(v, (int, float)) else str(v))
            mx = int(wcol[c] / 9.2)
            if len(s) > mx:
                s = s[:mx - 1] + "…"
            d.rectangle([xs[c], y, xs[c] + wcol[c], y + 28], outline=(200, 200, 200))
            d.text((xs[c] + 5, y + 14), s.replace(".", ","), fill="black", font=fnt, anchor="lm")
    mk = []
    for n, celdas, et, cl in marcas:
        c0, r0 = re.match(r"([A-Z]+)(\d+)", celdas[0]).groups(); c1, r1 = re.match(r"([A-Z]+)(\d+)", celdas[-1]).groups()
        mk.append((n, (xs[c0] + 2, ys[int(r0)] + 2, xs[c1] + wcol[c1] - 2, ys[int(r1)] + 26), et, cl))
    fig(f, titulo, "Planilla FLAMA «MP_Abastecimiento_Almacenamiento_FLAMA_OpcionD.xlsx»", f"Hoja «{hoja}», celdas marcadas "
        "(reproducción de los valores de las celdas)", im, mk, 1.0, tipo="Planilla MP FLAMA", nota=nota)


planilla("F36", "Geometria Acero", [30, 31, 32, 33, 34, 35, 37, 38, 43, 44], list("ABCDEFGHIJ"), [
    (1, ["C35", "C35"], "C35: caño del 1 kg cortado a 255 mm", ["MP:C35"]),
    (2, ["D34", "H35"], "D34:H35: desarrollo (fila 34) y alto (fila 35) del recorte del cuerpo 2,5 a 50 kg", [f"MP:{c}{r}" for c in "DEFGH" for r in (34, 35)]),
    (3, ["C38", "F38"], "C38:F38: Ø del disco de cúpula 1 a 10 kg", [f"MP:{c}38" for c in "CDEF"]),
    (4, ["C44", "F44"], "C44:F44: Ø del disco de fondo 1 a 10 kg", [f"MP:{c}44" for c in "CDEF"]),
    (5, ["J32", "J32"], "J32: espesor de chapa 4,75 de 70 / 100 kg", ["MP:J32"]),
], "Planilla MP FLAMA — geometría del acero por recipiente",
    nota="Las filas 34-35 de 70 y 100 kg (1212 × 680 / 900) NO se usan: no corresponden al plano (ver BOM). "
         "Las filas 37 y 43 del 1 kg y del 10 kg difieren de los planos (cúpula 10 kg 1,6; fondo 1 kg 1,25): corregir la planilla.")
def planilla_v(f, hoja, fila, cols, fila_cab, marcas, titulo, nota=""):
    from openpyxl import load_workbook
    wb = load_workbook("/home/user/planos_industrial_flama/referencias/MP_Abastecimiento_Almacenamiento_FLAMA_OpcionD.xlsx",
                       data_only=True)
    ws = wb[hoja]
    fnt = ImageFont.truetype(MONO, 15); fb = ImageFont.truetype(FONT, 15)
    W = 1300; H = 30 + 30 * len(cols)
    im = Image.new("RGB", (W, H), "white"); d = ImageDraw.Draw(im)
    d.rectangle([0, 0, W, 30], fill=(230, 230, 230))
    for x, t in ((40, "Celda"), (130, f"Encabezado (fila {fila_cab})"), (560, "Valor")):
        d.text((x, 15), t, fill="black", font=fb, anchor="lm")
    ys = {}
    for i, c in enumerate(cols):
        y = 30 + 30 * i; ys[c] = y
        d.rectangle([0, y, W, y + 30], outline=(200, 200, 200))
        d.text((40, y + 15), f"{c}{fila}", fill="black", font=fb, anchor="lm")
        d.text((130, y + 15), str(ws[f"{c}{fila_cab}"].value or "")[:46], fill=(80, 80, 80), font=fnt, anchor="lm")
        d.text((560, y + 15), str(ws[f"{c}{fila}"].value or "")[:80], fill="black", font=fnt, anchor="lm")
    mk = [(n, (36, ys[c0] + 2, W - 4, ys[c1] + 28), et, cl) for n, (c0, c1), et, cl in marcas]
    fig(f, titulo, "Planilla FLAMA «MP_Abastecimiento_Almacenamiento_FLAMA_OpcionD.xlsx»", f"Hoja «{hoja}», fila {fila} "
        "(reproducción de los valores de las celdas)", im, mk, 1.0, tipo="Planilla MP FLAMA", nota=nota)


planilla_v("F37", "MP Abastecimiento", 25, list("ABCDEF"), 4, [
    (1, ("A", "B"), "MP-22: varilla de refuerzo de 70 / 100 kg", ["MP:VARILLA"]),
    (2, ("C", "C"), "«Barra de 6 m (medida a definir por cálculo)» → se adoptó Ø8 (decisión FLAMA, sin cálculo)", ["MP:VARILLA"]),
], "Planilla MP FLAMA — MP-22 varilla de refuerzo")


# =========================================================== código FLAMA (valores de diseño propios)
def codigo(f, archivo, rangos, resalt, claves, titulo, nota=""):
    lines = open(f"{REPO}/flama/{archivo}", encoding="utf-8").read().split("\n")
    fnt = ImageFont.truetype(MONO, 15)
    rows = []
    for a, b in rangos:
        if rows:
            rows.append((None, "…"))
        for i in range(a, b + 1):
            rows.append((i, lines[i - 1]))
    W = 1120; H = 10 + 21 * len(rows)
    im = Image.new("RGB", (W, H), (250, 250, 250)); d = ImageDraw.Draw(im)
    marks = []
    for j, (i, s) in enumerate(rows):
        if i in resalt:
            y = 5 + 21 * j
            d.rectangle([60, y, W - 54, y + 21], fill=(255, 242, 204))
    for j, (i, s) in enumerate(rows):
        y = 5 + 21 * j
        d.text((8, y + 2), "" if i is None else f"{i:4d}", fill=(130, 130, 130), font=fnt)
        d.text((66, y + 2), s if len(s) <= 104 else s[:103] + "…", fill=(20, 20, 20), font=fnt)
    # una marca por bloque de líneas resaltadas contiguas
    n = 1
    j = 0
    while j < len(rows):
        if rows[j][0] in resalt:
            k = j
            while k + 1 < len(rows) and rows[k + 1][0] in resalt:
                k += 1
            lin = f"línea {rows[j][0]}" if rows[j][0] == rows[k][0] else f"líneas {rows[j][0]}–{rows[k][0]}"
            lin += ": " + " ⏎ ".join(rows[q][1].strip() for q in range(j, k + 1))[:220]
            marks.append((n, (58, 5 + 21 * j - 1, W - 52, 5 + 21 * k + 22), lin, claves))
            n += 1
            j = k + 1
        else:
            j += 1
    a0, b0 = rangos[0][0], rangos[-1][1]
    fig(f, titulo, f"Código del generador de planos FLAMA, «flama/{archivo}» (commit {SHA[:7]})",
        f"Líneas {a0}–{b0} (resaltado: los valores tomados)", im, marks, 1.0,
        link=f"{GH}{archivo}#L{a0}-L{b0}", tipo="Diseño FLAMA (supuesto)", nota=nota, der=True)


def L(archivo, patron):
    for i, s in enumerate(open(f"{REPO}/flama/{archivo}", encoding="utf-8"), 1):
        if patron in s:
            return i
    raise KeyError(patron)


def HL(archivo, pats):
    return {L(archivo, p) for p in pats}


# Los valores de diseño FLAMA no llevan captura de código: van en la tabla del Word / Excel con el link a la línea.

for f in figs:
    im = Image.open(f["png"]).convert("RGB")
    if f["tipo"] in ("Catálogo Fadesa", "Relevamiento de mercado"):
        q = f["png"][:-4] + ".jpg"; im.save(q, "JPEG", quality=82, optimize=True)
    else:
        q = f["png"][:-4] + "_q.png"
        im.quantize(colors=48, method=Image.Quantize.MEDIANCUT, dither=Image.Dither.NONE).save(q, optimize=True)
    f["img"] = q
json.dump(figs, open(f"{SP}/jw/figs.json", "w"), ensure_ascii=False, indent=1)
print(len(figs), "figuras")
