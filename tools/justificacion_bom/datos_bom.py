"""Datos del Word «Justificación del BOM», leídos del propio FLAMA_BOM.xlsx ya recalculado: así el Word y el Excel
salen de las mismas fórmulas. Agrega los controles que no están en el Excel (gas real con CoolProp, masa contra la
referencia de plano o ficha) y los textos (textos.py).

Uso: python datos_bom.py <FLAMA_BOM.xlsx recalculado> <salida.json> [cache.pkl]
  El libro tiene que tener los valores calculados (abrirlo y guardarlo, o pasarlo por LibreOffice:
  soffice --headless --convert-to xlsx --outdir <carpeta> FLAMA_BOM.xlsx).
"""
import json
import os
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
from flama import bom as B                       # noqa: E402
from flama import sustituto as SU                # noqa: E402
from flama.catalogo import MODELOS               # noqa: E402
import textos                                    # noqa: E402

try:
    from CoolProp.CoolProp import PropsSI
except ImportError:                              # sin CoolProp no se hace el control de gas real
    PropsSI = None

XLSX, OUT = sys.argv[1], sys.argv[2]
CACHE = __import__("pickle").load(open(sys.argv[3], "rb")) if len(sys.argv) > 3 else None   # faja del sustituto
wb = load_workbook(XLSX, data_only=True)
FLUIDO = {"GAS_N2": "Nitrogen", "GAS_AIRE": "Air", "GAS_AR": "Argon"}
M_REAL = {"GAS_N2": 28.014, "GAS_AIRE": 28.965, "GAS_AR": 39.948}       # g/mol


def r(x, n=4):
    return round(x, n) if isinstance(x, (int, float)) else x


def pct(a, b):
    return None if not b or a is None else round(100 * (a - b) / b, 1)


# ------------------------------------------------------------------ REFERENCIAS
ws = wb["REFERENCIAS"]
refs, ref = [], {}
for fila in ws.iter_rows(min_row=3, values_only=True):
    if fila[1] and str(fila[1]).startswith("Tolerancia general de los planos"):
        break
    if fila[0] and fila[1] and fila[0] != "ID":
        d = dict(id=fila[0], dato=fila[1], valor=r(fila[2], 6), min=r(fila[3], 6), max=r(fila[4], 6), unidad=fila[5],
                 fuente=fila[6], nota=fila[7] or "")
        refs.append(d)
        ref[d["id"]] = d
    elif not fila[0] and fila[1] and not fila[2]:
        refs.append(dict(seccion=fila[1]))
iso = [dict(desde=f[1], hasta=f[2], tol=f[3]) for f in ws.iter_rows(min_row=3, values_only=True)
       if not f[0] and isinstance(f[1], (int, float)) and isinstance(f[3], (int, float))]

# ------------------------------------------------------------------ MODELOS
ws = wb["MODELOS"]
cab = [c.value for c in ws[2]]
modelos = {}
for fila in ws.iter_rows(min_row=3, values_only=True):
    if fila[0]:
        modelos[fila[0]] = {k: r(v, 6) for k, v in zip(cab, fila)}
gases = []
for gid in ("GAS_N2", "GAS_AIRE", "GAS_AR"):
    rho_n = ref[gid]["valor"]
    real = M_REAL[gid] / 1000 / ref["VM"]["valor"]
    gases.append(dict(id=gid, dato=ref[gid]["dato"], rho_n=rho_n, rho_real=r(real, 6), dif=pct(rho_n, real)))
T, PATM = ref["T_REF"]["valor"], ref["P_ATM"]["valor"]
cargas = []
for m in MODELOS:
    x = modelos[m.codigo]
    d = dict(codigo=m.codigo, nombre=m.nombre, x=x)
    if x["Ref. gas"] and PropsSI is not None and x["V libre (L)"]:
        real = PropsSI("D", "T", T, "P", (x["Ps (MPa)"] + PATM) * 1e6, FLUIDO[x["Ref. gas"]]) * x["V libre (L)"] / 1000
        d.update(gas_real=r(real, 5), dif_real=pct(x["Gas (kg)"], real))
    cargas.append(d)


# ------------------------------------------------------------------ hojas del BOM
def hoja(nombre):
    ws = wb[nombre]
    cab = [c.value for c in ws[1]]
    return [dict(zip(cab, f)) for f in ws.iter_rows(min_row=2, values_only=True) if f[1]]


H = {n: hoja(n) for n in ("Fabricados", "Revendidos", "Cilindros", "Recargas")}
POR_COD = {n: {f["Código"]: f for f in filas} for n, filas in H.items()}

masas = []
for hojanom in ("Fabricados", "Revendidos"):
    for m in MODELOS:
        f = POR_COD[hojanom].get(m.codigo)
        if not f:
            continue
        sop = next((g for c, g in POR_COD[hojanom].items() if c.endswith("-S8") and
                    c.split("-")[0] == B.P.codigo_pieza(m, 0)[:-3].split("-")[0]), None)
        s = [sop.get(k) or 0 for k in ("Masa total (kg)", "Masa mín. (kg)", "Masa máx. (kg)")] if sop else [0, 0, 0]
        tot, lo, hi = f["Masa total (kg)"], f["Masa mín. (kg)"], f["Masa máx. (kg)"]
        carg, cmin, cmax = tot - s[0], lo - s[1], hi - s[2]
        ref_p, fuente = B.peso_ref(m)
        fuente = "Catálogo Fadesa 3 (orientativo)" if fuente.startswith("Catálogo") else fuente
        masas.append(dict(codigo=m.codigo, nombre=m.nombre, propio=hojanom == "Fabricados", total=r(tot, 3),
                          soporte=r(s[0], 3), bom=r(carg, 2), bom_min=r(cmin, 2), bom_max=r(cmax, 2),
                          tol_m=pct(cmin, carg), tol_p=pct(cmax, carg), ref=ref_p, fuente=fuente,
                          dif=pct(carg, ref_p), dentro=cmin <= ref_p <= cmax,
                          catalogo=float(m.spec["Peso cargado (kg)"].replace(",", ".")),
                          sin_masa=sorted({g["Pieza"] for g in H[hojanom] if g["Nivel"] >= 2 and not g["Ref."] and
                                           g["Masa total (kg)"] is None and g["Código"].startswith(
                                               B.P.codigo_pieza(m, 0)[:-3].split("-")[0] + "-")})))

CLAVE = ("Cuerpo", "Cúpula", "Fondo", "Cuello roscado", "Válvula HZ armada", "Manguera", "Soporte de pared",
         "Soporte vehicular")
piezas = []
for m in MODELOS:
    if m.codigo not in B.PROPIOS:
        continue
    base = B.P.codigo_pieza(m, 0)[:-3].split("-")[0] + "-"
    fila = {"codigo": m.codigo}
    for f in H["Fabricados"]:
        if not f["Código"].startswith(base) or f["Nivel"] != 2:
            continue
        for c in CLAVE:
            if f["Pieza"].startswith(c) and c not in fila and f["Masa total (kg)"]:
                mt = f["Masa total (kg)"]
                fila[c] = dict(masa=r(mt, 3), tol=r(100 * (f["Masa máx. (kg)"] / mt - 1), 1),
                               e=r(f["Espesor eq. (mm)"], 2), tiso=f["Tol. ISO 2768-m (± mm)"],
                               tvol=r(100 * (f["Tol. volumen (± %)"] or 0), 1), ref=f["Ref."])
    piezas.append(fila)

cilindros = []
for f in H["Cilindros"]:
    if f["Nivel"] == 0:
        cilindros.append(dict(codigo=f["Código"], nombre=f["Pieza"], masa=r(f["Masa total (kg)"], 3),
                              min=r(f["Masa mín. (kg)"], 3), max=r(f["Masa máx. (kg)"], 3)))

# ejemplos de tolerancia de pieza (para la sección de tolerancias)
ejemplos = []
for cod in ("ABC5-01", "ABC5-02", "ABC5-07", "ABC5-17", "ABC70-01", "ABC70-09", "ABC5-20"):
    for hn in ("Fabricados",):
        f = POR_COD[hn].get(cod)
        if f and f["Masa total (kg)"]:
            ejemplos.append(dict(codigo=cod, pieza=f["Pieza"], material=f["Material / agente"], ref=f["Ref."],
                                 vu=r(f["Vol. unit. (L)"], 5), sup=r(f["Superficie (dm²)"], 3),
                                 e=r(f["Espesor eq. (mm)"], 2), tiso=f["Tol. ISO 2768-m (± mm)"],
                                 tvol=r(100 * (f["Tol. volumen (± %)"] or 0), 1), rho=f["Densidad (kg/L)"],
                                 rmin=f["ρ mín. (kg/L)"], rmax=f["ρ máx. (kg/L)"], mt=r(f["Masa total (kg)"], 4),
                                 mmin=r(f["Masa mín. (kg)"], 4), mmax=r(f["Masa máx. (kg)"], 4)))

sustitutos = []
for m in MODELOS:
    filas = CACHE["sus"][m.codigo] if CACHE else SU.bom_sustituto(m)
    faja = next((x for x in filas if x["codigo"].endswith("-01")), None)
    eq = SU.equivalencia(m)[3].replace("igual clasificación e igual capacidad (mismo modelo)", "").strip("; ")
    alto = float(faja["med"].split(" × ")[0]) if faja else 0
    sustitutos.append(dict(codigo=m.codigo, nombre=m.nombre, faja=faja["med"].split(" (")[0] if faja else "",
                           nota="; ".join(x for x in [eq, "alto < 40: lo limita la oblea PBA" if alto < 40 else ""]
                                          if x)))

# capturas reusadas del Word de planos (masas de referencia) y capturas propias de normas y fichas
COMP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "justificacion")
figs = {x["id"]: x for x in json.load(open(os.path.join(COMP, "compacto.json")))["figs"]}
reusadas = []
for fid in ("F07", "F08", "F09", "F10", "F11", "F12", "F13", "F24", "F25", "F21"):
    x = figs[fid]
    ext = os.path.splitext(x["crop"])[1]
    reusadas.append(dict(id=fid, titulo=x["titulo"], tipo=x["tipo"], png=os.path.join("..", "justificacion", "figs",
                                                                                          fid + "_c" + ext),
                         w=x["cw"], h=x["ch"], marcas=x["marcas"]))
capturas = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "capturas.json")))

recargas = []
for m in MODELOS:
    x = modelos[m.codigo]
    rk = m.codigo.replace("FL_MAT_", "RK_")
    serv = [(f["Pieza"], r(f["Masa total (kg)"], 4), r(f["Masa mín. (kg)"], 4), r(f["Masa máx. (kg)"], 4))
            for f in H["Recargas"] if f["Nivel"] == 1 and f["Código"].startswith(rk + "-")]
    recargas.append(dict(codigo=m.codigo, nombre=m.nombre, x=x, servicios=serv))

D = dict(referencias=refs, iso=iso, modelos=modelos, gases=gases, cargas=cargas, masas=masas, piezas=piezas,
         cilindros=cilindros, ejemplos=ejemplos, sustitutos=sustitutos, recargas=recargas, capturas=capturas,
         reusadas=reusadas)
D["notas"] = textos.notas(D)
json.dump(D, open(OUT, "w"), ensure_ascii=False, indent=1, default=str)
print("ok", OUT, len(cargas), "modelos", len(refs), "referencias")
